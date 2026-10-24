"""Самопроверка установки MINT.

    python check_install.py              # всё, включая короткий запрос к языковой модели
    python check_install.py --no-llm     # без запроса к языковой модели
    python check_install.py --skip-docling --skip-page

Проверяет версию Python, пакеты, разбор тестового .docx (Docling), файл .env, связь с языковой моделью,
эмбеддингами и реранкером, папки базы и логов, настройки Streamlit и в конце открывает страницу MINT на
настоящем Streamlit без браузера (смоук-тест: страница рисуется без ошибок). Настройки и документы не меняет,
ключи не печатает — вывод можно целиком переслать разработчику.
"""
import argparse
import importlib
import json
import os
import platform
import sys
import tempfile
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from mint.i18n import N_, default_lang, set_lang, tr  # noqa: E402  (язык вывода: MINT_LANG или язык системы)
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(errors="replace")   # консоль Windows в cp866/cp1251 не падает на символах
    except Exception:
        pass

MIN_PYTHON = (3, 10)
COUNTS = {"ok": 0, "warn": 0, "fail": 0}
MARKS = {"ok": "[ OK ]", "warn": "[ !! ]", "fail": "[FAIL]", "info": "[ .. ]"}


def say(status, title, detail=""):
    if status in COUNTS:
        COUNTS[status] += 1
    print(f"{MARKS[status]} {title}" + (f"\n         {detail}" if detail else ""), flush=True)


def section(title):
    print(f"\n=== {title} ===", flush=True)


ZPS_HINT = (N_("Похоже, включена замкнутая программная среда (ЗПС): библиотеки из pip не подписаны и не "
               "загружаются. Это решает администратор (подпись библиотек или исключение) — обходить ЗПС не "
               "нужно."))


def zps_enabled():
    """Astra Linux: ЗПС включена, если в digsig_initramfs.conf стоит DIGSIG_ELF_MODE=1."""
    conf = Path("/etc/digsig/digsig_initramfs.conf")
    try:
        return any(line.strip().replace(" ", "") == "DIGSIG_ELF_MODE=1"
                   for line in conf.read_text(encoding="utf-8", errors="replace").splitlines())
    except OSError:
        return False


def check_system():
    section(tr("Система"))
    say("info", tr("ОС: {v}", v=platform.platform()))
    astra = Path("/etc/astra_version")
    if astra.is_file():
        say("info", f"Astra Linux: {astra.read_text(encoding='utf-8', errors='replace').strip()}")
        if zps_enabled():
            say("warn", tr("ЗПС включена"), tr(ZPS_HINT))
    py = sys.version_info
    in_venv = sys.prefix != getattr(sys, "base_prefix", sys.prefix)
    detail = f"{sys.executable}" + (tr(" (виртуальное окружение)") if in_venv else "")
    if py[:2] >= MIN_PYTHON:
        say("ok", f"Python {platform.python_version()}", detail)
    else:
        say("fail", tr("Python {v} — нужен {v2}.{v3} или новее", v=platform.python_version(), v2=MIN_PYTHON[0], v3=MIN_PYTHON[1]),
            tr("Astra Linux 1.7 поставляется с Python 3.7: нужна Astra 1.8 (Python 3.11) или отдельно "
               "установленный Python. В Windows — Python 3.11/3.12 с python.org (вариант «для текущего "
               "пользователя»)."))
        return False
    return True


PACKAGES = [
    # модуль, пакет pip, обязателен, зачем
    ("streamlit", "streamlit", True, N_("интерфейс")),
    ("docling", "docling-slim", True, N_("разбор .docx")),
    ("faiss", "faiss-cpu", True, N_("векторный индекс")),
    ("numpy", "numpy", True, ""),
    ("requests", "requests", True, N_("запросы к API")),
    ("docx", "python-docx", True, N_("замечания в Word")),
    ("pdfplumber", "pdfplumber", False, N_("разбор .pdf (без моделей)")),
    ("yaml", "PyYAML", True, N_("промпты моделей и config.yaml")),
    ("dotenv", "python-dotenv", False, N_("без него .env читается встроенным разбором — это нормально")),
    ("truststore", "truststore", False, N_("корпоративные сертификаты из хранилища ОС (важно для Windows)")),
]


def _version(dist):
    try:
        from importlib.metadata import version
        return version(dist)
    except Exception:
        return "?"


def check_packages():
    section(tr("Пакеты Python"))
    ok = True
    for mod, dist, required, why in PACKAGES:
        try:
            importlib.import_module(mod)
        except Exception as e:   # ImportError, а при ЗПС — OSError «Operation not permitted»
            msg = f"{type(e).__name__}: {e}"
            install = '"docling-slim[format-docx]"' if dist == "docling-slim" else dist
            hint = tr(ZPS_HINT) if ("not permitted" in msg or "failed to map segment" in msg) else \
                tr("Установите: pip install {dist}", dist=install) + (f" ({tr(why)})" if why else "")
            say("fail" if required else "warn", tr("{dist} не загружается", dist=dist), f"{msg[:300]}\n         {hint}")
            ok = ok and not required
            continue
        ver = _version(dist)
        if dist == "docling-slim" and ver == "?":          # полный Docling вместо облегчённой сборки — тоже годится
            dist, ver = "docling", _version("docling")
        if dist == "python-docx" and ver != "?" and tuple(int(x) for x in ver.split(".")[:2]) < (1, 2):
            say("warn", f"python-docx {ver}", tr("Нужен 1.2+, иначе замечания в Word вставляются в текст, а не примечаниями: pip install -U "
                                                 "python-docx"))
            continue
        if dist == "truststore" and ver == "?":
            continue
        say("ok", f"{dist} {ver}")
    try:
        import torch
        say("info", tr("torch {version} установлен — его тянет Docling, MINT он не нужен: .docx и .pdf разбираются без "
                       "моделей", version=torch.__version__))
    except Exception:
        say("info", tr("torch не установлен — MINT он не нужен"))
    return ok


def _make_docx(path):
    import docx
    d = docx.Document()
    d.add_heading("1. Общие положения", level=1)
    d.add_paragraph("Тестовый документ самопроверки MINT: шифрование передаваемых данных обязательно.")
    t = d.add_table(rows=2, cols=2)
    for r, row in enumerate([["Показатель", "Значение"], ["Доступность", "99,9 %"]]):
        for c, val in enumerate(row):
            t.cell(r, c).text = val
    d.save(path)


def check_docling():
    section(tr("Разбор .docx (Docling)"))
    try:
        from mint.parsing import compile_docx_to_markdown
    except Exception as e:
        say("fail", tr("модуль разбора MINT не загружается"), f"{type(e).__name__}: {e}")
        return
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "проверка.docx"
        try:
            _make_docx(path)
        except Exception as e:
            say("fail", tr("не удалось создать тестовый .docx (python-docx)"), f"{type(e).__name__}: {e}")
            return
        t0 = time.time()
        try:
            md = compile_docx_to_markdown(path.name, path.read_bytes())
        except Exception as e:
            msg = f"{type(e).__name__}: {e}"
            hint = ""
            low = msg.lower()
            if any(w in low for w in ("huggingface", "hf_hub", "connection", "resolve", "timed out", "proxy")):
                hint = (tr("Docling пытается скачать модели из интернета. Скачайте их на машине с доступом (docling-tools "
                           "models download), перенесите папку и укажите её в DOCLING_ARTIFACTS_PATH."))
            elif "not permitted" in low or "failed to map segment" in low:
                hint = tr(ZPS_HINT)
            say("fail", tr("Docling не разобрал тестовый документ"), msg[:500] + (f"\n         {hint}" if hint else ""))
            return
    dt = time.time() - t0
    if "Общие положения" in md and "Доступность" in md:
        say("ok", tr("тестовый .docx разобран за {dt:.1f} с", dt=dt),
            tr("заголовок и таблица на месте") + (tr(" (первый запуск обычно дольше)") if dt > 20 else ""))
    else:
        say("warn", tr("Docling вернул неожиданный текст"), md[:300].replace("\n", " | "))


def _nearest_existing(p):
    p = Path(p)
    while not p.exists() and p != p.parent:
        p = p.parent
    return p


def check_settings_and_dirs():
    section(tr("Настройки (.env, config.yaml)"))
    try:
        from mint.config import check_settings, config_file_path, env_file_candidates, load_settings
        from mint.services import apply_tls_policy
    except Exception as e:
        say("fail", tr("настройки MINT не загружаются"), f"{type(e).__name__}: {e}")
        return None, None
    cfg, notes = load_settings()
    if cfg.env_files:
        files = cfg.env_files if isinstance(cfg.env_files, str) else "; ".join(cfg.env_files)
        say("ok", tr(".env прочитан: {files}", files=files), tr("разбор: {env_parser}", env_parser=cfg.env_parser))
    else:
        say("warn", tr(".env не найден"), tr("Искал: ") + "; ".join(str(p) for p in env_file_candidates())
            + tr("\n         Скопируйте .env.example в .env рядом с app.py и заполните (имя файла — ровно «.env»)."))
    cfgp = config_file_path()
    say("info", f"config.yaml: {cfgp if cfgp.is_file() else tr('нет (используются значения по умолчанию)')}")
    for n in notes:
        say("warn", f"config.yaml: {n}")
    for level, msg in check_settings(cfg):
        say("fail" if level == "error" else "warn", msg)
    say("info", tr("LLM: {v}  модель: {v2}", v=cfg.llm_url or '—', v2=cfg.llm_model or tr("(первая из списка сервера)")))
    say("info", tr("Эмбеддинги: {v}  модель: {embed_model}", v=cfg.embed_url or '—', embed_model=cfg.embed_model))
    say("info", tr("Реранкер: {v}", v=cfg.rerank_url or tr("не подключён")) + (tr("  модель: {rerank_model}", rerank_model=cfg.rerank_model) if cfg.rerank_url else ""))
    say("info", tr("Ключ LLM_API_KEY: ") + (tr("задан") if cfg.llm_api_key else tr("не задан (нормально для серверов без авторизации)")))
    say("info", apply_tls_policy(cfg))

    section(tr("Папки"))
    for title, p in ((tr("База знаний (KB_DIR)"), cfg.kb_path), (tr("Логи и отзывы (LOG_DIR)"), cfg.log_path)):
        base = _nearest_existing(p)
        if os.access(base, os.W_OK):
            say("ok", f"{title}: {p}", "" if p.exists() else tr("папки ещё нет — MINT создаст её при первой записи"))
        else:
            say("fail", tr("{title}: {p} — нет прав на запись", title=title, p=p), tr("Проверьте права на {base} или укажите другую папку в .env", base=base))
    kb_dim = None
    manifest = cfg.kb_path / "manifest.json"
    if manifest.is_file():
        try:
            m = json.loads(manifest.read_text(encoding="utf-8"))
            say("info", tr("В базе: документов {v}, векторов {v2}, размерность {v3}, записана MINT {v4}", v=len(m.get('documents') or {}), v2=m.get('total_vectors'), v3=m.get('embed_dim'), v4=m.get('mint_version', '?')))
            kb_dim = m.get("embed_dim")
        except Exception as e:
            say("warn", tr("manifest.json базы не читается"), str(e))
    return cfg, kb_dim


def check_services(cfg, ask, kb_dim=None):
    section(tr("Связь с моделями"))
    import requests
    from mint.audit import is_llm_error
    from mint.limits import detect_server_info
    from mint.services import (_auth_headers, _embed_input, _parse_embedding_json, ask_llm, describe_http_status,
                               describe_request_error, probe_reranker)

    if cfg.llm_url:
        info = detect_server_info(cfg)
        if info.get("llm_models"):
            models = info["llm_models"]
            say("ok", tr("LLM: сервер отвечает, моделей: {n_models}", n_models=len(models)), ", ".join(models[:8]) + (" …" if len(models) > 8 else ""))
            if cfg.llm_model and cfg.llm_model not in models:
                say("warn", tr("модели {llm_model} нет в списке сервера", llm_model=cfg.llm_model), tr("Проверьте LLM_MODEL в .env"))
        else:
            say("fail" if info.get("llm_error") else "warn", tr("LLM: список моделей не получен"),
                info.get("llm_error") or tr("Сервер не отдаёт /models — это допустимо, если ответ модели ниже работает."))
        if ask:
            t0 = time.time()
            ans = ask_llm("Answer with one word.", "Answer with one word: does it work?", cfg, max_tokens=300)
            if is_llm_error(ans):
                say("fail", tr("LLM: ответ не получен"), ans.strip()[:400])
            else:
                say("ok", tr("LLM ответила за {v:.1f} с", v=time.time() - t0), ans.strip()[:80])

    if cfg.embed_url:
        try:
            r = requests.post(cfg.embed_url, json={"input": _embed_input("проверка связи / connection check", cfg), "model": cfg.embed_model},
                              headers=_auth_headers(cfg.embed_key), verify=cfg.tls_verify, timeout=30)
            if r.status_code != 200:
                say("fail", tr("Эмбеддинги: HTTP {status_code}", status_code=r.status_code), describe_http_status(r.status_code, cfg.embed_url))
            else:
                vec = _parse_embedding_json(r.json())
                if not vec:
                    say("fail", tr("Эмбеддинги: ответ без вектора"), r.text[:200])
                else:
                    say("ok", tr("Эмбеддинги: вектор размерности {n_vec}", n_vec=len(vec)))
                    if kb_dim and kb_dim != len(vec):
                        say("warn", tr("база построена другой моделью эмбеддингов (размерность {kb_dim})", kb_dim=kb_dim),
                            tr("Верните прежнюю EMBED_MODEL или переиндексируйте базу в интерфейсе."))
        except Exception as e:
            say("fail", tr("Эмбеддинги: нет связи"), describe_request_error(e, cfg.embed_url))

    code, desc = probe_reranker(cfg)
    if code is None:
        say("info", tr("Реранкер: {desc}", desc=desc))
    elif code == 200:
        say("ok", tr("Реранкер: {desc}", desc=desc))
    else:
        say("warn", tr("Реранкер: {desc}", desc=desc), tr("Поиск работает и без реранкера, но точность ниже."))


def check_streamlit_config():
    section(tr("Запуск"))
    conf = HERE / ".streamlit" / "config.toml"
    text = conf.read_text(encoding="utf-8", errors="replace") if conf.is_file() else ""
    if 'address = "localhost"' in text:
        say("ok", tr("Streamlit слушает только этот компьютер (localhost)"), str(conf))
    else:
        say("warn", tr('нет .streamlit/config.toml с address = "localhost"'),
            tr("Без него страница MINT доступна по сети всем, кто знает адрес компьютера."))
    if 'gatherUsageStats = false' not in text:
        say("warn", tr("статистика Streamlit не отключена"), tr("В .streamlit/config.toml: [browser] gatherUsageStats = false"))
    if Path.cwd().resolve() != HERE:
        say("info", tr("Запускайте MINT из папки {HERE} (run_mint.sh / run_mint.bat) — иначе Streamlit не увидит "
                       ".streamlit/config.toml", HERE=HERE))


def check_page():
    """Смоук-тест: страница MINT целиком на настоящем Streamlit, без браузера (streamlit.testing)."""
    section(tr("Пробный запуск страницы"))
    try:
        from streamlit.testing.v1 import AppTest
    except Exception as e:
        say("info", tr("пропущено: в этой установке нет streamlit.testing"), f"{type(e).__name__}: {str(e)[:200]}")
        return
    t0 = time.time()
    try:
        at = AppTest.from_file(str(HERE / "app.py"), default_timeout=180)
        at.run()
    except Exception as e:
        say("fail", tr("страница не запустилась"), f"{type(e).__name__}: {str(e)[:500]}")
        return
    if at.exception:
        exc = at.exception[0]
        trace = "\n         ".join(str(x) for x in (getattr(exc, "stack_trace", None) or [])[-8:])
        say("fail", tr("страница открылась с ошибкой"), f"{str(exc.message)[:500]}\n         {trace}")
        return
    titles = " ".join(str(t.value) for t in at.title)
    try:
        n_tabs = len(at.tabs)
    except Exception:
        n_tabs = "?"
    say("ok", tr("страница открывается без ошибок ({v:.1f} с)", v=time.time() - t0), tr("{titles}; вкладок: {n_tabs}", titles=titles, n_tabs=n_tabs))


def main():
    try:
        from mint.config import _load_env_files
        _load_env_files()                    # MINT_LANG может быть задан в .env
    except Exception:
        pass
    set_lang(default_lang())
    ap = argparse.ArgumentParser(description=tr("Самопроверка установки MINT"))
    ap.add_argument("--no-llm", action="store_true", help=tr("не отправлять тестовый запрос языковой модели"))
    ap.add_argument("--skip-docling", action="store_true", help=tr("не проверять разбор .docx"))
    ap.add_argument("--skip-page", action="store_true", help=tr("не открывать страницу MINT для проверки"))
    args = ap.parse_args()
    try:
        from mint.version import MINT_VERSION
    except Exception:
        MINT_VERSION = "?"
    print(tr("Самопроверка MINT {MINT_VERSION} — {v}", MINT_VERSION=MINT_VERSION, v=time.strftime('%Y-%m-%d %H:%M')))
    import logging
    mint_log = logging.getLogger("MINT")
    quiet = logging.NullHandler()
    quiet._mint_handler = True      # журнал MINT не открывается и не пишет в консоль во время проверки
    mint_log.addHandler(quiet)
    mint_log.setLevel(logging.ERROR)
    if check_system():
        pkgs_ok = check_packages()
        if not args.skip_docling:
            check_docling()
        cfg, kb_dim = check_settings_and_dirs() if pkgs_ok else (None, None)
        if cfg is not None and (cfg.llm_url or cfg.embed_url):
            check_services(cfg, ask=not args.no_llm, kb_dim=kb_dim)
        check_streamlit_config()
        if pkgs_ok and cfg is not None and not args.skip_page:
            check_page()
    print(tr("\nИтог: в порядке {ok}, предупреждений {warn}, ошибок {fail}.", ok=COUNTS['ok'], warn=COUNTS['warn'], fail=COUNTS['fail']))
    if COUNTS["fail"]:
        print(tr("Исправьте ошибки по подсказкам выше. Если непонятно — перешлите этот вывод разработчику "
                 "(ключей в нём нет)."))
    else:
        print(tr("Можно запускать: run_mint.sh (Linux) или run_mint.bat (Windows)."))
    sys.exit(1 if COUNTS["fail"] else 0)


if __name__ == "__main__":
    main()
