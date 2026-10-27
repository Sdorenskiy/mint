"""Одинаково ли MINT разбирает .docx разными способами — бэкендом Word из Docling напрямую (по умолчанию, без
PyTorch) и через DocumentConverter (как до 21.1), — и в разных установках (облегчённый Docling и полный).

Сравнивается всё, что MINT берёт из документа: Markdown для аудита (document_to_markdown) и структура для базы
знаний (parse_document: текстовые блоки с путём заголовков, таблицы, сведения о разборе).

    python tools/docx_parity.py папка_с_docx                       # оба способа здесь и сейчас: одинаково ли
    python tools/docx_parity.py папка_с_docx --dump out/A          # сохранить разбор (для сравнения установок)
    python tools/docx_parity.py --compare out/A out/B              # сравнить два сохранённых разбора

Документы никуда не отправляются; в отчёт попадают только имена файлов и, при расхождении, несколько строк
различий — запускайте у себя, отчёт с различиями передавайте только туда, куда можно передавать сами документы.
"""
import argparse
import difflib
import hashlib
import importlib.metadata as md
import json
import os
import sys
import traceback
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def _version(dist):
    try:
        return md.version(dist)
    except md.PackageNotFoundError:
        return None


def environment():
    return {"python": sys.version.split()[0], "docling": _version("docling"), "docling-slim": _version("docling-slim"),
            "docling-core": _version("docling-core"), "torch": _version("torch")}


def parse_one(path, parser):
    """Разбор одного файла способом parser («backend» или «converter»): {markdown, structure} или {error}."""
    os.environ["MINT_DOCX_PARSER"] = parser
    from mint.parsing import document_to_markdown, parse_document
    data = Path(path).read_bytes()
    try:
        markdown = document_to_markdown(Path(path).name, data)
        doc = parse_document(Path(path).name, data)
        structure = {"units": doc["units"], "tables": doc["tables"], "info": doc["info"],
                     "structured": doc["structured"], "pages_real": doc["pages_real"]}
        return {"markdown": markdown, "structure": json.loads(json.dumps(structure, ensure_ascii=False, default=str))}
    except Exception as e:
        return {"error": f"{type(e).__name__}: {e}", "trace": traceback.format_exc()[-3000:]}


def _digest(result):
    return hashlib.sha256(json.dumps(result, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()[:12]


def _diff(a, b, n=12):
    """Несколько первых строк различий Markdown или структуры."""
    if "error" in a or "error" in b:
        return [f"A: {a.get('error', 'ok')}", f"B: {b.get('error', 'ok')}"]
    for key in ("markdown", "structure"):
        if a[key] != b[key]:
            la = a[key].splitlines() if key == "markdown" else json.dumps(a[key], ensure_ascii=False, indent=1).splitlines()
            lb = b[key].splitlines() if key == "markdown" else json.dumps(b[key], ensure_ascii=False, indent=1).splitlines()
            lines = [ln for ln in difflib.unified_diff(la, lb, "A", "B", n=0, lineterm="") if not ln.startswith("@@")]
            return [f"[{key}]"] + lines[2:2 + n]
    return []


def make_samples(folder):
    """Пара .docx с тем, что бывает в регламентах и ТЗ: уровни заголовков, нумерованные и маркированные списки,
    таблица с объединёнными ячейками, текст на русском. Для проверки в CI вместе с тестовыми документами Docling."""
    import docx
    out = Path(folder)
    out.mkdir(parents=True, exist_ok=True)
    d = docx.Document()
    d.add_heading("Требования к сервису обмена", 0)
    d.add_heading("1. Общие положения", 1)
    d.add_paragraph("Документ описывает сервис обмена файлами между подсистемами.")
    d.add_heading("1.1 Хранение", 2)
    for s in ("Файлы хранятся 30 дней.", "Резервные копии — ежедневно.", "Журналы — не менее трёх лет."):
        d.add_paragraph(s, style="List Number")
    d.add_heading("1.2 Участники", 2)
    for s in ("Оператор портала", "Администратор платформы интеграции"):
        d.add_paragraph(s, style="List Bullet")
    d.add_heading("2. Таблица операций", 1)
    tb = d.add_table(rows=4, cols=3)
    tb.style = "Table Grid"
    for i, row in enumerate([("Шаг", "Исполнитель", "Результат"), ("1", "Оператор", "Файл загружен"),
                             ("2", "Платформа", "Файл передан"), ("3", "Получатель", "Файл принят")]):
        for j, v in enumerate(row):
            tb.cell(i, j).text = v
    tb.cell(3, 1).merge(tb.cell(3, 2))
    d.add_paragraph("Таблица 1 — Операции обмена")
    d.save(out / "mint_sample_requirements.docx")
    d2 = docx.Document()
    d2.add_paragraph("Документ без стилей заголовков: только обычные абзацы.")
    d2.add_paragraph("2.1 Нумерованная строка как заголовок")
    d2.add_paragraph("Текст пункта 2.1.")
    d2.save(out / "mint_sample_plain.docx")


def files_of(paths):
    out = []
    for p in map(Path, paths):
        out += sorted(p.rglob("*.docx")) if p.is_dir() else [p]
    return [f for f in out if not f.name.startswith("~$")]


def run_here(paths, dump=None):
    """Оба способа в этой установке. Возвращает (строки отчёта, есть ли расхождения или ошибки)."""
    env = environment()
    lines = [f"Установка: {json.dumps(env, ensure_ascii=False)}", ""]
    parsers = ["backend", "converter"]
    bad = False
    results = {}
    for f in files_of(paths):
        res = {p: parse_one(f, p) for p in parsers}
        results[f.name] = res
        a, b = res["backend"], res["converter"]
        if "error" in a:
            bad = True
            lines.append(f"❌ {f.name}: бэкенд — {a['error']}")
        elif "error" in b:
            lines.append(f"ℹ️ {f.name}: бэкенд — ok; DocumentConverter — {b['error']}")
        elif a == b:
            lines.append(f"✅ {f.name}: одинаково ({_digest(a)})")
        else:
            bad = True
            lines.append(f"⚠️ {f.name}: РАЗЛИЧАЕТСЯ")
            lines += ["    " + ln for ln in _diff(b, a)]
    lines += ["", f"torch загружен в этом процессе: {'да' if 'torch' in sys.modules else 'нет'}"]
    if dump:
        out = Path(dump)
        out.mkdir(parents=True, exist_ok=True)
        (out / "environment.json").write_text(json.dumps({**env, "torch_loaded": "torch" in sys.modules},
                                                         ensure_ascii=False, indent=1), encoding="utf-8")
        for name, res in results.items():
            (out / f"{name}.json").write_text(json.dumps(res, ensure_ascii=False, indent=1), encoding="utf-8")
    return lines, bad


def compare_dumps(dir_a, dir_b, parser_a="backend", parser_b="backend"):
    """Два сохранённых разбора (например, облегчённый Docling и полный): одинаково ли."""
    a_dir, b_dir = Path(dir_a), Path(dir_b)
    env = lambda d: json.dumps(json.loads((d / "environment.json").read_text(encoding="utf-8")), ensure_ascii=False)  # noqa: E731
    lines = [f"A: {dir_a} ({parser_a}) — {env(a_dir)}", f"B: {dir_b} ({parser_b}) — {env(b_dir)}", ""]
    bad = False
    names = sorted({p.name for p in a_dir.glob("*.docx.json")} | {p.name for p in b_dir.glob("*.docx.json")})
    for name in names:
        pa, pb = a_dir / name, b_dir / name
        if not (pa.exists() and pb.exists()):
            bad = True
            lines.append(f"❌ {name}: только в {'A' if pa.exists() else 'B'}")
            continue
        a = json.loads(pa.read_text(encoding="utf-8"))[parser_a]
        b = json.loads(pb.read_text(encoding="utf-8"))[parser_b]
        a.pop("trace", None), b.pop("trace", None)
        if a == b and "error" not in a:
            lines.append(f"✅ {name[:-5]}: одинаково ({_digest(a)})")
        else:
            bad = True
            lines.append(f"⚠️ {name[:-5]}: РАЗЛИЧАЕТСЯ" if "error" not in a and "error" not in b else f"❌ {name[:-5]}")
            lines += ["    " + ln for ln in _diff(a, b)]
    return lines, bad


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("paths", nargs="*", help="файлы .docx или папки с ними")
    ap.add_argument("--dump", help="сохранить разбор в эту папку")
    ap.add_argument("--compare", nargs=2, metavar=("A", "B"), help="сравнить два сохранённых разбора")
    ap.add_argument("--parsers", default="backend,backend", help="для --compare: способ в A и в B (backend/converter)")
    ap.add_argument("--make-samples", metavar="DIR", help="создать в папке примеры .docx и выйти")
    args = ap.parse_args()
    if args.make_samples:
        make_samples(args.make_samples)
        return
    if args.compare:
        pa, pb = args.parsers.split(",")
        lines, bad = compare_dumps(*args.compare, parser_a=pa, parser_b=pb)
    elif args.paths:
        lines, bad = run_here(args.paths, args.dump)
    else:
        ap.error("укажите файлы .docx или --compare")
    print("\n".join(lines))
    print("\nИтог:", "есть расхождения или ошибки" if bad else "одинаково")
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
