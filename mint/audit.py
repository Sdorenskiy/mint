"""Аудит документа: нормы из базы, черновик замечаний, самокритика, заключение, Word.

Весь аудит документа без интерфейса — run_audit(); им пользуется вкладка «Аудит документа».
"""

import copy
import io
import logging
import re
import threading
import time
from collections import Counter, defaultdict
from dataclasses import replace

import numpy as np

try:
    import docx
    from docx.shared import Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.text.run import Run as DocxRun
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

from .bm25 import BM25_STOP_WORDS, bm25_tokenize, normalize_token
from .citations import check_citations
from .config import FALLBACK_LIMITS
from .i18n import N_, tr
from .limits import AUDIT_PROMPT_KEYS, AUDIT_RESERVE_CHARS, LIMIT_SAFETY
from .prompts import (
    audit_map_prompt, autonomous_note, card_prompt, checklist_prompt, critique_system, document_block,
    focus_answer_prompt, generic_audit_focus, location_line, no_norms_note, norms_block, outline_block, pinned_block,
    prompt_extra_chars, prompt_text, prompts_version, render, basis_block,
)
from .basis import basis_materials, basis_references, find_basis_documents, reference_label
from .kb import kb_default_top_k
from .search import search_parent_child
from .deep_checks import (build_deep_tasks, check_formatting, conflict_tasks, find_processes, limit_tasks,
                          process_summary, run_tasks)
from .expert import CHECKLIST_MAX, load_expert_materials, query_vector, select_pinned
from .parallel import parallel_map
from .sections import (
    document_blocks, document_outline, part_path_at, section_label, split_document_parts, whole_document_part,
)
from .services import (
    answer_label, ask_llm, model_reasons, rate_limit_stats, reasoning_blocked_by_extra_body, reasoning_on,
    record_answers,
)
from .text import compact_markdown_tables, extract_json, norm_for_match, quote_in_text, smart_truncate
from .version import MINT_VERSION

logger = logging.getLogger("MINT")


_SEV_RE = re.compile(r"\*{0,2}\[(КРИТИЧНЫЙ|СРЕДНИЙ|НИЗКИЙ|CRITICAL|HIGH|MEDIUM|LOW)\]\*{0,2}", re.IGNORECASE)

_LIST_MARKER_RE = re.compile(r"^\s*(?:[-–—•*·]+|\d{1,2}[.)](?=\s))\s*")

_SEVERITIES = (N_("КРИТИЧНЫЙ"), N_("СРЕДНИЙ"), N_("НИЗКИЙ"))     # внутренние коды критичности (показ — через tr)

# Английские ответы модели (промпты на английском) приводятся к тем же кодам
_SEVERITY_ALIASES = {"CRITICAL": "КРИТИЧНЫЙ", "HIGH": "КРИТИЧНЫЙ", "MEDIUM": "СРЕДНИЙ", "LOW": "НИЗКИЙ"}

_FINDING_HEADER_WORDS = ("аудит главы", "выявленные нарушения", "section review", "identified violations")

CLEAN_PHRASES = ("часть чистая", "section is clean")


def normalize_severity(value, default="СРЕДНИЙ"):
    """«критичный», «CRITICAL», «High» → внутренний код («КРИТИЧНЫЙ» / «СРЕДНИЙ» / «НИЗКИЙ»)."""
    v = str(value or "").strip().upper().replace("Ё", "Е")
    v = _SEVERITY_ALIASES.get(v, v)
    return v if v in _SEVERITIES else default


def severity_label(code):
    """Код критичности для показа на языке интерфейса: «КРИТИЧНЫЙ» → «CRITICAL» в английском интерфейсе."""
    return tr(code)


def is_llm_error(text):
    """ask_llm возвращает ошибки обычной строкой, начинающейся с «❌»."""
    return (text or "").lstrip().startswith("❌")


def is_clean_report(text):
    low = (text or "").lower()
    return any(p in low for p in CLEAN_PHRASES)


def strip_list_marker(text):
    """Убирает маркер списка («- », «• », «1. »), не трогая цифры в начале текста («152-ФЗ …»)."""
    return _LIST_MARKER_RE.sub("", text or "", count=1).strip()


def _indent(line):
    return len(line) - len(line.lstrip())


def _is_continuation(s, indent, prev, prev_indent):
    """Подпункт замечания, а не новое замечание: предыдущее кончается двоеточием («не описано:»), строка
    с бо́льшим отступом или начинается со строчной буквы («сроки/этапы перехода;»)."""
    body = strip_list_marker(s)
    return (prev.rstrip().endswith(":") or indent > prev_indent
            or (bool(body) and body[0].isalpha() and body[0].islower()))


def parse_findings(draft):
    """Разбирает черновик замечаний на отдельные пункты (строки со знаком «-», «•», «*» или «1.»).
    Если у замечаний есть теги критичности, пункт списка без тега — подпункт предыдущего замечания (список
    внутри замечания): он присоединяется к замечанию, а не становится отдельным «обрывком». Строка без знака
    списка, но с тегом в начале («шаг 3 → шаг 4: [СРЕДНИЙ] …» — так пишут некоторые модели) — тоже замечание.
    Строки делятся по любому переводу строки (модель может прислать \\r или разделитель строк Unicode)."""
    out, indents, in_list = [], [], False
    tagged_draft = bool(_SEV_RE.search(draft or ""))
    for line in (draft or "").splitlines():
        s = line.strip()
        if not s or any(w in s.lower() for w in _FINDING_HEADER_WORDS):
            continue
        is_item = (s.startswith(("-", "•", "*", "–", "—")) or re.match(r"^\d{1,2}[.)]\s", s)
                   or (tagged_draft and _SEV_RE.search(s[:TAG_NEAR_START])))
        if (tagged_draft and out and is_item and not _SEV_RE.search(s)
                and (in_list or _is_continuation(s, _indent(line), out[-1], indents[-1]))):
            sep = " " if out[-1].rstrip().endswith(":") else "; "
            out[-1] = out[-1].rstrip().rstrip(";") + sep + strip_list_marker(s)
            in_list = True                       # начатый список подпунктов продолжается до следующего тега
            continue
        if is_item:
            out.append(s)
            indents.append(_indent(line))
            in_list = False
    if not out and (draft or "").strip() and not is_clean_report(draft):
        out = [draft.strip()]
    unique, seen = [], set()
    for item in out:                     # модель зациклилась — одно и то же замечание много раз
        key = norm_for_match(_SEV_RE.sub(" ", strip_list_marker(item)))
        if key and key not in seen:
            seen.add(key)
            unique.append(item)
    return unique


TAG_NEAR_START = 80        # тег критичности не дальше этого от начала строки — строка начинает замечание
SECTION_FINDINGS_MAX = 12   # замечаний из одного ответа проверки части (весь документ одним запросом — WHOLE_FINDINGS_MAX)
WHOLE_FINDINGS_MAX = 40
SECTION_PROMPT_FINDINGS = 10    # «не больше N замечаний» в запросе проверки части: модель обычно доходит до лимита
WHOLE_PROMPT_FINDINGS = 30      # то же, когда весь документ проверяется одним запросом
FOCUS_FINDINGS_MAX = 10     # из ответа на пункт «особого внимания»
DEEP_FINDINGS_MAX = 6       # из ответа на узкий вопрос глубокой проверки
# по видам вопросов: модель заполняет шаблон — на каждый переход по замечанию на каждый из 4 вопросов промпта,
# хотя по делу обычно одно-два; оставляются самые критичные
DEEP_KIND_MAX = {"step": 2, "requirement": 2, "boundary": 3, "conflict": 2}


def cap_findings(findings, limit):
    """Не больше limit замечаний из одного ответа: сначала критичные, затем средние и низкие; порядок ответа
    сохраняется. → (оставленные, сколько отброшено)."""
    if len(findings) <= limit:
        return findings, 0
    order = sorted(range(len(findings)), key=lambda i: (_RANK.get(findings[i]["severity"], 1), i))
    keep = set(order[:limit])
    return [f for i, f in enumerate(findings) if i in keep], len(findings) - limit


def finding_from_line(line, part=None):
    """Строка «- [КРИТИЧНЫЙ] текст» → замечание {severity, text, quote, part}."""
    m = _SEV_RE.search(line or "")
    severity = normalize_severity(m.group(1)) if m else "СРЕДНИЙ"
    text = strip_list_marker(_SEV_RE.sub(" ", line or "", count=1))
    text = re.sub(r"^[\s*:—–-]+(?=\S)", "", text).strip()
    text = re.sub(r"[ \t]{2,}", " ", text)        # «шаг 3 → шаг 4:   текст» — лишние пробелы после метки
    return {"severity": severity, "text": text, "quote": "", "part": part}


def format_finding(f):
    return f"- [{severity_label(f['severity'])}] {f['text']}"


def _critique_user_content(section_text, norms_text, numbered, special, doc_block=""):
    norms_part = render("critique_norms", norms=norms_text) if norms_text else prompt_text("critique_no_norms")
    special_part = render("critique_special", special=special.strip()) + "\n\n" if (special or "").strip() else ""
    return render("critique_user", section=section_text, norms=norms_part, special=special_part, findings=numbered,
                  document=doc_block or "")


CRITIQUE_BATCH = 15     # замечаний в одном запросе самокритики


def critique_findings(section_text, norms_text, findings, cfg, special="", doc_block=""):
    """Самокритика с доказательствами: вердикт по КАЖДОМУ замечанию с опорой на главу документа и нормы, а также
    критичность и дословная цитата из главы (цитата проверяется по тексту). doc_block — документ целиком или
    оглавление (как при проверке раздела): самокритика снимает «не описано», если это есть в другом разделе.
    Возвращает (оставленные, отклонённые). При сбое или неразборчивом ответе замечания остаются как в черновике."""
    if not findings:
        return findings, []
    if len(findings) > CRITIQUE_BATCH:     # длинный список — пачками: ответ на всё сразу обрывается и не разбирается
        kept, dropped = [], []
        for i in range(0, len(findings), CRITIQUE_BATCH):
            k, d = critique_findings(section_text, norms_text, findings[i:i + CRITIQUE_BATCH], cfg, special, doc_block)
            kept += k
            dropped += d
        return kept, dropped
    numbered = "\n".join(f"{i}. {f['text']}" for i, f in enumerate(findings, 1))
    user = _critique_user_content(section_text, norms_text, numbered, special, doc_block)
    raw = ask_llm(critique_system(), user, cfg, guard=True)
    if is_llm_error(raw):
        return findings, []
    data = extract_json(raw, kind="array")
    if not isinstance(data, list):
        return findings, []

    verdicts = {}
    for d in data:
        if isinstance(d, dict):
            try:
                verdicts[int(d.get("n"))] = d
            except (TypeError, ValueError):
                pass

    kept, dropped = [], []
    for i, f in enumerate(findings, 1):
        v = verdicts.get(i)
        if v is None:  # модель пропустила замечание — оставляем как есть
            kept.append(f)
            continue
        if str(v.get("verdict", "")).strip().upper().startswith(("УДАЛ", "DROP", "REMOVE", "DELETE", "REJECT")):
            reason = str(v.get("reason") or "").strip() or tr("без пояснения")
            dropped.append(f"{f['text']} — {reason}")
            continue
        quote = str(v.get("quote") or "").strip().strip("«»\"")
        kept.append({
            **f,
            "severity": normalize_severity(v.get("severity"), default=f["severity"]),
            "text": strip_list_marker(str(v.get("text") or "").strip()) or f["text"],
            "quote": quote if quote_in_text(quote, section_text) else "",
        })
    return kept, dropped


SPECIAL_FOCUS_NORMS = 2     # сколько фрагментов норм по «особому вниманию» добавляется к каждому разделу


def special_focus_norms(special, cfg, glossary=None):
    """Один поиск по тексту «особого внимания» на весь аудит: до SPECIAL_FOCUS_NORMS фрагментов норм, которые
    добавляются к нормам каждого раздела. При низкой уверенности поиска ничего не добавляется (чтобы не шуметь)."""
    if not (special or "").strip():
        return []
    diag = {}
    try:
        items, srcs = search_parent_child(special.strip(), cfg, k_neighbors=SPECIAL_FOCUS_NORMS, return_items=True,
                                          glossary=glossary, diagnostics=diag)
    except Exception as e:
        logger.warning(f"Поиск норм по «особому вниманию» не удался: {str(e)[:120]}")
        return []
    if not srcs or isinstance(items, str) or diag.get("low_conf"):
        return []
    return list(items)[:SPECIAL_FOCUS_NORMS]


def build_norms_context_multi(queries, cfg, glossary=None, extra_items=None, per_query=2, skip_ids=()):
    """Нормы для документа целиком: поиск по каждому разделу (per_query лучших), по очереди — сначала лучшие
    каждого раздела; дедупликация по parent_id, лимит audit_context_max_chars. skip_ids — фрагменты, которые
    модель уже получает (правила экспертов и карточки систем). → (текст норм, источники)."""
    found = []
    for q in queries:
        try:
            items, _srcs = search_parent_child(smart_truncate(q, max_chars=cfg.audit_query_chars), cfg,
                                               k_neighbors=per_query, return_items=True, glossary=glossary)
        except Exception as e:
            logger.warning(f"Поиск нормативного контекста не удался: {str(e)[:120]}")
            continue
        if not isinstance(items, str):
            found.append(list(items or []))
    ordered = list(extra_items or []) + [lst[r] for r in range(per_query) for lst in found if r < len(lst)]
    blocks, shorts, seen, total = [], [], set(skip_ids or ()), 0
    for it in ordered:
        ids = set(it.get("ids") or [it["id"]])
        if ids <= seen:
            continue
        if blocks and total + len(it["block"]) > cfg.audit_context_max_chars:
            continue                      # длинный не влез — короткий следующего раздела ещё может
        seen |= ids
        blocks.append(it["block"])
        shorts.append(it["short"])
        total += len(it["block"])
    return "\n\n".join(blocks), shorts


def build_norms_context(section_text, cfg, k, glossary=None, extra_items=None, skip_ids=()):
    """ОДИН поиск по базе на раздел документа (или замечание). Дедупликация по parent_id, лимит объёма.
    extra_items — фрагменты по «особому вниманию»: идут первыми; skip_ids — фрагменты, которые модель уже
    получает (правила экспертов и карточки систем). Возвращает (текст_норм, список_источников); пустой
    текст — нормативный контекст не найден."""
    query = smart_truncate(section_text, max_chars=cfg.audit_query_chars)
    try:
        items, srcs = search_parent_child(query, cfg, k_neighbors=k, return_items=True, glossary=glossary)
    except Exception as e:
        logger.warning(f"Поиск нормативного контекста не удался: {str(e)[:120]}")
        items, srcs = [], []
    if isinstance(items, str):
        items = []
    items = list(extra_items or []) + list(items or [])
    if not items:
        return "", []
    blocks, shorts, seen, total = [], [], set(skip_ids or ()), 0
    for it in items:
        ids = set(it.get("ids") or [it["id"]])     # склеенный блок — несколько parent-фрагментов
        if ids <= seen:
            continue
        seen |= ids
        if blocks and total + len(it["block"]) > cfg.audit_context_max_chars:
            break
        blocks.append(it["block"])
        shorts.append(it["short"])
        total += len(it["block"])
    return "\n\n".join(blocks), shorts


# ---------- Пункты документа: замечание о другой части ----------
_CLAUSE_LINE_RE = re.compile(r"^[ \t]*(?:#{1,6}[ \t]+)?\**(\d+(?:\.\d+)*)\.?\**[ \t]")
_TOC_TAIL_RE = re.compile(r"(?:\t|\.{3,}|…)\s*\d{1,3}\**\s*$")
_FINDING_CLAUSE_RE = re.compile(r"^\W*(?:(?:пп?|пункт\w*|подраздел\w*|раздел\w*|clause|section)\.?\s*)?"
                                r"(\d+(?:\.\d+)*)\.?(?=[\s,:;)]|$)", re.IGNORECASE)
_PATH_NO_RE = re.compile(r"^(\d+(?:\.\d+)*)\.?\s")


def text_clauses(text):
    """Номера пунктов, с которых начинаются строки текста: заголовки и нумерованные абзацы («2.2.4.1. …»), без
    строк содержания («2.1 Название … 7»)."""
    out = set()
    for line in (text or "").split("\n"):
        if _TOC_TAIL_RE.search(line):
            continue
        m = _CLAUSE_LINE_RE.match(line)
        if m:
            out.add(m.group(1))
    return out


def part_clauses(part):
    """Пункты одной части документа: из её текста и из путей заголовков, с которых она начинается."""
    own = text_clauses(part.get("text", ""))
    for _start, path in part.get("spans") or []:
        for comp in (path or "").split(" → "):
            m = _PATH_NO_RE.match(comp.strip() + " ")
            if m:
                own.add(m.group(1))
    return own


def finding_clause(text):
    """Пункт документа в начале замечания: «2.2.3, шаг 4: …» → "2.2.3", «п. 4.3: …» → "4.3"; иначе None."""
    m = _FINDING_CLAUSE_RE.match(text or "")
    return m.group(1) if m else None


def foreign_clause(text, own, known):
    """Пункт ДРУГОЙ части документа, с которого начинается замечание, или None. Пункт относится к той части, где
    стоит самый глубокий известный документу номер, совпадающий с ним или объемлющий его («2.2.4.2» → «2.2.4.2»
    или «2.2.4»); часть, где есть его подпункты, тоже своя («2.2» для части с «2.2.3»). Номер, которого в документе
    нет (пункт нормы, год), — без отсева; у части без своих номеров (содержание, «Введение») чужие все пункты."""
    c = finding_clause(text)
    if not c:
        return None
    owner = max((k for k in known if c == k or c.startswith(k + ".")), key=len, default=None)
    if owner is None or owner in own or any(o.startswith(owner + ".") for o in own):
        return None
    return c


def pinned_ids(pinned):
    return {i for it in pinned or [] for i in (it.get("ids") or [it["id"]])}


def norms_with_pinned(pinned, found_text, found_sources):
    """Документы-основания, правила экспертов и карточки систем (pinned) + найденные нормы → (промпт норм, текст
    норм, источники). Документы-основания (basis) — своим блоком: «документ должен им соответствовать»."""
    basis_text = "\n\n".join(it["block"] for it in pinned or [] if it.get("basis"))
    pinned_text = "\n\n".join(it["block"] for it in pinned or [] if not it.get("basis"))
    prompt = basis_block(basis_text) + pinned_block(pinned_text) + (
        norms_block(found_text) if found_text else ("" if pinned_text or basis_text else no_norms_note()))
    text = "\n\n".join(x for x in (basis_text, pinned_text, found_text) if x)
    return prompt, text, [it["short"] for it in pinned or []] + list(found_sources or [])


def audit_section(sec, cfg, focus, use_wiki, k, glossary=None, part=None, special="", focus_items=None,
                  doc_block="", location="", max_findings=SECTION_FINDINGS_MAX, norm_queries=None, pinned=None,
                  clauses=None, prompt_max=SECTION_PROMPT_FINDINGS):
    """Проверка одного раздела документа: нормы из базы → черновик замечаний → самокритика.
    special — «особое внимание» пользователя, focus_items — нормы по нему (special_focus_norms); doc_block —
    документ целиком или оглавление (для справки: связи между разделами), location — где раздел в документе;
    pinned — правила экспертов и карточки систем (expert.select_pinned): идут модели сверх лимита норм.
    clauses — (номера пунктов этой части, номера пунктов всего документа): замечание, которое начинается с пункта
    другой части («2.2.3, шаг 4: …» при проверке раздела 1), отбрасывается — его проверяет своя часть.
    Возвращает {report (текст для лога), findings (замечания), failed, sources, dropped, norms_text}."""
    norms_text, sources_used = "", []
    if use_wiki:
        skip = pinned_ids(pinned)
        if norm_queries:
            found, srcs = build_norms_context_multi(norm_queries, cfg, glossary, extra_items=focus_items, skip_ids=skip)
        else:
            found, srcs = build_norms_context(sec, cfg, k, glossary, extra_items=focus_items, skip_ids=skip)
        norms_prompt, norms_text, sources_used = norms_with_pinned(pinned, found, srcs)
    else:
        norms_prompt = autonomous_note()

    user = (doc_block or "") + (location or "") + render("audit_user", section=sec)
    draft = ask_llm(audit_map_prompt(norms_prompt, focus, special, max_findings=prompt_max), user, cfg, guard=True)
    dropped, findings = [], []
    failed = is_llm_error(draft)
    if failed:
        report = draft  # сбой шлюза — НЕ «чистый раздел»
    elif is_clean_report(draft) or len(draft.strip()) < 15:
        report = draft
    else:
        findings = [finding_from_line(line, part) for line in parse_findings(draft)]
        findings = [f for f in findings if len(f["text"]) > 15]
        if clauses:
            own, known = clauses
            keep = []
            for f in findings:
                c = foreign_clause(f["text"], own, known)
                if c:
                    dropped.append(f"{f['text']} — " + tr("о пункте {c}: его проверяет своя часть документа", c=c))
                else:
                    keep.append(f)
            findings = keep
        findings, over = cap_findings(findings, max_findings)
        if over:
            dropped.append(tr("…и ещё {n} замечаний сверх потолка {cap} на один ответ (сначала оставлены критичные)",
                              n=over, cap=max_findings))
        if cfg.use_self_critique and findings:
            findings, rejected = critique_findings(sec, norms_text, findings, cfg, special=special, doc_block=doc_block)
            dropped = rejected + dropped
        report = "\n".join(format_finding(f) for f in findings) if findings else tr("Часть чистая")
    return {"report": report, "findings": findings, "failed": failed, "sources": sources_used, "dropped": dropped,
            "norms_text": norms_text}


def _finding_key(text):
    return frozenset(t for t in bm25_tokenize(text) if t not in BM25_STOP_WORDS)


_RANK = {s: i for i, s in enumerate(_SEVERITIES)}


_STEP_LABEL_MAX = 100       # метка места («2.2.3, шаг 4: …») — не дальше этого от начала замечания
_STEP_RUN_RE = re.compile(r"(?i)\b(?:шаг\w*|step\w*)\s*(\d+(?:\.\d+)?)((?:\s*(?:→|->|—|–|-|,|/|\bи\b|\band\b)\s*"
                          r"(?:шаг\w*\s*|step\w*\s*)?\d+(?:\.\d+)?)*)")


def finding_steps(text):
    """Шаги процесса из метки места в начале замечания («шаг 2 → шаг 3: …», «2.2.3, шаги 6.2–7.1: …») —
    множество номеров; пустое, если метки нет. Номера пунктов документа после «↔» сюда не попадают."""
    head = (text or "")[:_STEP_LABEL_MAX]
    cut = head.find(": ")
    if cut == -1:
        return frozenset()
    steps = set()
    for m in _STEP_RUN_RE.finditer(head[:cut]):
        steps.add(m.group(1))
        steps.update(re.findall(r"\d+(?:\.\d+)?", m.group(2) or ""))
    return frozenset(steps)


def steps_conflict(a, b):
    """Замечания о разных шагах процесса: у обоих указаны шаги, и общих нет (2 → 3 и 4 → 5). Такие не сводятся:
    место в процессе — часть замечания."""
    sa, sb = finding_steps(a), finding_steps(b)
    return bool(sa and sb and not sa & sb)


def _absorb(u, f):
    """Повтор f вливается в оставшееся замечание u: наибольшая критичность, цитата (если своей нет),
    разделы обоих, пункт «особого внимания»."""
    if _RANK.get(f["severity"], 1) < _RANK.get(u["severity"], 1):
        u["severity"] = f["severity"]
    u["quote"] = u.get("quote") or f.get("quote", "")
    sections = list(u.get("sections") or [])
    for s in f.get("sections") or []:
        if s and s not in sections:
            sections.append(s)
    u["sections"] = sections
    origins = list(u.get("origins") or [])
    u["origins"] = origins + [o for o in f.get("origins") or [] if o not in origins]
    if u.get("focus_item") is None and f.get("focus_item") is not None:
        u["focus_item"] = f["focus_item"]
        u["point_kind"] = f.get("point_kind") or "focus"


DUP_SIMILARITY = 0.8        # повтор — совпадение основ слов не меньше этого
RUN_DUP_SIMILARITY = 0.4    # … у замечаний из кусков одного длинного раздела (таблица на десятки строк, разрезанная
                            # на части): модель пишет об одном недостатке в каждом куске своими словами


def _same_run(u, f):
    """Оба замечания — из разных кусков одного длинного раздела (не из одного ответа модели)."""
    ru, rf = u.get("run"), f.get("run")
    return (ru is not None and ru == rf and u.get("part") != f.get("part")
            and not set(u.get("origins") or []) & set(f.get("origins") or []))


def collect_unique_findings(findings, stats=None):
    """Убирает повторы замечаний из разных разделов (совпадение основ слов ≥ DUP_SIMILARITY; из кусков одного
    длинного раздела — ≥ RUN_DUP_SIMILARITY). У оставшегося замечания берётся наибольшая критичность, а цитата — от
    повтора, если своей нет. stats["run_merged"] — сколько склеено по кускам одного раздела."""
    unique, keys = [], []
    for f in findings:
        key = _finding_key(f["text"])
        if not key:
            continue
        for i, k2 in enumerate(keys):
            sim = len(key & k2) / len(key | k2)
            run = sim < DUP_SIMILARITY and _same_run(unique[i], f)
            if (sim >= DUP_SIMILARITY or (run and sim >= RUN_DUP_SIMILARITY)) \
                    and not steps_conflict(unique[i]["text"], f["text"]):
                _absorb(unique[i], f)
                if run and stats is not None:
                    stats["run_merged"] = stats.get("run_merged", 0) + 1
                break
        else:
            unique.append(dict(f))
            keys.append(key)
    return unique


def finding_display(f):
    """Текст замечания для заключения: с разделом документа (или разделами, если замечание повторялось)."""
    sections = [s for s in (f.get("sections") or []) if s]
    if not sections:
        return f["text"]
    if len(sections) == 1:
        return f"{f['text']} — *{sections[0]}*"
    return f"{f['text']} — *" + tr("разделы: {v}", v="; ".join(sections[:6])) + "*"


GENERAL_SECTION = N_("Документ в целом")


def section_order(blocks):
    """Разделы документа в порядке текста: название раздела (как у замечаний) → номер по порядку."""
    order = {}
    for b in blocks:
        if b["level"]:
            order.setdefault(section_label(b["path"]), len(order))
    return order


def corroboration(f):
    """Сколько разных проверок нашли это замечание: ответов модели, чьи повторы слились в него (проверка без
    модели — одна)."""
    return max(1, len(f.get("origins") or []))


def group_by_section(findings, order):
    """Замечания по разделам документа, в порядке документа; внутри раздела — сначала те, что нашли несколько
    проверок независимо (так находятся главные недостатки), затем по критичности. Замечание с несколькими
    разделами — в первом из них по порядку документа, без раздела — в «Документ в целом» (в конце).
    → [(раздел или GENERAL_SECTION, [замечания])]."""
    groups = {}
    for f in findings:
        secs = [x for x in (f.get("sections") or []) if x]
        home = min(secs, key=lambda x: order.get(x, len(order))) if secs else GENERAL_SECTION
        groups.setdefault(home, []).append(f)
    keys = sorted(groups, key=lambda k: (k == GENERAL_SECTION, order.get(k, len(order)), k))
    return [(k, sorted(groups[k], key=lambda f: (-corroboration(f), _RANK.get(f["severity"], 1)))) for k in keys]


def _found_by(f):
    n = corroboration(f)
    if n < 2:
        return ""
    origins = f.get("origins") or []
    passes = {o.split("#p")[1] if "#p" in o else "1" for o in origins}
    if len(passes) < 2:
        return " — *" + tr("нашли проверок: {n}", n=n) + "*"
    checks = len({o.split("#p")[0] for o in origins})
    return " — *" + tr("найдено: {n} (проверок — {c}, проходов — {p})", n=n, c=checks, p=len(passes)) + "*"


def _also_sections(f, home):
    """« — также: …» для замечания, которое относится к нескольким разделам."""
    other = [x for x in (f.get("sections") or []) if x and x != home]
    return (" — *" + tr("также: {v}", v="; ".join(other[:5])) + "*") if other else ""


_SEV_PREFIX = {"КРИТИЧНЫЙ": N_("🔥 **[КРИТИЧНЫЙ]**"), "СРЕДНИЙ": N_("⚠️ **[СРЕДНИЙ]**"), "НИЗКИЙ": N_("📝 **[НИЗКИЙ]**")}


def _focus_block(focus_answers, findings):
    """Блок заключения «Ответы на пункты особого внимания»: ответ модели и замечания по каждому пункту."""
    if not focus_answers:
        return ""
    lines = [tr("## 🎯 ОТВЕТЫ НА ПУНКТЫ ОСОБОГО ВНИМАНИЯ"), ""]
    for i, fa in enumerate(focus_answers):
        lines.append(f"**{i + 1}. {fa['item']}**")
        if fa["failed"]:
            lines += [tr("❌ Ответ не получен: {err}", err=fa["report"].strip()), ""]
            continue
        mine = sorted((f for f in findings if f.get("focus_item") == i), key=lambda f: _RANK.get(f["severity"], 1))
        if fa["answer"] or not mine:
            lines.append(fa["answer"] or tr("Замечаний по пункту нет."))
        for n, f in enumerate(mine, 1):
            lines.append(f"{n}. {tr(_SEV_PREFIX.get(f['severity'], _SEV_PREFIX['СРЕДНИЙ']))} {finding_display(f)}")
        lines.append("")
    return "\n".join(lines) + "\n"


def _numbered_findings(findings):
    mine = sorted(findings, key=lambda f: _RANK.get(f["severity"], 1))
    return [f"{n}. {tr(_SEV_PREFIX.get(f['severity'], _SEV_PREFIX['СРЕДНИЙ']))} {finding_display(f)}"
            for n, f in enumerate(mine, 1)]


def _checklist_block(entries, findings, doc_type=""):
    """Блок «Чек-лист»: ответ на каждый пункт (да / частично / нет / не относится) и замечания по нему."""
    if not entries:
        return ""
    counts = Counter(e.get("verdict", "unknown") for e in entries if not e["failed"])
    n_failed = sum(1 for e in entries if e["failed"])
    title = tr("## ✅ ЧЕК-ЛИСТ «{t}»", t=doc_type) if doc_type else tr("## ✅ ЧЕК-ЛИСТ")
    summary = tr("Да — {y}, частично — {p}, нет — {n}, не относится — {na}", y=counts["yes"], p=counts["partial"],
                 n=counts["no"], na=counts["na"])
    if counts["unknown"]:
        summary += tr(", без ясного ответа — {u}", u=counts["unknown"])
    if n_failed:
        summary += tr(", ответ не получен — {f}", f=n_failed)
    lines = [title, summary + ".", ""]
    for i, e in enumerate(entries):
        if e["failed"]:
            lines += [f"❔ **{i + 1}. {e['item']}**", tr("❌ Ответ не получен: {err}", err=e["report"].strip()), ""]
            continue
        lines.append(f"{VERDICT_ICONS.get(e.get('verdict', 'unknown'), '❔')} **{i + 1}. {e['item']}**  ")
        lines.append(e["answer"] or tr("Ответ без пояснения."))
        lines += _numbered_findings([f for f in findings if f.get("focus_item") == i]) + [""]
    return "\n".join(lines) + "\n"


def _systems_block(entries, findings):
    """Блок «Что документ требует от систем»: по карточке каждой упомянутой системы — требования со статусами
    и замечания."""
    if not entries:
        return ""
    lines = [tr("## 🧩 ЧТО ДОКУМЕНТ ТРЕБУЕТ ОТ СИСТЕМ"),
             tr("По карточкам систем из базы знаний: что система умеет сейчас, что в планах, где нужна доработка."), ""]
    for i, e in enumerate(entries):
        lines.append(f"### {e['item']}")
        lines.append(tr("_Карточка: {card}; упоминаний в документе: {n}._", card=e.get("card", ""), n=e.get("mentions", 0))
                     + "  ")
        if e["failed"]:
            lines += [tr("❌ Ответ не получен: {err}", err=e["report"].strip()), ""]
            continue
        counts = Counter(c for c, _t in e["statuses"])
        if counts:
            lines.append(" · ".join(f"{tr(STATUS_LABELS[c])}: {counts[c]}" for c in STATUS_ORDER if counts[c]))
            lines.append("")
            for c in STATUS_ORDER:
                lines += [f"- {tr(STATUS_LABELS[c])} — {t}" for code, t in e["statuses"] if code == c]
        elif e["answer"]:
            lines.append(e["answer"])
        mine = _numbered_findings([f for f in findings if f.get("focus_item") == i])
        if mine:
            lines += ["", tr("**Замечания:**")] + mine
        lines.append("")
    return "\n".join(lines) + "\n"


def _unverified_block(unverified):
    """Блок «Не подтверждено источником»: замечания со ссылками и цифрами, которых нет ни в документе, ни в нормах."""
    if not unverified:
        return ""
    lines = [tr("## ❓ НЕ ПОДТВЕРЖДЕНО ИСТОЧНИКОМ — ПРОВЕРЬТЕ"),
             tr("Ссылки и цифры из этих замечаний не найдены ни в документе, ни в нормативных фрагментах, которые "
                "получила модель: возможно, модель их придумала."), ""]
    for f in unverified:
        missing = ", ".join(f"«{m}»" for m in f.get("unverified") or [])
        lines.append(f"- {tr(_SEV_PREFIX.get(f['severity'], _SEV_PREFIX['СРЕДНИЙ']))} {finding_display(f)} "
                     + tr("(не найдено: {v})", v=missing))
    return "\n".join(lines) + "\n"


_TYPO_RE = re.compile(r"«[^»\n]{1,80}»\s*→\s*«[^»\n]{1,80}»|\bопечатк|орфограф|\btypo|spelling", re.IGNORECASE)
_PLACE_RE = re.compile(r"\b\d+(?:\.\d+)+\b|\bшаг\w*\s*\d|\bstep\s*\d|\bтабл|\btable\s*\d|\bп\.\s*\d", re.IGNORECASE)


def split_minor_findings(findings):
    """Главное отдельно от мелочей: (главные, мелкие и общие, опечатки). Мелкое — низкая критичность без цитаты и
    без указания места в документе («Указать, каким образом обеспечивается …»); опечатки — отдельно. Проверки
    без модели остаются в главном."""
    main, minor, typos = [], [], []
    for f in findings:
        if is_typo(f):
            typos.append(f)
        elif (f["severity"] == "НИЗКИЙ" and f.get("source") != "check" and not f.get("quote")
              and not _PLACE_RE.search(f["text"])):
            minor.append(f)
        else:
            main.append(f)
    return main, minor, typos


def _plain_block(title, findings):
    if not findings:
        return ""
    lines = [title.format(n=len(findings)), ""] + [f"- {finding_display(f)}" for f in findings]
    return "\n".join(lines) + "\n"


# ---------- «Главное» в начале заключения ----------
# Самые важные замечания одной строкой каждое — указателями на номера замечаний ниже, а не их копиями: число замечаний
# в заключении не меняется. Выбирает модель-редактор (audit_summary_system); не ответила — по подтверждению и
# критичности.
SUMMARY_MAX_CHARS = 30000       # список замечаний в запросе; длиннее — сначала найденные несколькими проверками
SUMMARY_GIST_CHARS = 160        # строка «Главного», не длиннее


def _by_severity(findings):
    return sorted(findings, key=lambda f: _RANK.get(f["severity"], 1))


def summary_candidates(by_section, focus=(), focus_findings=(), checklist=(), checklist_findings=(), systems=(),
                       system_findings=()):
    """Замечания, из которых выбирается «Главное», с их номерами в заключении (нумерация как в
    build_final_report): главные по разделам — «№12»; по пунктам — «особое внимание 2, №1» и т. п.
    → [{"ref", "section", "finding"}]."""
    out, n = [], 1
    for home, items in by_section or []:
        for f in items:
            out.append({"ref": tr("№{n}", n=n), "section": home, "finding": f})
            n += 1
    for entries, findings, label in ((focus, focus_findings, N_("особое внимание {i}, №{n}")),
                                     (checklist, checklist_findings, N_("чек-лист {i}, №{n}"))):
        for i, e in enumerate(entries or []):
            if e["failed"]:
                continue
            mine = _by_severity(f for f in findings or [] if f.get("focus_item") == i)
            out += [{"ref": tr(label, i=i + 1, n=k), "section": "", "finding": f} for k, f in enumerate(mine, 1)]
    for i, e in enumerate(systems or []):
        if e["failed"]:
            continue
        mine = _by_severity(f for f in system_findings or [] if f.get("focus_item") == i)
        out += [{"ref": tr("система «{s}», №{n}", s=e["item"], n=k), "section": "", "finding": f}
                for k, f in enumerate(mine, 1)]
    return out


def _summary_priority(cands):
    """Порядок важности без модели: найденные несколькими проверками, затем по критичности, затем по порядку."""
    return sorted(range(len(cands)), key=lambda i: (-corroboration(cands[i]["finding"]),
                                                     _RANK.get(cands[i]["finding"]["severity"], 1), i))


def _summary_fallback(cands, limit):
    picked = [i for i in _summary_priority(cands) if cands[i]["finding"]["severity"] != "НИЗКИЙ"][:limit]
    return [{**cands[i], "gist": _short(cands[i]["finding"]["text"], SUMMARY_GIST_CHARS)} for i in picked]


def select_summary(cands, cfg, special="", limit=12):
    """«Главное»: до limit замечаний из cands (summary_candidates) с сутью одной строкой. → (пункты [{"ref",
    "section", "finding", "gist"}], ошибка — пустая строка, если выбрала модель)."""
    if limit <= 0 or not cands:
        return [], ""
    keep, size = [], 0
    for i in _summary_priority(cands):
        length = len(cands[i]["finding"]["text"]) + 60      # с номером, критичностью и разделом
        if size + length > SUMMARY_MAX_CHARS:
            continue
        keep.append(i)
        size += length
    keep.sort()
    listing = []
    for k, i in enumerate(keep, 1):
        c = cands[i]
        place = c["section"] or "; ".join(x for x in c["finding"].get("sections") or [] if x)
        n = corroboration(c["finding"])
        listing.append(f"{k}. [{severity_label(c['finding']['severity'])}] {c['finding']['text']}"
                       + (f" — {tr(place)}" if place else "") + (" — " + tr("нашли проверок: {n}", n=n) if n > 1 else ""))
    special_block = render("audit_summary_special", special=special) + "\n\n" if (special or "").strip() else ""
    raw = ask_llm(render("audit_summary_system", max=limit),
                  render("audit_summary_user", special=special_block, findings="\n".join(listing)), cfg, guard=True)
    if is_llm_error(raw):
        return _summary_fallback(cands, limit), raw.strip()
    data = extract_json(raw, kind="array")
    if not isinstance(data, list):
        return _summary_fallback(cands, limit), tr("ответ модели не разобран")
    out, seen = [], set()
    for d in data:
        if not isinstance(d, dict):
            continue
        try:
            k = int(d.get("n"))
        except (TypeError, ValueError):
            continue
        if not 1 <= k <= len(keep) or k in seen:
            continue
        seen.add(k)
        c = cands[keep[k - 1]]
        gist = " ".join(str(d.get("gist") or "").split()) or c["finding"]["text"]
        out.append({**c, "gist": _short(gist, SUMMARY_GIST_CHARS)})
        if len(out) >= limit:
            break
    if data and not out:
        return _summary_fallback(cands, limit), tr("в ответе модели нет номеров из списка")
    return out, ""


_SEV_ICON = {"КРИТИЧНЫЙ": "🔥", "СРЕДНИЙ": "⚠️", "НИЗКИЙ": "📝"}


def _summary_block(summary, by_model=True):
    """Блок «Главное»: суть замечания одной строкой → номер замечания в заключении."""
    if not summary:
        return ""
    lines = [tr("## ⭐ ГЛАВНОЕ"),
             (tr("_Самые важные замечания, по мнению модели-редактора; подробности и цитаты — в замечаниях по "
                 "номерам ниже._") if by_model else
              tr("_Самые важные замечания — найденные несколькими проверками и критичные; подробности — в "
                 "замечаниях по номерам ниже._")), ""]
    for k, it in enumerate(summary, 1):
        place = tr(it["section"]) if it.get("section") and it["section"] != GENERAL_SECTION else ""
        lines.append(f"{k}. {_SEV_ICON.get(it['finding']['severity'], '⚠️')} {it['gist']} → **{it['ref']}**"
                     + (f" · *{place}*" if place else ""))
    return "\n".join(lines) + "\n\n"


def build_final_report(filename, sections, failed_sections, n_sections, special="", focus_answers=None,
                       focus_findings=None, unverified=None, deep_note="", minor=None, typos=None, checklist=None,
                       checklist_findings=None, doc_type="", systems=None, system_findings=None, summary=None,
                       summary_by_model=True):
    """Заключение (Markdown). sections — главные замечания по разделам документа: [(раздел, [замечания])] из
    group_by_section; summary — «Главное» (select_summary): указатели на замечания по номерам."""
    formatted_blocks, has_any, counter = [], False, 1
    for home, items in sections or []:
        if not items:
            continue
        has_any = True
        formatted_blocks.append(f"### {tr(home)}")
        for f in items:
            formatted_blocks.append(f"{counter}. {tr(_SEV_PREFIX.get(f['severity'], _SEV_PREFIX['СРЕДНИЙ']))} "
                                    f"{f['text']}{_also_sections(f, home)}{_found_by(f)}")
            counter += 1
        formatted_blocks.append("")

    focus_findings = list(focus_findings or [])
    checklist_findings, system_findings = list(checklist_findings or []), list(system_findings or [])
    point_findings = focus_findings + checklist_findings + system_findings
    minor, typos = list(minor or []), list(typos or [])
    if has_any or point_findings or minor or typos:
        risks_block = ("\n".join(formatted_blocks).strip()
                       or (tr("Замечания — в ответах на пункты выше.") if point_findings
                           else tr("Главных замечаний нет — только мелкие и опечатки (ниже).")))
        has_critical = (any(f["severity"] == "КРИТИЧНЫЙ" for _h, items in sections or [] for f in items)
                        or any(f["severity"] == "КРИТИЧНЫЙ" for f in point_findings))
        verdict_text = (tr("⚠️ ДОКУМЕНТ ТРЕБУЕТ ДОРАБОТКИ. ")
                        + (tr("Выявлены критические замечания.") if has_critical else tr("Выявлены замечания.")))
    elif failed_sections:
        risks_block = tr("### ❓ В проверенных разделах замечаний не обнаружено.")
        verdict_text = ""
    else:
        risks_block = tr("### ✅ Замечаний не обнаружено.")
        verdict_text = tr("✅ ДОКУМЕНТ СОГЛАСОВАН.")

    failed_block = ""
    if failed_sections:
        verdict_text = (verdict_text + " " if verdict_text else "") + (
            tr("⚠️ ПРОВЕРЕНО НЕ ПОЛНОСТЬЮ: не удалось проверить {n_failed_sections} из {n_sections} разделов — "
               "результат нельзя считать согласованием.", n_failed_sections=len(failed_sections), n_sections=n_sections)
        )
        failed_block = (
            tr("\n## ⚠️ РАЗДЕЛЫ, КОТОРЫЕ НЕ УДАЛОСЬ ПРОВЕРИТЬ\n")
            + "\n".join(tr("- Часть {n}/{n_sections} «{prev}»: {err}", n=n, n_sections=n_sections, prev=prev, err=err) for n, prev, err in failed_sections)
            + "\n"
        )

    header = (
        tr("# ЭКСПЕРТНОЕ ЗАКЛЮЧЕНИЕ\n**Объект анализа**: документ «{filename}»\n", filename=filename)
        + (tr("**Тип документа**: {v}\n", v=doc_type) if doc_type else "")
        + (tr("**Особое внимание**: {v}\n", v=special.strip()) if (special or "").strip() else "")
        + tr("**Статус**: {verdict_text}\n\n", verdict_text=verdict_text)
    )
    header = re.sub(r"(?<=[^\n])\n(?=\*\*)", "  \n", header)     # строки шапки — каждая с новой строки в Markdown
    return (
        header
        + _summary_block(summary, summary_by_model)
        + _focus_block(focus_answers, focus_findings)
        + _checklist_block(checklist, checklist_findings, doc_type)
        + _systems_block(systems, system_findings)
        + tr("## 📋 ЗАМЕЧАНИЯ ПО РАЗДЕЛАМ ДОКУМЕНТА\n{risks_block}\n{failed_block}",
             risks_block=risks_block, failed_block=failed_block)
        + ("\n" + _unverified_block(unverified) if unverified else "")
        + ("\n" + _plain_block(tr("## 📝 МЕЛКИЕ И ОБЩИЕ ЗАМЕЧАНИЯ ({n}) — без цитаты и указания места"), minor)
           if minor else "")
        + ("\n" + _plain_block(tr("## ✏️ ОПЕЧАТКИ И ГРАММАТИКА ({n})"), typos) if typos else "")
        + ("\n" + tr("_Глубокая проверка: {v}._", v=deep_note) + "\n" if deep_note else "")
    )


# ---------- Проверки без модели: сокращения ----------
# Сокращение — слово из заглавных букв (ХРО, ОФ) или со строчными связками (ОТиП). Расшифрованным считается
# сокращение из раздела «Аббревиатуры и сокращения» / «Термины и определения» (первая колонка таблицы или текст до
# тире) или со скобками в тексте: «Отдел технической приёмки … (ОТиП)».
# Не сокращение — часть имени сервера, учётной записи, пути или кода документа: «CORP\\Admins», «srv.CORP.local»,
# «ОРГ.10.ТЗ.2», «index=OSLOG», «user@HOST»
_ABBR_RE = re.compile(r"(?<![\w/\\.@=-])([А-ЯЁA-Z]{2,}(?:[а-яё]{1,2}[А-ЯЁ]+)*(?:/[А-ЯЁA-Z]{2,})?)"
                      r"(?![\w/\\@=-]|\.\w)")
_DEF_HEADING_RE = re.compile(r"(?i)сокращени|аббревиатур|обозначени|термин|определени|глоссари|"
                             r"abbreviation|acronym|glossary|\bterms\b|definitions")
_ROMAN_RE = re.compile(r"^[IVXLCM]+$")
KNOWN_ABBREVIATIONS = frozenset(
    "РФ ГОСТ ISO IEC МЭК URL URI HTTP HTTPS API REST SOAP XML JSON PDF ZIP CSV XLSX DOCX ID IT ИТ OK ФИО ИНН ОГРН "
    "КПП БИК США ЕС ООН USB IP TCP UDP DNS SQL UTC GMT МСК ПК "
    # общеизвестные в ИТ: протоколы, форматы, оборудование
    "SSH NTP TLS SSL LDAP LDAPS SMTP IMAP FTP SFTP SNMP ICMP DHCP NAT VPN VLAN LAN WAN RDP SMB NFS SSO MFA OTP PKI CRL "
    "OCSP JWT SAML WSDL XSD HTML CSS YAML UTF ASCII PNG JPG JPEG GIF SVG TXT SDK GUI CLI UI UX CPU RAM ROM SSD HDD OS "
    "VM DB DBMS ПО ОС БД СУБД ВМ ЛВС "
    # слова и значения заглавными буквами, а не сокращения
    "ДА НЕТ НЕ ИЛИ ВСЕ ВСЁ ТЕСТ YES NO ON OFF TRUE FALSE NULL NONE ALL ANY".split())
_PRODUCT_WORD_RE = re.compile(r"(?<![\w.])[A-Z][a-z]{2,}[ \t]+$")
_CAPS_WORD_MIN = 3          # «ТЕСТ», если в тексте есть «тест» или «Тест», — слово заглавными, а не сокращение
ABBR_LIST_MAX = 15


def _abbr_tokens(text, stop):
    out = set()
    for tok in _ABBR_RE.findall(text or ""):
        if tok in stop or _ROMAN_RE.match(tok) or (len(tok) >= 8 and tok.isupper()):
            continue                       # УТВЕРЖДАЮ, СОГЛАСОВАНО — не сокращения
        if "/" in tok and all(x in stop for x in tok.split("/")):
            continue                       # «TCP/UDP» — оба общеизвестные
        out.add(tok)
    return out


def _definition_terms(text):
    """Термины из раздела сокращений: первая содержательная ячейка строки таблицы или текст до тире."""
    terms = []
    for line in (text or "").split("\n"):
        s = line.strip()
        if s.startswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
        else:
            cells = [re.split(r"\s[–—-]\s|:\s", s, maxsplit=1)[0]]
        cells = [c for c in cells if c and not re.fullmatch(r"[\d.№п/\s]+|:?-{2,}:?", c)]
        if cells and len(cells[0]) <= 60:
            terms.append(cells[0])
    return terms


def check_abbreviations(blocks, known=""):
    """Сокращения, которые используются без расшифровки, и расшифрованные, но неиспользуемые — замечания
    без модели. blocks — document_blocks(); known — свои общеизвестные сокращения через запятую."""
    stop = KNOWN_ABBREVIATIONS | {x.strip() for x in re.split(r"[,;\s]+", known or "") if x.strip()}
    def_heading, def_terms, defined, used, first = "", [], set(), Counter(), {}
    for b in blocks:
        last = (b["path"] or "").split(" → ")[-1]
        if not b["level"] and last and _DEF_HEADING_RE.search(last):
            def_heading = def_heading or last
            for term in _definition_terms(b["text"]):
                toks = _abbr_tokens(term, stop)
                if toks:
                    defined |= toks
                    def_terms.append((term, toks))
            continue
        for m in re.finditer(r"\(([^()]{2,20})\)", b["text"]):          # «… (ОТиП)» — расшифровка в тексте
            defined |= _abbr_tokens(m.group(1), stop)
        for m in _ABBR_RE.finditer(b["text"]):
            tok = m.group(1)
            if tok.isascii() and _PRODUCT_WORD_RE.search(b["text"], max(0, m.start() - 40), m.start()):
                continue                   # «Astra Linux SE», «Kaspersky Endpoint EDR» — часть названия продукта
            if tok in _abbr_tokens(tok, stop):
                used[tok] += 1
                first.setdefault(tok, b["path"])
    # слово, написанное заглавными («ВАЖНО», «ТЕСТ» в ячейках таблиц), если то же слово есть в тексте обычным
    # написанием («Важно:», «тест»)
    lower = {w.lower() for w in re.findall(r"(?<!\w)[А-ЯЁа-яё][а-яё]{%d,}(?!\w)" % (_CAPS_WORD_MIN - 1),
                                           "\n".join(b["text"] for b in blocks))}
    undefined = [(t, n) for t, n in used.most_common()
                 if t not in defined and not (len(t) >= _CAPS_WORD_MIN and t.lower() in lower)]
    out = []

    def listed(items):
        parts = []
        for tok, n in items[:ABBR_LIST_MAX]:
            where = section_label(first.get(tok)) or tr("начало документа")
            parts.append(tr("{tok} (×{n}, впервые — «{where}»)", tok=tok, n=n, where=where))
        more = len(items) - ABBR_LIST_MAX
        return "; ".join(parts) + (tr(" и ещё {n}", n=more) if more > 0 else "")

    if def_heading:
        if undefined:
            out.append({"severity": "СРЕДНИЙ", "quote": "", "part": None,
                        "sections": [section_label(def_heading)], "source": "check",
                        "text": tr("Сокращения используются без расшифровки — их нет в разделе «{h}»: {v}. Добавьте их "
                                   "в перечень.", h=section_label(def_heading), v=listed(undefined))})
        unused = [term for term, toks in def_terms if not any(t in used for t in toks)]
        if unused:
            out.append({"severity": "НИЗКИЙ", "quote": "", "part": None,
                        "sections": [section_label(def_heading)], "source": "check",
                        "text": tr("В разделе «{h}» расшифрованы сокращения, которых нет в тексте документа: {v}.",
                                   h=section_label(def_heading), v=", ".join(unused[:ABBR_LIST_MAX]))})
    elif len(undefined) >= 3:
        out.append({"severity": "НИЗКИЙ", "quote": "", "part": None, "sections": [],
                    "source": "check",
                    "text": tr("В документе нет перечня сокращений; без расшифровки используются: {v}.",
                               v=listed(undefined))})
    return out


# ---------- Проверки без модели: ссылки и цифры в замечаниях ----------
# Ссылка на таблицу, пункт или акт («Таблица №5», «п. 3.5.2», «99-ФЗ») и цифры с единицами («4 Гб», «60 мин») из
# замечания ищутся в документе и в нормативных фрагментах, которые получила модель. Не нашлось — замечание уходит
# в блок «Не подтверждено источником». Вычисленные моделью значения («≈ 1 350 Гб», «в 27 раз») не проверяются.
_REF_TABLE_RE = re.compile(r"(?i)\b(?:табл(?:ица|ице|ицы|ицу|ицей|\.)?|table)\s*(?:№\s*|no\.?\s*)?(\d+(?:\.\d+)*)")
_REF_CLAUSE_RE = re.compile(r"(?i)(?:\bпп?\.|\bпункт\w*|\bподпункт\w*|\bраздел\w*|\bподраздел\w*|\bглав\w*|"
                            r"\bsection|\bclause)\s*(\d+(?:\.\d+)+)")
_REF_ACT_RE = re.compile(r"(?i)(?<![\w-])(\d{1,4})\s*-\s*(п|у|фз)(?![а-яё\w])")
_UNITS = {"мб": "(?:мб|мбайт|мегабайт\\w*|mb)", "mb": "(?:мб|мбайт|мегабайт\\w*|mb)",
          "гб": "(?:гб|гбайт|гигабайт\\w*|gb)", "gb": "(?:гб|гбайт|гигабайт\\w*|gb)",
          "кб": "(?:кб|кбайт|килобайт\\w*|kb)", "kb": "(?:кб|кбайт|килобайт\\w*|kb)",
          "тб": "(?:тб|тбайт|терабайт\\w*|tb)", "tb": "(?:тб|тбайт|терабайт\\w*|tb)",
          "мин": "(?:мин|минут\\w*)", "час": "(?:ч\\.?|час\\w*)", "сек": "(?:сек\\w*|с\\b)", "%": "%"}
_NUM_UNIT_RE = re.compile(r"(?i)(?<![\d.,])(\d+(?:[.,]\d+)?)\s?(мб|мбайт|гб|гбайт|кб|тб|mb|gb|kb|tb|мин(?:ут\w*)?|"
                          r"час(?:ов|а)?|сек(?:унд\w*)?|%)")
_COMPUTED_RE = re.compile(r"(?i)(?:≈|~|×|\*|=|около|примерно|порядка|в\s+\d+[.,]?\d*\s+раз)\s*$")


def _norm_src(text):
    return re.sub(r"\s+", " ", (text or "").lower().replace("ё", "е").replace(" ", " "))


def unverified_refs(text, sources_norm, check_clauses=True):
    """Ссылки и цифры из замечания, которых нет в источниках (sources_norm — _norm_src(документ + нормы))."""
    missing = []
    for m in _REF_TABLE_RE.finditer(text or ""):
        n = re.escape(m.group(1))
        if not re.search(rf"(?:табл\w*\.?|table)\s*(?:№\s*|no\.?\s*)?{n}(?!\d|\.\d)", sources_norm):
            missing.append(m.group(0).strip())
    if check_clauses:
        for m in _REF_CLAUSE_RE.finditer(text or ""):
            n = re.escape(m.group(1))
            if not re.search(rf"(?<![\d.]){n}(?!\d|\.\d)", sources_norm):
                missing.append(m.group(0).strip())
    for m in _REF_ACT_RE.finditer(text or ""):
        if not re.search(rf"(?<!\d){m.group(1)}\s*-\s*{m.group(2).lower()}(?![а-яё\w])", sources_norm):
            missing.append(m.group(0).strip())
    nums, absent = [], []
    for m in _NUM_UNIT_RE.finditer(text or ""):
        if _COMPUTED_RE.search((text or "")[max(0, m.start() - 20):m.start()]):
            continue
        num = re.escape(m.group(1)).replace(",", "[.,]").replace("\\.", "[.,]")
        unit = m.group(2).lower()
        key = next((k for k in _UNITS if unit.startswith(k)), unit)
        nums.append(m.group(0))
        if not re.search(rf"(?<![\d.,]){num}\s?{_UNITS.get(key, re.escape(unit))}", sources_norm):
            absent.append(m.group(0))
    if nums and len(absent) == len(nums):          # все цифры замечания нигде не встречаются
        missing += absent
    return list(dict.fromkeys(missing))


# ---------- Сведение повторов моделью ----------
MERGE_MAX_CHARS = 30000     # список замечаний в одном запросе сведения; длиннее — пачками
_PAIR_RE = re.compile(r"[«\"“]([^»\"”\n]{1,1000})[»\"”]\s*→\s*[«\"“]([^»\"”\n]{1,1000})[»\"”]")
_NO_ERROR_RE = re.compile(r"(?i)нет\s+ошибк|ошибок\s+(?:не\s+найдено|нет)|не\s+(?:является\s+)?ошибк|"
                          r"\bno\s+(?:error|mistake|typo)|not\s+an?\s+(?:error|mistake|typo)")


_NO_PROBLEM_RE = re.compile(
    r"(?i)(?:проблем\w*|противоречи\w*|несоответстви\w*|нарушени\w*|замечани\w*)(?:\s+[\w-]+){0,2}?\s+"
    r"(?:нет|не\s+(?:выявлен|обнаружен|найден|установлен)\w*)(?:\s+[\w-]+){0,4}\s*[.!]?\s*$|"
    r"\bнет\s+(?:проблем|противоречи|несоответстви)\w*(?:\s+[\w-]+){0,4}\s*[.!]?\s*$|"
    r"\bno\s+(?:[\w-]+\s+)?(?:problems?|contradictions?|inconsistenc\w+)(?:\s+[\w-]+){0,4}\s*[.!]?\s*$")
_BUT_RE = re.compile(r"(?i)\bно\b|однако|\bbut\b|however|\?")


def is_no_problem(text):
    """Замечание, которое само говорит, что проблемы нет («…, что соответствует. Проблем нет по этому пункту.»,
    «…поэтому противоречий между ними нет»): последняя фраза — такой вывод, без «но» и «однако»."""
    sentences = [x for x in re.split(r"(?<=[.!?])\s+", (text or "").strip()) if x.strip()]
    last = sentences[-1] if sentences else ""
    return bool(_NO_PROBLEM_RE.search(last)) and not _BUT_RE.search(last)


def _short(text, n=110):
    text = re.sub(r"\s+", " ", text or "").strip()
    return text if len(text) <= n else text[:n].rstrip() + "…"


def merge_log_line(members, apart=(), other_steps=()):
    """Строка лога о группе сведения: что осталось, что влилось, что оставлено отдельно (тот же ответ модели или
    другой шаг процесса)."""
    line = tr("Оставлено: «{t}»", t=_short(members[0]["text"]))
    if len(members) > 1:
        line += " ← " + tr("влито: {v}", v="; ".join(f"«{_short(f['text'])}»" for f in members[1:]))
    if apart:
        line += " · " + tr("отдельно (тот же ответ модели): {v}", v="; ".join(f"«{_short(f['text'])}»" for f in apart))
    if other_steps:
        line += " · " + tr("отдельно (другой шаг процесса): {v}",
                           v="; ".join(f"«{_short(f['text'])}»" for f in other_steps))
    return line


def is_typo(f):
    """Опечатка или грамматика: проверка языка или низкая критичность и «как написано» → «как правильно»."""
    return f.get("check") == "language" or (f["severity"] == "НИЗКИЙ" and bool(_TYPO_RE.search(f["text"])))


def _typo_norm(text):
    return re.sub(r"[\s\u00a0]+", " ", (text or "").replace("ё", "е").lower()).strip(" .,;")


def clean_typos(findings):
    """Опечатки без мусора и повторов: строка, где «как написано» и «как правильно» совпадают и сама модель
    пишет «ошибки нет», убирается; одна и та же исправляемая фраза из разных проверок — одна строка."""
    out, seen = [], {}
    for f in findings:
        pairs = _PAIR_RE.findall(f["text"])
        if (len(pairs) == 1 and _typo_norm(pairs[0][0]) == _typo_norm(pairs[0][1])
                and _NO_ERROR_RE.search(f["text"])):
            continue
        key = _typo_norm(pairs[0][0]) if pairs else None
        if key and key in seen:
            _absorb(seen[key], f)
            continue
        f = dict(f)
        out.append(f)
        if key:
            seen[key] = f
    return out


def merge_duplicate_findings(findings, cfg, log=None):
    """Последний шаг аудита: модель группирует замечания об одном и том же недостатке (из разных разделов и
    ответов на «особое внимание»), группа сводится в одно замечание. Длинный список — пачками (замечания одного
    раздела — в одной пачке), затем, если поместится, ещё один проход по всем. Возвращает (замечания, групп,
    ошибка): при сбое модели или неразборчивом ответе замечания остаются как были.

    Опечатки модели не передаются: у каждой своя фраза, повторы среди них убирает clean_typos. Замечания из
    одного ответа модели (одна часть документа, один вопрос глубокой проверки, один пункт «особого внимания») не
    сливаются друг с другом: модель сама перечислила их как разные проблемы, а при слиянии осталась бы одна
    формулировка. log — список, куда дописывается состав каждой группы (для лога проверки)."""
    typos = [f for f in findings if is_typo(f)]
    if not typos:
        return _merge_rest(findings, cfg, log)
    merged, groups, err = _merge_rest([f for f in findings if not is_typo(f)], cfg, log)
    if not groups:
        return findings, 0, err
    return merged + typos, groups, err


def _merge_rest(findings, cfg, log=None):
    if len(_merge_listing(findings)) <= MERGE_MAX_CHARS:
        return _merge_once(findings, cfg, log)
    order = sorted(range(len(findings)), key=lambda i: ((findings[i].get("sections") or [""])[0], i))
    chunks, cur, cur_len = [], [], 0
    for i in order:
        n = len(_merge_listing([findings[i]])) + 1
        if cur and cur_len + n > MERGE_MAX_CHARS * 2 // 3:
            chunks.append(cur)
            cur, cur_len = [], 0
        cur.append(findings[i])
        cur_len += n
    if cur:
        chunks.append(cur)
    merged, groups, errors = [], 0, []
    for chunk in chunks:
        out, g, err = _merge_once(chunk, cfg, log)
        merged += out
        groups += g
        if err:
            errors.append(err)
    if len(_merge_listing(merged)) <= MERGE_MAX_CHARS:
        merged, g, err = _merge_once(merged, cfg, log)
        groups += g
        if err:
            errors.append(err)
    return merged, groups, "; ".join(dict.fromkeys(errors))


def _merge_listing(findings):
    lines = []
    for i, f in enumerate(findings, 1):
        where = "; ".join(s for s in (f.get("sections") or []) if s)
        lines.append(f"{i}. [{severity_label(f['severity'])}] " + (f"({where}) " if where else "") + f["text"])
    return "\n".join(lines)


def _merge_once(findings, cfg, log=None):
    """Один запрос сведения повторов (список помещается)."""
    if len(findings) < 3:
        return findings, 0, ""
    lines = []
    for i, f in enumerate(findings, 1):
        where = "; ".join(s for s in (f.get("sections") or []) if s)
        lines.append(f"{i}. [{severity_label(f['severity'])}] " + (f"({where}) " if where else "") + f["text"])
    listing = "\n".join(lines)
    raw = ask_llm(prompt_text("audit_merge_system"), render("audit_merge_user", findings=listing), cfg, guard=True)
    if is_llm_error(raw):
        return findings, 0, raw.strip()
    data = extract_json(raw, kind="array")
    if not isinstance(data, list):
        return findings, 0, tr("ответ модели не разобран")
    taken, groups = set(), []
    for g in data:
        if isinstance(g, dict):
            g = g.get("items")
        if not isinstance(g, list):
            continue
        idx = []
        for x in g:
            try:
                n = int(x)
            except (TypeError, ValueError):
                continue
            if 1 <= n <= len(findings) and n not in taken and n not in idx:
                idx.append(n)
        if len(idx) >= 2:
            taken.update(idx)
            groups.append(idx)
    if not groups:
        return findings, 0, ""
    rep_of, n_groups = {}, 0
    for idx in groups:
        # остаётся подтверждённая формулировка (без выдуманных ссылок), затем самая критичная, затем ответ на
        # «особое внимание», затем самая полная; критичность — наибольшая в группе, разделы — всех повторов.
        # Замечание из того же ответа модели, что и уже вошедшие в группу, остаётся отдельным.
        idx = sorted(idx, key=lambda n: (bool(findings[n - 1].get("unverified")),
                                         _RANK.get(findings[n - 1]["severity"], 1),
                                         findings[n - 1].get("focus_item") is None, -len(findings[n - 1]["text"])))
        rep = dict(findings[idx[0] - 1])
        used = set(rep.get("origins") or [])
        members = [idx[0]]
        apart, other_steps = [], []
        for n in idx[1:]:
            own = set(findings[n - 1].get("origins") or [])
            if own & used:
                apart.append(n)
                continue
            if steps_conflict(rep["text"], findings[n - 1]["text"]):   # модель свела разные шаги процесса
                other_steps.append(n)
                continue
            _absorb(rep, findings[n - 1])
            used |= own
            members.append(n)
        if log is not None and len(idx) >= 2:
            log.append(merge_log_line([findings[n - 1] for n in members], [findings[n - 1] for n in apart],
                                      [findings[n - 1] for n in other_steps]))
        if len(members) < 2:
            continue
        n_groups += 1
        for n in members:
            rep_of[n] = rep
    out, emitted = [], set()
    for i, f in enumerate(findings, 1):
        rep = rep_of.get(i)
        if rep is None:
            out.append(f)
        elif id(rep) not in emitted:
            emitted.add(id(rep))
            out.append(rep)
    return out, n_groups, ""


# ---------- Документ целиком в проверке раздела ----------
def document_budget_chars(cfg, extra_chars=0):
    """Сколько символов документа можно добавить к проверке части: окно модели минус ответ, промпт, нормы,
    сама часть (до SOFT_FACTOR × audit_map_size) и черновик замечаний для самокритики."""
    window = cfg.llm_context_tokens or FALLBACK_LIMITS["llm_context_tokens"]
    cpt = cfg.chars_per_token or FALLBACK_LIMITS["chars_per_token"]
    total = int((window - cfg.max_tokens) * cpt * LIMIT_SAFETY)
    used = (AUDIT_RESERVE_CHARS + prompt_extra_chars(AUDIT_PROMPT_KEYS) + cfg.audit_context_max_chars
            + int(cfg.audit_map_size * 1.5) + 2000 + extra_chars)
    return total - used


def document_mode(md_text, parts, cfg, extra_chars=0):
    """Как модель видит документ при проверке части: "single" — документ в одну часть (он и есть часть),
    "full" — документ целиком для справки, "outline" — только оглавление (документ не помещается в окно
    или audit_full_document выключен)."""
    if len(parts) <= 1:
        return "single"
    if cfg.audit_full_document and len(md_text) <= document_budget_chars(cfg, extra_chars):
        return "full"
    return "outline"


def _part_keys(text):
    return {t for t in bm25_tokenize(text) if t not in BM25_STOP_WORDS and len(t) >= 3}


def related_parts(query, parts, exclude=None, top=2, min_shared=2):
    """Части документа, ближе всего к тексту query по словам (вес — редкость слова в документе)."""
    keys = [_part_keys(p["text"]) for p in parts]
    df = Counter(k for ks in keys for k in ks)
    q = _part_keys(query)
    scored = []
    for i, ks in enumerate(keys):
        if i == exclude:
            continue
        shared = q & ks
        if len(shared) >= min_shared:
            scored.append((sum(np.log(1 + len(parts) / df[k]) for k in shared), i))
    scored.sort(key=lambda x: (-x[0], x[1]))
    return [i for _s, i in scored[:top]]


def _fragments(parts, indices, max_chars):
    out = []
    for i in indices:
        label = section_label(parts[i]["path"]) or tr("начало документа")
        out.append(f"=== {label} ===\n" + smart_truncate(parts[i]["text"], max_chars=max_chars))
    return "\n\n".join(out)


# ---------- Сверка «не описано» с другими разделами (документ целиком не помещается) ----------
_ABSENCE_RE = re.compile(
    r"(?i)\bне\s+(?:описан|указан|определ|привед|раскрыт|задан|установлен|предусмотрен|отражен|отражён|"
    r"зафиксирован|регламентирован|конкретизирован)\w*|\bотсутству\w*|\bнет\s+(?:описани|указани|информаци|"
    r"сведени|требовани)\w*|\bnot\s+(?:described|specified|defined|stated|given|covered)\b|\bmissing\b|\babsent\b")
ABSENCE_MAX = 15            # замечаний «не описано» на аудит, которые сверяются с другими разделами


def verify_absence(res, idx, ctx, cfg):
    """Замечания «не описано X» части idx сверяются с ближайшими по словам другими частями документа: если модель
    находит X там (и цитата подтверждается текстом), замечание снимается. Нужна, когда модель видит только
    оглавление; с документом целиком её роль выполняют проверка раздела и самокритика."""
    parts = ctx["parts"]
    lock = ctx.setdefault("lock", threading.Lock())
    kept = []
    for f in res["findings"]:
        if not _ABSENCE_RE.search(f["text"]):
            kept.append(f)
            continue
        near = related_parts(f["text"], parts, exclude=idx)
        with lock:                       # разделы проверяются параллельно — общий счётчик сверок
            allowed = bool(near) and ctx["absence_left"] > 0
            if allowed:
                ctx["absence_left"] -= 1
        if not allowed:
            kept.append(f)
            continue
        fragments = _fragments(parts, near, max_chars=cfg.audit_map_size)
        raw = ask_llm(prompt_text("audit_absence_system"), render("audit_absence_user", finding=f["text"],
                                                                  fragments=fragments), cfg, guard=True)
        data = extract_json(raw, kind="object") if not is_llm_error(raw) else None
        quote = str((data or {}).get("quote") or "").strip().strip("«»\"") if isinstance(data, dict) else ""
        if isinstance(data, dict) and data.get("found") is True and quote_in_text(quote, fragments):
            res["dropped"].append(f"{f['text']} — " + tr("есть в другом разделе документа: «{q}»", q=quote[:200]))
        else:
            kept.append(f)
    if len(kept) != len(res["findings"]):
        res["findings"] = kept
        res["report"] = "\n".join(format_finding(f) for f in kept) if kept else tr("Часть чистая")


# ---------- Ответы на пункты: «особое внимание», чек-лист, требования к системам ----------
# Каждый пункт — отдельный вопрос модели по всему документу (или по разделам на его тему, если документ не
# помещается). Замечания пункта помечаются focus_item (номер пункта) и point_kind ("focus", "checklist",
# "system") и показываются в заключении под своим пунктом.
FOCUS_POINTS_MAX = 8
SYSTEMS_MAX = 5             # систем с карточками на аудит: по каждой — вопрос «что документ требует от системы»
SYSTEM_MIN_MENTIONS = 2     # столько раз система должна упоминаться в документе, чтобы задать вопрос
_ANSWER_RE = re.compile(r"^\s*\**\s*(?:Ответ|Answer)\s*\**\s*:\s*\**\s*(.*)$", re.IGNORECASE)
_VERDICT_WORDS = (("не относится", "na"), ("неприменим", "na"), ("not applicable", "na"), ("n/a", "na"),
                  ("частично", "partial"), ("partially", "partial"), ("partly", "partial"), ("partial", "partial"),
                  ("да", "yes"), ("yes", "yes"), ("нет", "no"), ("no", "no"))
VERDICT_ICONS = {"yes": "✅", "partial": "⚠️", "no": "❌", "na": "➖", "unknown": "❔"}
_STATUS_RE = re.compile(
    r"^\s*(?:[-–—•*·]+|\d{1,2}[.)])?\s*\**\s*\[\s*(ЕСТЬ|В\s+ПЛАНАХ|ДОРАБОТКА|ПРОТИВОРЕЧИЕ|НЕЯСНО|HAVE|PLANNED|CHANGE|"
    r"CONFLICT|UNCLEAR)\s*\]\s*\**\s*[:—–-]?\s*(.+)$", re.IGNORECASE)
_STATUS_CODES = {"ЕСТЬ": "have", "HAVE": "have", "В ПЛАНАХ": "planned", "PLANNED": "planned", "ДОРАБОТКА": "change",
                 "CHANGE": "change", "ПРОТИВОРЕЧИЕ": "conflict", "CONFLICT": "conflict", "НЕЯСНО": "unclear",
                 "UNCLEAR": "unclear"}
STATUS_ORDER = ("conflict", "change", "unclear", "planned", "have")       # сначала то, что требует внимания
STATUS_LABELS = {"have": N_("✅ есть"), "planned": N_("🗓️ в планах"), "change": N_("🛠️ нужна доработка"),
                 "conflict": N_("⛔ противоречие"), "unclear": N_("❓ неясно")}
NO_REQUIREMENTS = ("требований к системе нет", "no requirements for the system")


def focus_points(special):
    """Пункты «особого внимания»: через «;», с новой строки или вопросы («…?»)."""
    items = [x.strip(" \t-•*–—") for x in re.split(r"(?<=\?)\s+|[;\n]+", special or "")]
    return [x for x in items if len(x) >= 3][:FOCUS_POINTS_MAX]


def checklist_verdict(answer):
    """Ответ модели на пункт чек-листа → "yes" / "partial" / "no" / "na" / "unknown" (по первому слову)."""
    low = (answer or "").strip().lower().lstrip("«\"'*_ ")
    for word, code in _VERDICT_WORDS:
        if low.startswith(word) and (len(low) == len(word) or not low[len(word)].isalnum()):
            return code
    return "unknown"


def _focus_context(point, ctx, cfg):
    if ctx["mode"] in ("full", "single", "whole"):
        return document_block(ctx["md"])
    near = related_parts(point, ctx["parts"], top=3, min_shared=1)
    fragments = _fragments(ctx["parts"], near, max_chars=cfg.audit_map_size * 2) if near else ""
    return ((render("audit_fragments", fragments=fragments) + "\n\n") if fragments else "") + outline_block(ctx["outline"])


def _point_pinned(point, ctx, cfg):
    """Документы-основания, правила экспертов и карточки систем, упомянутых в документе, — ближайшие к пункту, если
    всё не помещается."""
    return (_basis_pinned(ctx, point, cfg)
            + select_pinned(ctx.get("expert"), point, ctx.get("pinned_budget", 0), mention_text=ctx["md"]))


_ITEM_START_RE = re.compile(r"^\s*(?:[-–—•·*]|\d{1,2}[.)])\s+\S")     # «- текст», «• текст», «2. текст»


def _split_answer(raw):
    """Ответ модели на пункт → (строка «Ответ: …», остальные строки — замечания). Если тегов критичности в
    ответе нет совсем, замечания — пункты списка после строки «Ответ:» (модель забыла теги; раньше такие пункты
    склеивались в один абзац ответа и замечаниями не считались)."""
    answer, rest = [], []
    tagless = not _SEV_RE.search(raw)
    for line in raw.splitlines():
        m = _ANSWER_RE.match(line)
        if m and not answer:
            answer.append(m.group(1).strip())
        elif (_SEV_RE.search(line) or (rest and line.strip())
              or (tagless and answer and _ITEM_START_RE.match(line) and len(strip_list_marker(line)) > 15)):
            rest.append(line)
        elif answer and line.strip() and not rest:
            answer.append(line.strip())         # ответ в несколько строк до первого замечания
    return " ".join(answer).strip() or ("" if rest else raw.strip()), rest


def _point_findings(lines, i, kind):
    out = []
    for line in parse_findings("\n".join(lines)):
        f = finding_from_line(line)
        if len(f["text"]) > 15:
            out.append({**f, "focus_item": i, "point_kind": kind, "sections": []})
    return cap_findings(out, FOCUS_FINDINGS_MAX)[0]


def answer_focus_point(i, point, ctx, cfg, attempts=1, kind="focus"):
    """Ответ на один пункт «особого внимания» (kind="focus") или чек-листа (kind="checklist") по всему документу
    (или по разделам на его тему, если документ не помещается): строка «Ответ: …» и замечания. Возвращает запись
    для заключения и лога (у пункта чек-листа — verdict: yes / partial / no / na / unknown)."""
    t0 = time.time()
    norms_text, sources = "", []
    if ctx["use_wiki"]:
        pinned = _point_pinned(point, ctx, cfg)
        found, srcs = build_norms_context(point, cfg, ctx["k"], ctx.get("glossary"), skip_ids=pinned_ids(pinned))
        norms_prompt, norms_text, sources = norms_with_pinned(pinned, found, srcs)
    else:
        norms_prompt = autonomous_note()
    context = _focus_context(point, ctx, cfg)
    found = ctx.get("found_so_far") or ""
    found_block = ("\n\n" + render("audit_found_so_far", findings=found)) if found else ""
    if kind == "checklist":
        raw = ask_llm(checklist_prompt(norms_prompt, ctx.get("doc_type", "")),
                      render("audit_checklist_user", item=point, context=context) + found_block, cfg, guard=True)
    else:
        raw = ask_llm(focus_answer_prompt(norms_prompt),
                      render("audit_focus_user", item=point, context=context) + found_block, cfg, guard=True)
    entry = {"kind": kind, "index": i, "item": point, "failed": is_llm_error(raw), "answer": "", "findings": [],
             "report": raw.strip() if is_llm_error(raw) else "", "sources": sources, "norms_text": norms_text,
             "dropped": [], "no_norms": bool(ctx["use_wiki"] and not sources), "attempts": attempts}
    if not entry["failed"]:
        entry["answer"], rest = _split_answer(raw)
        if rest:
            entry["findings"] = _point_findings(rest, i, kind)
        entry["report"] = "\n".join([tr("Ответ: {v}", v=entry["answer"])] + [format_finding(f) for f in entry["findings"]])
    if kind == "checklist":
        entry["verdict"] = checklist_verdict(entry["answer"]) if not entry["failed"] else "unknown"
    entry["seconds"] = time.time() - t0
    return entry


def _system_context(card, ctx, cfg):
    """Документ для вопроса о системе: целиком, если помещается, иначе — разделы, где система упоминается."""
    if ctx["mode"] in ("full", "single", "whole"):
        return document_block(ctx["md"])
    pattern = card.get("pattern")
    hits = sorted(((len(pattern.findall(p["text"])) if pattern else 0, -i) for i, p in enumerate(ctx["parts"])),
                  reverse=True)
    near = [-i for n, i in hits if n > 0][:4]
    fragments = _fragments(ctx["parts"], near, max_chars=cfg.audit_map_size * 2) if near else ""
    return ((render("audit_fragments", fragments=fragments) + "\n\n") if fragments else "") + outline_block(ctx["outline"])


def answer_system(i, card, ctx, cfg, attempts=1, mentions=0):
    """Что документ требует от системы (по её карточке из базы): строки со статусами (есть / в планах / доработка
    / противоречие / неясно) и замечания к документу. Возвращает запись для заключения и лога."""
    t0 = time.time()
    name = card["names"][0] if card.get("names") else card["source"]
    card_text = ctx["expert"].card_text(card, max(4000, int(ctx.get("pinned_budget") or 0)))
    raw = ask_llm(card_prompt(name), render("audit_card_user", system=name, card=card_text,
                                            context=_system_context(card, ctx, cfg)), cfg, guard=True)
    failed = is_llm_error(raw)
    entry = {"kind": "system", "index": i, "item": name, "card": card["source"], "names": list(card.get("names") or []),
             "mentions": mentions, "failed": failed, "answer": "", "statuses": [], "findings": [],
             "report": raw.strip() if failed else "", "sources": [f"🗂️ {card['source']}"], "norms_text": card_text,
             "dropped": [], "no_norms": False, "attempts": attempts}
    if not failed:
        rest = []
        for line in raw.split("\n"):
            m = _STATUS_RE.match(line)
            if m:
                code = _STATUS_CODES.get(re.sub(r"\s+", " ", m.group(1).upper()), "unclear")
                entry["statuses"].append((code, m.group(2).strip()))
            elif line.strip():
                rest.append(line)
        if any(_SEV_RE.search(x) for x in rest):
            entry["findings"] = _point_findings(rest, i, "system")
        if not entry["statuses"] and not entry["findings"]:
            entry["answer"] = (tr("Требований к системе нет.") if any(w in raw.lower() for w in NO_REQUIREMENTS)
                               else raw.strip())
        entry["report"] = "\n".join([f"- {tr(STATUS_LABELS[c])}: {t}" for c, t in entry["statuses"]]
                                    + [format_finding(f) for f in entry["findings"]]) or entry["answer"]
    entry["seconds"] = time.time() - t0
    return entry


# ---------- Лог проверки по разделам ----------
def audit_log_entry(part, total, preview, res, use_wiki):
    """Запись лога одного раздела: хранится в сеансе, чтобы не пропадать после скачивания файлов.
    seconds — сколько шла проверка раздела; attempts — сколько раз он проверялся (больше 1 — после сбоя)."""
    return {"part": part, "total": total, "preview": preview, "failed": bool(res["failed"]),
            "report": (res["report"] or "").strip(), "sources": list(res["sources"] or []),
            "dropped": list(res["dropped"] or []), "no_norms": bool(use_wiki and not res["sources"]),
            "seconds": res.get("seconds"), "attempts": res.get("attempts", 1)}


def section_timing(entry):
    """«42 с» или «42 с, проверен со 2-й попытки» — для лога раздела."""
    if entry.get("seconds") is None:
        return ""
    text = tr("{s} с", s=round(entry["seconds"]))
    if entry.get("attempts", 1) > 1:
        text += ", " + tr("проверен с {n}-й попытки", n=entry["attempts"])
    return text


def _log_entry_lines(e):
    lines = []
    if section_timing(e):
        lines += [f"⏱ {section_timing(e)}", ""]
    if e["failed"]:
        return lines + [tr("**Не проверено:** {report}", report=e["report"])]
    lines.append(e["report"] or "—")
    if e["sources"]:
        lines += ["", tr("**Нормативные источники:** ") + "; ".join(e["sources"])]
    elif e["no_norms"]:
        lines += ["", tr("**Нормативные источники:** в базе не найдены")]
    if e["dropped"]:
        lines += ["", tr("**Отклонено:**")] + [f"- {d}" for d in e["dropped"]]
    return lines


def point_title(e, level=""):
    """Заголовок ответа на пункт в логе: особое внимание, чек-лист или система."""
    if e.get("kind") == "checklist":
        return level + tr("✅ Чек-лист: «{item}»", item=e["item"])
    if e.get("kind") == "system":
        return level + tr("🧩 Система «{item}» (карточка {card})", item=e["item"], card=e.get("card", ""))
    return level + tr("🎯 Особое внимание: «{item}»", item=e["item"])


def audit_log_markdown(filename, log, meta, focus_log=None, notes=None):
    """Лог проверки одним файлом: настройки прогона и по каждому разделу — результат, источники, отклонённое;
    ответы на «особое внимание» и итоговые шаги (notes — [(заголовок, строки)])."""
    lines = [tr("# Лог проверки: {filename}", filename=filename), ""]
    for k, v in meta.items():
        if v not in (None, ""):
            lines.append(f"- **{k}:** {v}")
    for e in log:
        lines += ["", tr("## Часть {part}/{total} — «{preview}»", part=e['part'], total=e['total'], preview=e['preview']), ""]
        lines += _log_entry_lines(e)
    for e in focus_log or []:
        lines += ["", point_title(e, level="## "), ""]
        lines += _log_entry_lines(e)
    for title, body in notes or []:
        lines += ["", f"## {title}", ""] + list(body)
    return "\n".join(lines) + "\n"


# ---------- Аудит документа целиком ----------
RETRY_TIMEOUT_MAX = 600     # с: потолок таймаута при повторной проверке раздела


def retry_timeout(cfg):
    """Таймаут для повторной проверки разделов, не проверенных с первого раза: вдвое больше, не меньше 180 с."""
    return int(min(RETRY_TIMEOUT_MAX, max(180, 2 * int(cfg.llm_timeout))))


def _locate_sections(findings, part):
    """Раздел документа у каждого замечания части: по месту цитаты; без цитаты — по пункту в начале замечания
    («2.2.3, шаг 4: …» — раздел 2.2.3 или ближайший объемлющий заголовок части), затем по совпадению слов с
    разделами части (часть может объединять несколько мелких разделов), иначе — раздел, с которого начинается
    часть."""
    spans = part.get("spans") or [(0, part.get("path", ""))]
    numbered = []                                   # (номер заголовка, путь) — заголовки части
    for _s, pth in spans:
        m = _PATH_NO_RE.match((pth or "").split(" → ")[-1].strip() + " ")
        if m:
            numbered.append((m.group(1), pth))
    text = part["text"]
    segments = [text[s:(spans[i + 1][0] if i + 1 < len(spans) else len(text))] for i, (s, _p) in enumerate(spans)]
    seg_keys = [_part_keys(seg) for seg in segments]
    norm, imap = None, None
    for f in findings:
        path = part.get("path", "")
        if f.get("quote"):
            if norm is None:
                norm, imap = _norm_with_map(text)
            k = norm.find(norm_for_match(f["quote"]))
            if k != -1:
                path = part_path_at(part, imap[k])
        elif numbered and (c := finding_clause(f["text"])) and (
                best := max(((n, pth) for n, pth in numbered if c == n or c.startswith(n + ".")),
                            key=lambda x: len(x[0]), default=None)):
            path = best[1]
        elif len(segments) > 1:
            fk = _part_keys(f["text"])
            scored = sorted(((len(fk & ks), -i) for i, ks in enumerate(seg_keys)), reverse=True)
            if scored and scored[0][0] >= 2 and (len(scored) == 1 or scored[0][0] > scored[1][0]):
                path = spans[-scored[0][1]][1]
        label = section_label(path)
        f["sections"] = [label] if label else []


def _basis_pinned(ctx, text, cfg):
    """Фрагменты документов-оснований, ближайшие к тексту раздела или пункта — по словам и по смыслу."""
    basis, budget = ctx.get("basis"), ctx.get("basis_budget", 0)
    if basis is None or budget <= 0:
        return []
    return select_pinned(basis, text, budget, query_vec=query_vector(basis, text, cfg))


def _check_part(sections, idx, cfg, ctx, attempts=1):
    """Проверка одного раздела с замером времени: (результат audit_section, запись лога)."""
    with answer_label(tr("Часть {n}/{t}", n=idx + 1, t=len(sections))):
        return _check_part_labeled(sections, idx, cfg, ctx, attempts)


def _check_part_labeled(sections, idx, cfg, ctx, attempts=1):
    sec = sections[idx]
    part = ctx["parts"][idx] if ctx.get("parts") else {"text": sec, "path": "", "spans": [], "continued": False}
    t0 = time.time()
    if ctx.get("whole_tokens"):        # весь документ одним запросом: длинный ответ и долгий запрос
        cfg = replace(cfg, max_tokens=max(cfg.max_tokens, ctx["whole_tokens"]),
                      llm_timeout=max(cfg.llm_timeout, WHOLE_TIMEOUT))
    pinned = _basis_pinned(ctx, sec, cfg) + select_pinned(ctx.get("expert"), sec, ctx.get("pinned_budget", 0))
    whole = bool(ctx.get("whole_tokens"))
    res = audit_section(sec, cfg, ctx["focus"], ctx["use_wiki"], ctx["k"], glossary=ctx.get("glossary"), part=idx + 1,
                        special=ctx["special"], focus_items=ctx["focus_items"], doc_block=ctx.get("doc_block", ""),
                        location=location_line(part["path"], part["continued"]),
                        max_findings=WHOLE_FINDINGS_MAX if whole else SECTION_FINDINGS_MAX,
                        prompt_max=WHOLE_PROMPT_FINDINGS if whole else SECTION_PROMPT_FINDINGS,
                        norm_queries=ctx.get("norm_queries"), pinned=pinned,
                        clauses=None if whole or len(sections) < 2 else (part_clauses(part), ctx.get("clauses") or set()))
    if not res["failed"]:
        if ctx.get("mode") == "outline":
            verify_absence(res, idx, ctx, cfg)
        _locate_sections(res["findings"], part)
    res["seconds"], res["attempts"] = time.time() - t0, attempts
    preview = sec[:120].replace("\n", " ").strip() + "..."
    return res, audit_log_entry(idx + 1, len(sections), preview, res, ctx["use_wiki"])


# Сбои, которые проходят со временем: таймаут, перегрузка шлюза, сеть (на двух языках интерфейса). Ошибка
# доступа, настройки или «контекст не помещается» от повтора не исчезнет.
_TRANSIENT = ("[ТАЙМАУТ", "[MINT TIMEOUT", "ИНФРАСТРУКТУРЫ", "INFRASTRUCTURE", "СЕТЕВОЙ СБОЙ", "NETWORK FAILURE")


def _recheck(state, cfg, emit, timeout=None, automatic=False):
    """Повторная проверка разделов и ответов на «особое внимание», не полученных из-за сбоя модели (обычно
    таймаут), с бо́льшим таймаутом. automatic — повтор в конце аудита: только преходящие сбои и только если хоть
    один раздел проверен (если не ответила ни на один, шлюз, скорее всего, недоступен — повтор лишь удвоил бы
    ожидание)."""
    failed = [i for i, r in enumerate(state["results"]) if r["failed"]]
    points = [(kind, state.setdefault(kind, [])) for kind in ("focus", "checklist", "systems")]
    failed_points = [(kind, lst, i) for kind, lst in points for i, e in enumerate(lst) if e["failed"]]
    deep = state.setdefault("deep", [])
    failed_deep = [i for i, d in enumerate(deep) if d["entry"]["failed"]]
    if automatic:
        if len(failed) == len(state["results"]):
            return
        failed = [i for i in failed if any(m in state["results"][i]["report"] for m in _TRANSIENT)]
        failed_points = [x for x in failed_points if any(m in x[1][x[2]]["report"] for m in _TRANSIENT)]
        failed_deep = [i for i in failed_deep if any(m in deep[i]["entry"]["report"] for m in _TRANSIENT)]
    if not failed and not failed_points and not failed_deep:
        return
    rcfg = replace(cfg, llm_timeout=timeout or retry_timeout(cfg), llm_retries=1)
    emit("retry", n=len(failed) + len(failed_points) + len(failed_deep), timeout=rcfg.llm_timeout)
    for n, idx in enumerate(failed):
        res, entry = _check_part(state["sections"], idx, rcfg, state["ctx"],
                                 attempts=state["results"][idx].get("attempts", 1) + 1)
        state["results"][idx], state["log"][idx] = res, entry
        emit("section", entry=entry, index=n, total=len(failed), retry=True)
    for kind, lst, i in failed_points:
        lst[i] = _answer_job((kind, i, _job_item(state["ctx"], kind, lst[i])), state["ctx"], rcfg,
                             attempts=lst[i].get("attempts", 1) + 1)
        emit(_POINT_EVENTS[kind], entry=lst[i], index=i, total=len(lst), retry=True)
    if failed_deep:
        def done(n, entry):
            d = deep[failed_deep[n]]
            entry["attempts"] = d["entry"].get("attempts", 1) + 1
            d["entry"] = entry
            emit("deep_task", entry=entry, index=n, total=len(failed_deep), retry=True)
        run_tasks([deep[i]["task"] for i in failed_deep], rcfg, on_done=done, parallel=cfg.audit_parallel)


_POINT_EVENTS = {"focus": "focus_answer", "checklist": "checklist_answer", "systems": "system_answer"}


def _job_item(ctx, kind, entry):
    """Что передать заново при повторе пункта: текст пункта или карточку системы с числом упоминаний."""
    if kind == "systems":
        return ctx["system_cards"][entry["index"]]
    return entry["item"]


def _answer_job(job, ctx, cfg, attempts=1):
    kind, i, item = job
    labels = {"focus": N_("Особое внимание {n}"), "checklist": N_("Чек-лист {n}"), "systems": N_("Карточка системы {n}")}
    with answer_label(tr(labels.get(kind, labels["focus"]), n=i + 1)):
        return _answer_job_labeled(kind, i, item, ctx, cfg, attempts)


def _answer_job_labeled(kind, i, item, ctx, cfg, attempts):
    if kind == "systems":
        card, mentions = item
        return answer_system(i, card, ctx, cfg, attempts=attempts, mentions=mentions)
    return answer_focus_point(i, item, ctx, cfg, attempts=attempts, kind="checklist" if kind == "checklist" else "focus")


def _run_points(state, cfg, emit, checklist=()):
    """Ответы на пункты «особого внимания», чек-листа и вопросы о системах — параллельно (audit_parallel)."""
    ctx = state["ctx"]
    jobs = ([("focus", i, p) for i, p in enumerate(focus_points(ctx["special"]))]
            + [("checklist", i, p) for i, p in enumerate(list(checklist)[:CHECKLIST_MAX])]
            + [("systems", i, c) for i, c in enumerate(ctx.get("system_cards") or [])])
    for kind in _POINT_EVENTS:
        state[kind] = [None] * sum(1 for j in jobs if j[0] == kind)
    if not jobs:
        return
    emit("points", focus=len(state["focus"]), checklist=len(state["checklist"]),
         systems=[c["names"][0] if c.get("names") else c["source"] for c, _n in ctx.get("system_cards") or []],
         parallel=cfg.audit_parallel)

    def done(n, entry):
        kind, i, _item = jobs[n]
        state[kind][i] = entry
        emit(_POINT_EVENTS[kind], entry=entry, index=i, total=len(state[kind]))

    parallel_map(lambda job: _answer_job(job, ctx, cfg), jobs, parallel=cfg.audit_parallel, on_done=done)


def _run_deep(state, cfg, emit):
    """Глубокая проверка (deep_checks): узкие вопросы модели по процессу, цифрам, утверждениям частей и опечаткам,
    затем — сравнение утверждений одной темы из разных разделов. Не больше audit_deep_max_calls вопросов."""
    ctx = state["ctx"]
    budget = int(cfg.audit_deep_max_calls or 0)
    t0 = time.time()
    # утверждения и опечатки — по разделам и в режиме «весь документ одним запросом»: один вопрос на весь документ
    # модель среднего размера решает хуже (и опечаток из него берётся не больше DEEP_FINDINGS_MAX)
    state["processes"] = process_summary(find_processes(ctx["blocks"]))
    tasks, skipped = limit_tasks(build_deep_tasks(ctx["blocks"], ctx.get("deep_parts") or ctx["parts"]), budget)
    emit("deep", n=len(tasks), skipped=skipped, parallel=cfg.audit_parallel)
    entries = run_tasks(tasks, cfg, parallel=cfg.audit_parallel,
                        on_done=lambda i, e: emit("deep_task", entry=e, index=i, total=len(tasks)))
    state["deep"] = [{"task": t, "entry": e} for t, e in zip(tasks, entries)]
    second = conflict_tasks([c for e in entries for c in e.get("claims") or []])
    if budget > 0:
        left = max(0, budget - len(tasks))
        skipped += max(0, len(second) - left)
        second = second[:left]
    if second:
        emit("deep", n=len(second), skipped=0, parallel=cfg.audit_parallel, stage=2)
        entries2 = run_tasks(second, cfg, parallel=cfg.audit_parallel,
                             on_done=lambda i, e: emit("deep_task", entry=e, index=i, total=len(second)))
        state["deep"] += [{"task": t, "entry": e} for t, e in zip(second, entries2)]
    state["deep_skipped"] = skipped
    state["deep_wall"] = time.time() - t0


PASSES_MAX = 3
PASS_TEMPERATURE = 0.4      # проходы 2–3: температура не ниже этой — модель смотрит на документ немного иначе
# что из глубокой проверки повторяется в проходах 2–3: утверждения частей нужны один раз (для противоречий),
# опечатки — тоже
PASS_DEEP_KINDS = ("boundary", "step", "requirement", "numbers", "conflict")


def audit_passes(cfg):
    return max(1, min(PASSES_MAX, int(cfg.audit_passes or 1)))


def _run_passes(state, cfg, emit):
    """Проходы 2–3 (audit_passes): проверка разделов и вопросы глубокой проверки о процессе, требованиях, цифрах
    и противоречиях — ещё раз, с температурой не ниже PASS_TEMPERATURE. Замечания проходов сводятся вместе с
    первым: то, что один проход случайно пропустил, находит другой, а найденное несколько раз выше в заключении.
    Разделы и вопросы, на которые модель в этих проходах не ответила, не перепроверяются — их проверил первый."""
    total = audit_passes(cfg)
    if total < 2:
        return
    ctx, sections = state["ctx"], state["sections"]
    pcfg = replace(cfg, temperature=max(float(cfg.temperature or 0), PASS_TEMPERATURE))
    deep_tasks = [d["task"] for d in state.get("deep") or [] if d["task"]["kind"] in PASS_DEEP_KINDS]
    for n in range(2, total + 1):
        t0 = time.time()
        ctx["absence_left"] = ABSENCE_MAX           # сверка «не описано» с другими разделами — заново на проход
        emit("pass", n=n, total=total, sections=len(sections), deep=len(deep_tasks), temperature=pcfg.temperature)
        done = []

        def section_done(idx, pair, n=n, done=done):
            done.append(idx)
            emit("section", entry=pair[1], index=len(done) - 1, total=len(sections), pass_no=n)

        with answer_label(tr("Проход {n}", n=n)):
            pairs = parallel_map(lambda idx: _check_part(sections, idx, pcfg, ctx), range(len(sections)),
                                 parallel=cfg.audit_parallel, on_done=section_done)
            entries = (run_tasks(deep_tasks, pcfg, parallel=cfg.audit_parallel,
                                 on_done=lambda i, e: emit("deep_task", entry=e, index=i, total=len(deep_tasks)))
                       if deep_tasks else [])
        state.setdefault("passes", []).append({
            "n": n, "results": [r for r, _e in pairs], "log": [e for _r, e in pairs],
            "deep": [{"task": t, "entry": e} for t, e in zip(deep_tasks, entries)], "seconds": time.time() - t0})


def _pass_note(passes, temperature):
    """Лог: что дали проходы 2–3 — по каждому разделу и вопросу глубокой проверки."""
    body = [tr("Температура проходов 2–3: {t}. Их замечания сведены вместе с первым проходом; «найдено: …» у "
               "замечания в заключении считает и проходы.", t=temperature), ""]
    for p in passes:
        failed = sum(1 for r in p["results"] if r["failed"]) + sum(1 for d in p["deep"] if d["entry"]["failed"])
        found = (sum(len(r["findings"]) for r in p["results"] if not r["failed"])
                 + sum(len(d["entry"]["findings"]) for d in p["deep"] if not d["entry"]["failed"]))
        body += [tr("### Проход {n}: разделов — {s}, вопросов глубокой проверки — {d}, замечаний — {f}, не выполнено "
                    "— {x}; {t} с", n=p["n"], s=len(p["results"]), d=len(p["deep"]), f=found, x=failed,
                    t=round(p["seconds"])), ""]
        for e in p["log"]:
            body += [tr("**Часть {part}/{total} — «{preview}»**", part=e["part"], total=e["total"], preview=e["preview"]),
                     ""] + _log_entry_lines(e) + [""]
        for d in p["deep"]:
            e = d["entry"]
            body += [f"**{e['title']}**", "",
                     tr("**Не выполнено:** {report}", report=e["report"]) if e["failed"] else (e["report"] or "—"), ""]
    return body


REASONING_TIMEOUT = 300         # с: аудит с рассуждениями (audit_reasoning) — таймаут не меньше
REASONING_EXTRA_TOKENS = 6000   # к длине ответа: рассуждения на шлюзах (vLLM, SGLang) входят в max_tokens


def audit_cfg(cfg):
    """Настройки запросов аудита. С рассуждениями (audit_reasoning) — без выключателя рассуждений, с запасом длины
    ответа на них и таймаутом не меньше REASONING_TIMEOUT; запросы при этом идут с включателем (reasoning_on)."""
    if not cfg.audit_reasoning:
        return cfg
    return replace(cfg, llm_no_reasoning=False, llm_timeout=max(int(cfg.llm_timeout), REASONING_TIMEOUT),
                   max_tokens=int(cfg.max_tokens) + REASONING_EXTRA_TOKENS)


def _reasoning_line(cfg):
    """Строка лога «Рассуждения модели»: включены ли для аудита и рассуждала ли модель на самом деле."""
    if not cfg.audit_reasoning:
        return tr("выключены") if cfg.llm_no_reasoning else tr("как настроено на сервере")
    line = tr("включены для аудита (audit_reasoning): ответ до {n} токенов вместе с рассуждениями, таймаут {t} с",
              n=cfg.max_tokens, t=cfg.llm_timeout)
    if reasoning_blocked_by_extra_body(cfg):
        return line + "; " + tr("но llm_extra_body их выключает — свои параметры запроса важнее")
    return line + "; " + (tr("модель рассуждала") if model_reasons(cfg) else
                          tr("рассуждений в ответах не было: шлюз мог их не включить — включатель задаётся в "
                             "llm_extra_body"))


def audit_basis_documents(md_text, filename, cfg, expert=None, use_wiki=True, chosen=None):
    """Документы-основания для аудита. chosen — список документов базы, выбранный в интерфейсе (пустой — без них);
    None — найти по ссылкам документа (если включено audit_basis_docs). Документ, уже отмеченный 📌 или 🗂️, и сам
    проверяемый документ не берутся. → {"docs": [{source, by, ref}], "missing": [ссылки], "refs": [ссылки], "manual",
    "removed": [найденные, но снятые в интерфейсе]}. У выбранного в интерфейсе документа, который MINT нашёл и сам, —
    как он найден (by: code | title); у добавленного вручную — by: manual."""
    info = {"docs": [], "missing": [], "refs": [], "manual": chosen is not None, "off": False, "removed": []}
    if not use_wiki or (chosen is None and not cfg.audit_basis_docs):
        info["off"] = True
        return info
    pinned = {d["source"] for d in (expert.always + expert.cards)} if expert else set()
    try:
        refs = basis_references(md_text)
        found = find_basis_documents(refs, cfg, exclude=[filename] + sorted(pinned))
    except Exception as e:                       # документы-основания — дополнение: без них аудит идёт как обычно
        logger.warning(f"Документы-основания не найдены: {str(e)[:160]}")
        refs, found = [], {"docs": [], "missing": []}
    info.update(refs=refs, missing=found["missing"])
    if chosen is None:
        info["docs"] = found["docs"]
        return info
    auto = {d["source"]: d for d in found["docs"]}
    info["docs"] = [auto.get(s) or {"source": s, "by": "manual", "ref": {}} for s in dict.fromkeys(chosen)
                    if s and s not in pinned]
    info["removed"] = [d for d in found["docs"] if d["source"] not in set(chosen)]
    return info


def basis_line(info):
    """Строка лога: какие документы-основания подключены и по какой ссылке, какие найденные сняты вручную, каких
    нет в базе."""
    if not info:
        return None
    if info.get("off"):
        return tr("выключено")
    by = {"code": N_("по шифру"), "title": N_("по названию"), "manual": N_("выбран вручную")}

    def label(d):
        return (f"{d['source']} ({tr(by.get(d['by'], d['by']))}"
                + (f" {reference_label(d['ref'])}" if d.get("ref") and d["by"] != "manual" else "") + ")")

    parts = [label(d) for d in info["docs"]]
    if info.get("removed"):
        parts.append(tr("найдены, но сняты вручную: {v}", v=", ".join(label(d) for d in info["removed"])))
    if info["missing"]:
        parts.append(tr("в базе нет: {v}", v="; ".join(reference_label(r) for r in info["missing"][:8])))
    if not parts:
        return tr("не выбраны") if info.get("manual") else (tr("ссылок на документы-основания в документе нет")
                                                           if not info.get("refs") else None)
    return "; ".join(parts)


WHOLE_MAX_TOKENS = 4000     # ответ модели, когда весь документ проверяется одним запросом
WHOLE_TIMEOUT = 300         # с: таймаут такого запроса — ответ длинный

_MODE_LABELS = {"whole": N_("весь документ одним запросом ({n} симв.)"), "full": N_("целиком ({n} симв.)"), "outline": N_("оглавление — документ целиком не помещается в окно "
                                                                  "модели или выключен в настройках"),
                "single": N_("документ — одна часть")}


def expert_line(materials, which):
    """Строка лога: документы базы «учитывать всегда» или карточки систем (с найденными названиями)."""
    if materials is None:
        return None
    if which == "always":
        return "; ".join(d["source"] for d in materials.always) or tr("нет")
    return "; ".join(f"{d['source']} ({', '.join(d['names'])})" for d in materials.cards) or tr("нет")


def _rate_line(rate0):
    """Строка лога: сколько раз за аудит шлюз ответил 429 и сколько длились общие паузы; None — ни разу."""
    if not rate0:
        return None
    now = rate_limit_stats()
    n, paused = now["count"] - rate0["count"], now["paused"] - rate0["paused"]
    if n <= 0:
        return None
    return tr("{n} раз, общие паузы {s} с — если повторяется, уменьшите «Аудит: запросов к модели одновременно»",
              n=n, s=round(paused))


def model_findings(state, points=()):
    """Замечания из ответов модели: проверка разделов, ответы на пункты, глубокая проверка, проходы 2–3. У каждого —
    origins: из какого ответа оно (part3, focus1, deep7; у проходов 2–3 — part3#p2): замечания одного ответа при
    сведении повторов друг с другом не сливаются, а число origins у сведённого — сколько проверок его нашли."""
    raw = []
    part_runs = section_runs(state["ctx"].get("parts") or [])
    for idx, res in enumerate(state["results"]):
        if not res["failed"]:
            raw.extend({**f, "origins": [f"part{idx}"], "run": part_runs.get(idx)} for f in res["findings"])
    for e in points:
        if not e["failed"]:
            raw.extend({**f, "origins": [f"{f.get('point_kind') or 'focus'}{f.get('focus_item')}"]}
                       for f in e["findings"])
    runs = [("", [], state.get("deep") or [])] + [(f"#p{p['n']}", p["results"], p["deep"])
                                                  for p in state.get("passes") or []]
    for suffix, results, deep in runs:
        for idx, res in enumerate(results):
            if not res["failed"]:
                raw.extend({**f, "origins": [f"part{idx}{suffix}"], "run": part_runs.get(idx)}
                           for f in res["findings"])
        for i, d in enumerate(deep):
            e = d["entry"]
            raw.extend({**f, "origins": [f"deep{i}{suffix}"]}
                       for f in (e["findings"] if not e["failed"] else e.get("fallback") or []))
    return raw


def section_runs(parts):
    """Куски одного длинного раздела: {номер части: номер первой части раздела} — для частей, которые продолжают
    предыдущую (тот же раздел, часть начинается не с заголовка). Части-одиночки не попадают."""
    runs, start = {}, None
    for i, p in enumerate(parts):
        prev = parts[i - 1] if i else None
        if prev is not None and p.get("continued") and p.get("path") and p.get("path") == prev.get("path"):
            start = i - 1 if start is None else start
            runs[i - 1] = runs[i] = start
        else:
            start = None
    return runs


FOUND_SO_FAR_MAX_CHARS = 12000      # замечания других проверок в запросе ответа на пункт


def found_so_far_block(state, max_chars=FOUND_SO_FAR_MAX_CHARS):
    """Замечания проверки разделов, глубокой проверки и проходов — для ответов на пункты «особого внимания» и
    чек-листа: ответ согласуется с уже найденным, а замечания пишутся только новые. Без опечаток и замечаний «проблемы
    нет», без повторов по словам; сначала найденные несколькими проверками и критичные. Пустая строка — нечего."""
    found = [f for f in model_findings(state) if not is_typo(f) and not is_no_problem(f["text"])]
    found = sorted(collect_unique_findings(found), key=lambda f: (-corroboration(f), _RANK.get(f["severity"], 1)))
    lines, size = [], 0
    for f in found:
        place = "; ".join(x for x in f.get("sections") or [] if x)
        line = format_finding(f) + (f" — {place}" if place else "")
        if size + len(line) > max_chars:
            break
        lines.append(line)
        size += len(line) + 1
    return "\n".join(lines)


def _assemble(state, filename, cfg, emit=None):
    """Итог аудита из результатов разделов и ответов на «особое внимание»: сверка ссылок и цифр с источниками,
    замечания без повторов (по словам, затем моделью), категории, заключение, лог."""
    emit = emit or (lambda *a, **k: None)
    ctx, sections, results = state["ctx"], state["sections"], state["results"]
    focus = [e for e in state.get("focus") or [] if e]
    checklist = [e for e in state.get("checklist") or [] if e]
    systems = [e for e in state.get("systems") or [] if e]
    points = focus + checklist + systems
    failed_sections = [(idx + 1, state["log"][idx]["preview"], res["report"].strip())
                       for idx, res in enumerate(results) if res["failed"]]
    raw = model_findings(state, points)
    deep = [d["entry"] for d in state.get("deep") or []]
    passes = state.get("passes") or []

    md = ctx.get("md", "")
    norms_all = "\n\n".join([r.get("norms_text") or "" for r in results] + [e.get("norms_text") or "" for e in points]
                             + [r.get("norms_text") or "" for p in passes for r in p["results"]]
                             + [it.get("block") or "" for it in ctx["focus_items"]])
    sources = _norm_src(md + "\n" + norms_all)
    check_clauses = bool(re.search(r"(?m)^#{1,6}\s+\d+(?:\.\d+)+", md))
    findings = []
    for f in raw:
        missing = unverified_refs(f["text"], sources, check_clauses)
        # [Источник: файл, раздел] — файл должен быть среди нормативных фрагментов, которые получила модель
        missing += [c["text"] for c in check_citations(f["text"], norms_all) if c["reason"] == "file"]
        findings.append({**f, "unverified": list(dict.fromkeys(missing))})
    no_problem = [f for f in findings if not is_typo(f) and is_no_problem(f["text"])]
    findings = [f for f in findings if not any(f is x for x in no_problem)]
    dup_stats = {}
    findings = collect_unique_findings(findings, stats=dup_stats)
    # опечатки сводятся без модели (одна исправляемая фраза — одна строка), остальное — моделью
    typo_all = [f for f in findings if is_typo(f)]
    typo_list = clean_typos(typo_all)
    findings = [f for f in findings if not is_typo(f)]
    merge_line, merge_groups = tr("Выключено в настройках."), []
    if cfg.audit_merge_duplicates:
        before = len(findings)
        with answer_label(tr("Сведение повторов")):
            findings, n_groups, err = merge_duplicate_findings(findings, cfg, log=merge_groups)
        merge_line = (tr("Не выполнено: {err}", err=err) if err else
                      tr("Объединено групп повторов: {g}; замечаний было {a}, стало {b}.", g=n_groups, a=before,
                         b=len(findings)))
        emit("merge", groups=n_groups, before=before, after=len(findings), error=err)
    if dup_stats.get("run_merged"):
        merge_line += " " + tr("Повторы в кусках одного длинного раздела склеены без модели: {n}.",
                               n=dup_stats["run_merged"])
    if typo_all:
        merge_line += " " + tr("Опечатки — без модели: было {a}, стало {b}.", a=len(typo_all), b=len(typo_list))
    if no_problem:
        merge_line += " " + tr("Убрано замечаний, где модель сама пишет, что проблемы нет: {n}.", n=len(no_problem))
    findings += typo_list
    unverified = [f for f in findings if f.get("unverified")]
    confirmed = [f for f in findings if not f.get("unverified")]
    by_kind = {"focus": [], "checklist": [], "system": []}
    for f in confirmed:
        if f.get("focus_item") is not None:
            by_kind.get(f.get("point_kind") or "focus", by_kind["focus"]).append(f)
    focus_findings = by_kind["focus"]
    main = [f for f in confirmed if f.get("focus_item") is None] + [dict(f) for f in state.get("checks") or []]
    main, minor, typos = split_minor_findings(main)
    by_section = group_by_section(main, section_order(ctx.get("blocks") or []))
    summary, summary_err, summary_line = [], "", tr("Выключено в настройках (audit_summary_max = 0).")
    limit = int(cfg.audit_summary_max or 0)
    if limit > 0:
        cands = summary_candidates(by_section, focus, focus_findings, checklist, by_kind["checklist"], systems,
                                   by_kind["system"])
        if cands:
            with answer_label(tr("Главное")):
                summary, summary_err = select_summary(cands, cfg, special=ctx["special"], limit=limit)
            emit("summary", n=len(summary), error=summary_err)
        summary_line = (tr("Не выбрано моделью: {err} — взяты найденные несколькими проверками и критичные: {n}.",
                           err=summary_err, n=len(summary)) if summary_err else
                        tr("Выбрано моделью: {n} из {m} замечаний.", n=len(summary), m=len(cands)))
    version = prompts_version()
    total = sum(r.get("seconds") or 0 for r in results) + sum(e.get("seconds") or 0 for e in points)
    deep_seconds = sum(e.get("seconds") or 0 for e in deep)
    mode = ctx.get("mode", "outline")
    notes = [
        (tr("Проверки без модели"), [format_finding(f) for f in state.get("checks") or []] or [tr("Замечаний нет.")]),
    ]
    deep_note = ""
    if ctx.get("deep"):
        n_failed = sum(1 for e in deep if e["failed"])
        deep_note = tr("вопросов модели: {n}; не выполнено: {f}; пропущено по лимиту: {s}; заняла {t} с (вопросы "
                       "шли параллельно, в сумме — {sum} с)", n=len(deep), f=n_failed, s=state.get("deep_skipped", 0),
                       t=round(state.get("deep_wall", deep_seconds)), sum=round(deep_seconds))
        body = [deep_note, ""]
        for e in deep:
            timing = section_timing(e)
            body += [f"### {e['title']}" + (f" · ⏱ {timing}" if timing else ""), ""]
            body.append(tr("**Не выполнено:** {report}", report=e["report"]) if e["failed"] else (e["report"] or "—"))
            body.append("")
        notes.append((tr("Глубокая проверка"), body))
    if passes:
        notes.append((tr("Проходы 2–{n}", n=passes[-1]["n"]),
                      _pass_note(passes, max(float(cfg.temperature or 0), PASS_TEMPERATURE))))
    notes.append((tr("Главное"), [summary_line] + [f"- {it['ref']}: {it['gist']}" for it in summary]))
    notes.append((tr("Сведение повторов"), [merge_line] + (["", tr("Группы повторов:")] + [f"- {g}" for g in merge_groups]
                                                          if merge_groups else [])))
    if unverified:
        notes.append((tr("Не подтверждено источником"),
                      [format_finding(f) + " — " + tr("не найдено: {v}", v=", ".join(f["unverified"]))
                       for f in unverified]))
    log_md = audit_log_markdown(filename, state["log"], {
        tr("Версия MINT"): MINT_VERSION,
        tr("Время"): time.strftime("%Y-%m-%d %H:%M"),
        tr("Модель"): cfg.llm_model,
        tr("Промпты"): version,
        tr("Нормативная база"): tr("да") if ctx["use_wiki"] else tr("нет"),
        tr("Самокритика"): tr("да") if cfg.use_self_critique else tr("нет"),
        tr("Рассуждения модели"): _reasoning_line(cfg),
        tr("Проходов проверки (audit_passes)"): audit_passes(cfg),
        tr("Таймаут ответа LLM, с"): cfg.llm_timeout,
        tr("Время аудита (по часам), с"): round(state["wall"]) if state.get("wall") else None,
        tr("Время проверки разделов (сумма по разделам), с"): round(total),
        tr("Запросов к модели одновременно"): cfg.audit_parallel,
        tr("Разделов"): len(sections),
        tr("Документ в проверке раздела"): tr(_MODE_LABELS[mode], n=len(md)),
        tr("Глубокая проверка"): deep_note or tr("выключена"),
        tr("Процесс в документе (глубокая проверка)"): (
            (state.get("processes") or tr("не найден: шаги процесса ищутся в таблицах операций, списках шагов "
                                          "и подразделах «Шаг N»")) if ctx.get("deep") else None),
        tr("Шлюз просил реже (HTTP 429)"): _rate_line(state.get("rate0")),
        tr("Размер раздела (audit_map_size)"): cfg.audit_map_size,
        "Top-K": ctx["k"],
        tr("Объём норм (audit_context_max_chars)"): cfg.audit_context_max_chars,
        tr("Особое внимание"): ctx["special"],
        tr("Нормы по особому вниманию"): "; ".join(it["short"] for it in ctx["focus_items"]),
        tr("Критерии"): ctx["focus"].replace("\n", " "),
        tr("Чек-лист"): (tr("{t}: пунктов {n}", t=ctx.get("doc_type") or tr("общие пункты"), n=len(checklist))
                         if checklist else tr("нет")),
        tr("Правила экспертов (учитывать всегда)"): expert_line(ctx.get("expert"), "always"),
        tr("Карточки систем"): expert_line(ctx.get("expert"), "cards"),
        tr("Системы, упомянутые в документе"): "; ".join(
            tr("{name} — {n} упом.", name=e["item"], n=e.get("mentions", 0)) for e in systems),
        tr("Документы-основания (audit_basis_docs)"): basis_line(ctx.get("basis_info")),
        tr("Правила и карточки на запрос (audit_pinned_max_chars)"): (cfg.audit_pinned_max_chars
                                                                      if ctx.get("expert") else None),
    }, focus_log=points, notes=notes)
    return {
        "findings": main + focus_findings + by_kind["checklist"] + by_kind["system"] + minor + typos,
        "minor": minor, "typos": typos, "unverified": unverified, "by_section": by_section,
        "focus_answers": focus, "checklist_answers": checklist, "system_answers": systems,
        "doc_type": ctx.get("doc_type", ""),
        "report": build_final_report(filename, by_section, failed_sections, len(sections), special=ctx["special"],
                                     focus_answers=focus, focus_findings=focus_findings, unverified=unverified,
                                     deep_note=deep_note, minor=minor, typos=typos, checklist=checklist,
                                     checklist_findings=by_kind["checklist"], doc_type=ctx.get("doc_type", ""),
                                     systems=systems, system_findings=by_kind["system"], summary=summary,
                                     summary_by_model=not summary_err),
        "summary": summary,
        "deep": deep,
        "log": state["log"], "log_md": log_md, "failed_sections": failed_sections, "sections": len(sections),
        "focus": ctx["focus"], "special": ctx["special"], "focus_items": ctx["focus_items"], "prompts": version,
        "document_mode": mode, "state": state, "basis": ctx.get("basis_info"),
    }


def run_audit(md_text, filename, cfg, focus="", special="", use_wiki=True, glossary=None, on_progress=None,
              retry_failed=True, deep=None, whole=False, checklist=None, doc_type="", systems=None, basis=None):
    """Аудит документа (см. _run_audit). Каждый запрос к модели и её ответ слово в слово — в state["raw"]
    (raw_answers_markdown)."""
    raw = []
    cfg = audit_cfg(cfg)
    with record_answers(raw), reasoning_on(cfg.audit_reasoning):
        res = _run_audit(md_text, filename, cfg, focus=focus, special=special, use_wiki=use_wiki, glossary=glossary,
                         on_progress=on_progress, retry_failed=retry_failed, deep=deep, whole=whole,
                         checklist=checklist, doc_type=doc_type, systems=systems, basis=basis)
    res["state"]["raw"] = raw
    return res


def _run_audit(md_text, filename, cfg, focus="", special="", use_wiki=True, glossary=None, on_progress=None,
               retry_failed=True, deep=None, whole=False, checklist=None, doc_type="", systems=None, basis=None):
    """Аудит документа (Markdown после разбора документа): части по разделам → нормы из базы → замечания →
    самокритика; ответы на пункты «особого внимания» по всему документу; проверки без модели (сокращения, ссылки
    и цифры в замечаниях); сведение повторов; заключение. focus — критерии проверки (пусто — общий аудит),
    special — «особое внимание», use_wiki — подбирать нормы из базы знаний. Если документ помещается в окно
    модели (и включено audit_full_document), каждая часть проверяется вместе с документом целиком — модель
    видит связи между разделами; иначе — с оглавлением, а замечания «не описано» сверяются с другими разделами.
    retry_failed — разделы, на которых модель не ответила (таймаут, сбой шлюза, сеть), в конце проверяются ещё
    раз с бо́льшим таймаутом (retry_timeout()) — если хоть один раздел проверен; recheck_failed_sections()
    перепроверяет по запросу пользователя.

    on_progress(событие, **данные) сообщает о ходе работы: "sections" (n, mode, chars), "checks" (findings),
    "focus_norms" (items), после каждого раздела — "section" (entry — запись лога, index, total; retry=True при
    повторе), после ответа на пункт «особого внимания» — "focus_answer" (entry, index, total), перед глубокой
    проверкой — "deep" (n, skipped, parallel; stage=2 — сравнение утверждений), после каждого её вопроса —
    "deep_task" (entry, index, total), перед повтором — "retry" (n, timeout), после сведения повторов — "merge"
    (groups, before, after, error).

    whole — весь документ одним запросом (для сравнения с проверкой по разделам): одна «часть», длинный ответ
    (WHOLE_MAX_TOKENS) и таймаут WHOLE_TIMEOUT; не помещается в окно модели — событие "whole_too_big" (chars,
    room) и проверка по разделам.

    deep — глубокая проверка (deep_checks: процесс по шагам, требования и границы процесса, противоречия между
    разделами, цифры, опечатки — узкие вопросы, на которые отвечает и модель среднего размера); None — как в
    настройках (audit_deep).

    checklist — пункты чек-листа (expert.checklist_items), doc_type — тип документа для промпта и заключения:
    каждый пункт — вопрос по всему документу ("checklist_answer"). Документы базы «учитывать всегда» и карточки
    систем (expert.load_roles) идут в каждую проверку при use_wiki; systems — вопрос «что документ требует от
    системы» по каждой системе с карточкой, упомянутой в документе не меньше SYSTEM_MIN_MENTIONS раз
    ("system_answer"; None — как в настройках audit_systems). Перед этими вопросами — событие "points"
    (focus, checklist, systems, parallel).

    Возвращает словарь: findings (подтверждённые замечания без повторов, вместе с ответами на «особое
    внимание»), unverified (замечания со ссылками и цифрами, которых нет в источниках), focus_answers, deep
    (записи вопросов глубокой проверки),
    by_section (главные замечания по разделам документа), report (заключение, Markdown), log и log_md (лог проверки), failed_sections, sections (число
    частей), document_mode, focus, special, focus_items (нормы по «особому вниманию»), prompts (версия набора
    промптов), state — для recheck_failed_sections()."""
    def emit(event, **data):
        if on_progress is not None:
            on_progress(event, **data)

    t_start = time.time()
    focus = (focus or "").strip() or generic_audit_focus()
    special = (special or "").strip()
    md_text = compact_markdown_tables(md_text or "")
    expert = None
    if use_wiki:
        try:
            expert = load_expert_materials(cfg)
        except Exception as e:                    # роли и карточки — дополнение: без них аудит идёт как обычно
            logger.warning(f"Правила экспертов и карточки систем не загружены: {str(e)[:160]}")
    pinned_budget = int(cfg.audit_pinned_max_chars or 0) if expert else 0
    basis_info = audit_basis_documents(md_text, filename, cfg, expert, use_wiki, basis)
    basis_mat = basis_materials(cfg, [d["source"] for d in basis_info["docs"]]) if basis_info["docs"] else None
    basis_budget = int(cfg.audit_basis_max_chars or 0) if basis_mat else 0
    if basis_info["docs"] or basis_info["missing"]:
        emit("basis", docs=basis_info["docs"], missing=basis_info["missing"])
    extra = (len(focus) + len(special) + min(pinned_budget, expert.total_chars() if expert else 0)
             + min(basis_budget, basis_mat.total_chars() if basis_mat else 0))
    whole_tokens = 0
    if whole:
        window = cfg.llm_context_tokens or FALLBACK_LIMITS["llm_context_tokens"]
        # длинный ответ, но не больше четверти окна модели: маленькому окну — меньше
        wcfg = replace(cfg, max_tokens=max(cfg.max_tokens, min(WHOLE_MAX_TOKENS, window // 4)))
        room = document_budget_chars(wcfg, extra) + int(cfg.audit_map_size * 1.5)
        if len(md_text) <= room:
            whole_tokens = wcfg.max_tokens
        else:
            emit("whole_too_big", chars=len(md_text), room=max(0, room))
    parts = [whole_document_part(md_text)] if whole_tokens else split_document_parts(md_text, cfg.audit_map_size)
    sections = [p["text"] for p in parts]
    blocks = document_blocks(md_text)
    outline = document_outline(blocks)
    mode = "whole" if whole_tokens else document_mode(md_text, parts, cfg, extra_chars=extra)
    doc_block = document_block(md_text) if mode == "full" else outline_block(outline) if mode == "outline" else ""
    emit("sections", n=len(sections), mode=mode, chars=len(md_text))
    checks = check_abbreviations(blocks, cfg.audit_known_abbreviations) + check_formatting(blocks)
    emit("checks", findings=checks)
    focus_items = special_focus_norms(special, cfg, glossary=glossary) if (special and use_wiki) else []
    if special and use_wiki:
        emit("focus_norms", items=focus_items)
    ctx = {"focus": focus, "special": special, "use_wiki": use_wiki, "k": cfg.top_k or kb_default_top_k(cfg),
           "focus_items": focus_items, "glossary": glossary, "parts": parts, "md": md_text, "outline": outline,
           "mode": mode, "doc_block": doc_block, "absence_left": ABSENCE_MAX, "blocks": blocks,
           "whole_tokens": whole_tokens,
           # весь документ одним запросом: нормы ищутся по каждому разделу, а не по началу документа
           "norm_queries": ([p["text"] for p in split_document_parts(md_text, cfg.audit_map_size)]
                            if whole_tokens else None),
           "deep_parts": split_document_parts(md_text, cfg.audit_map_size) if whole_tokens else None,
           "deep": cfg.audit_deep if deep is None else bool(deep),
           "expert": expert, "pinned_budget": pinned_budget, "doc_type": (doc_type or "").strip(),
           "basis": basis_mat, "basis_budget": basis_budget, "basis_info": basis_info,
           "clauses": text_clauses(md_text),
           "system_cards": (expert.mentioned_cards(md_text, SYSTEM_MIN_MENTIONS)[:SYSTEMS_MAX]
                            if expert and (cfg.audit_systems if systems is None else systems) else [])}
    state = {"sections": sections, "results": [], "log": [], "ctx": ctx, "focus": [], "checklist": [], "systems": [],
             "checks": checks, "deep": [], "passes": [], "rate0": rate_limit_stats()}
    done = []

    def section_done(idx, pair):
        done.append(idx)
        emit("section", entry=pair[1], index=len(done) - 1, total=len(sections))

    pairs = parallel_map(lambda idx: _check_part(sections, idx, cfg, ctx), range(len(sections)),
                         parallel=cfg.audit_parallel, on_done=section_done)
    state["results"] = [r for r, _e in pairs]
    state["log"] = [e for _r, e in pairs]
    if ctx["deep"]:
        _run_deep(state, cfg, emit)
    _run_passes(state, cfg, emit)
    # ответы на пункты — после остальных проверок и с их замечаниями: ответ не расходится с найденным
    ctx["found_so_far"] = found_so_far_block(state)
    _run_points(state, cfg, emit, checklist=[x for x in (checklist or []) if str(x).strip()])
    if retry_failed:
        _recheck(state, cfg, emit, automatic=True)
    state["wall"] = time.time() - t_start
    return _assemble(state, filename, cfg, emit)


def recheck_failed_sections(result, filename, cfg, on_progress=None, timeout=None):
    """Ещё раз проверяет разделы и пункты «особого внимания», которые аудит не смог проверить (result — итог
    run_audit), и собирает итог заново: новые замечания добавляются к прежним, заключение и лог обновляются.
    timeout — таймаут попытки, по умолчанию retry_timeout(cfg)."""
    def emit(event, **data):
        if on_progress is not None:
            on_progress(event, **data)

    state = result["state"]
    t0 = time.time()
    cfg = audit_cfg(cfg)
    with record_answers(state.setdefault("raw", [])), reasoning_on(cfg.audit_reasoning):
        _recheck(state, cfg, emit, timeout=timeout)
        state["wall"] = state.get("wall", 0) + time.time() - t0
        return _assemble(state, filename, cfg, emit)


# ---------- Сверка с эталоном ----------
# Эталон — замечания, которые аудит должен найти (эксперта, прошлого согласования). По нему видно, что находит
# выбранная модель и что пропускает: прогоны с разными моделями и настройками сравниваются на одном документе.
def reference_verdict(value):
    """Вердикт сверки в любом виде («НАЙДЕНО», «✅ найдено», «Не найдено», «Partially», «Да») → found / partial /
    missed / unknown. Шаблон из промпта («НАЙДЕНО|ЧАСТИЧНО|НЕТ») — unknown."""
    raw = str(value or "")
    if "|" in raw:
        return "unknown"
    t = re.sub(r"[^A-ZА-Я ]+", " ", raw.upper().replace("Ё", "Е")).strip()
    if re.search(r"\bЧАСТ|\bPART", t):
        return "partial"
    if re.search(r"\bНЕ\s*НАЙД|\bНЕТ\b|\bNOT\b|\bMISS|\bNO\b|\bОТСУТСТ|\bNONE\b", t):
        return "missed"
    if re.search(r"\bНАЙД|\bFOUND\b|\bYES\b|\bДА\b|\bЕСТЬ\b", t):
        return "found"
    return "unknown"
REFERENCE_VERDICT_LABELS = {"found": N_("✅ найдено"), "partial": N_("⚠️ частично"), "missed": N_("❌ не найдено"),
                            "unknown": N_("❓ не сверено")}


def parse_reference(text):
    """Эталонные замечания: по одному на строку («-», «1.» или без маркера); строка без маркера после пункта со
    списком — продолжение пункта."""
    items = []
    for line in (text or "").split("\n"):
        s = line.strip()
        if not s:
            continue
        marked = bool(re.match(r"^(?:[-•*–—]|\d{1,2}[.)])\s+", s))
        if items and not marked and items[-1].rstrip().endswith((":", ",")):
            items[-1] += " " + s
        else:
            items.append(strip_list_marker(s) if marked else s)
    return [x for x in items if len(x) >= 5]


REFERENCE_BATCH = 3         # пунктов эталона в запросе, если весь эталон за раз модель не осилила


def compare_with_reference(result, reference_text, cfg):
    """Какие замечания эталона нашёл аудит (result — итог run_audit). Сверяет модель: то же место и та же суть,
    формулировка может отличаться. → {rows: [{n, reference, verdict (found|partial|missed|unknown), found
    (тексты замечаний аудита), comment}], found, partial, missed, total, error}."""
    ref = parse_reference(reference_text)
    findings = list(result.get("findings") or []) + list(result.get("unverified") or [])
    out = {"rows": [], "found": 0, "partial": 0, "missed": 0, "total": len(ref), "error": "",
           "audit_total": len(findings), "matched": 0}
    if not ref:
        out["error"] = tr("эталон пуст: впишите замечания, по одному на строку")
        return out
    rows = [{"n": i, "reference": r, "verdict": "unknown", "found": [], "comment": ""} for i, r in enumerate(ref, 1)]
    out["rows"] = rows
    if not findings:
        for r in rows:
            r["verdict"] = "missed"
        out["missed"] = len(rows)
        return out
    listing = "\n".join(f"{i}. [{severity_label(f['severity'])}] {finding_display(f)}" for i, f in enumerate(findings, 1))

    def ask(part):
        raw = ask_llm(prompt_text("audit_reference_system"),
                      render("audit_reference_user", reference="\n".join(f"{r['n']}. {r['reference']}" for r in part),
                             findings=listing), cfg, guard=True)
        data = extract_json(raw, kind="array") if not is_llm_error(raw) else None
        return raw, (data if isinstance(data, list) and data else None)

    odd = []                                     # вердикты, которые не удалось распознать (для ошибки)

    def apply(data, allowed):
        for d in data or []:
            if not isinstance(d, dict):
                continue
            try:
                n = int(d.get("n"))
            except (TypeError, ValueError):
                continue
            if n not in allowed:
                continue
            verdict = reference_verdict(d.get("verdict"))
            if verdict == "unknown":
                odd.append(str(d.get("verdict") or "")[:40])
                continue
            nums = []
            for x in d.get("found") or []:
                try:
                    k = int(x)
                except (TypeError, ValueError):
                    continue
                if 1 <= k <= len(findings) and k not in nums:
                    nums.append(k)
            if verdict == "found" and not nums:
                verdict = "partial"                  # «найдено», но без номера замечания — не засчитываем полностью
            rows[n - 1].update(verdict=verdict, found=[findings[k - 1]["text"] for k in nums], found_ids=nums,
                               comment=str(d.get("comment") or "").strip())

    state = result.get("state") if isinstance(result.get("state"), dict) else {}
    with record_answers(state.setdefault("raw", [])), answer_label(tr("Сверка с эталоном")):
        return _compare_rows(ask, apply, odd, rows, out)


def _compare_rows(ask, apply, odd, rows, out):
    """Сверка по пунктам эталона: весь эталон одним запросом, пропущенное — пачками (см. compare_with_reference)."""
    raw, data = ask(rows)
    if is_llm_error(raw):
        out["error"] = raw.strip()
        return out
    apply(data, {r["n"] for r in rows})
    pending = [r for r in rows if r["verdict"] == "unknown"]
    bad = None if data else raw
    # пункты без ответа (модель не осилила весь эталон, ответ оборван, вердикт не по форме) — ещё раз, пачками
    for k in range(0, len(pending), REFERENCE_BATCH):
        part = pending[k:k + REFERENCE_BATCH]
        braw, bdata = ask(part)
        if not is_llm_error(braw):
            apply(bdata, {r["n"] for r in part})
        if any(r["verdict"] == "unknown" for r in part):
            bad = bad or braw
    left = [str(r["n"]) for r in rows if r["verdict"] == "unknown"]
    if left:
        why = (tr("непонятный вердикт: {v}", v=", ".join(f"«{x}»" for x in dict.fromkeys(odd)))
               if odd and not bad else tr("ответ модели не разобран: «{v}»", v=_short(bad or raw, 200)))
        out["error"] = tr("не сверены пункты {n} — {why}", n=", ".join(left), why=why)
    matched = set()
    for r in rows:
        if r["verdict"] in ("found", "partial", "missed"):
            out[r["verdict"]] += 1
        matched.update(r.get("found_ids") or [])
    out["matched"] = len(matched)
    return out


# ---------- Сырые ответы модели (файл для разбора) ----------
RAW_HEAD, RAW_TAIL = 1500, 1500   # от запроса — начало и конец, символов; ответ модели — целиком
RAW_FENCE = "~~~~~"


def _raw_clip(text, document):
    """Запрос для файла сырых ответов: документ целиком — пометкой (он одинаковый почти во всех запросах), длинное —
    началом и концом (вопрос обычно в конце запроса)."""
    text = text or ""
    if document and len(document) > 500 and document in text:
        text = text.replace(document, tr("[… документ целиком — {n} симв., как в загруженном файле …]", n=len(document)))
    if len(text) <= RAW_HEAD + RAW_TAIL + 200:
        return text
    return (text[:RAW_HEAD] + "\n" + tr("[… пропущено {n} симв. …]", n=len(text) - RAW_HEAD - RAW_TAIL) + "\n"
            + text[-RAW_TAIL:])


def raw_answers_markdown(result, filename):
    """Сырые ответы модели: каждый запрос аудита (и сверки с эталоном, если была) — какая проверка спросила, что
    спросили (начало и конец) и что модель ответила, слово в слово."""
    state = result.get("state") or {}
    raw = sorted(state.get("raw") or [], key=lambda r: r.get("started", 0))
    document = (state.get("ctx") or {}).get("md", "")
    models = ", ".join(dict.fromkeys(r.get("model", "") for r in raw)) or "—"
    lines = [tr("# Сырые ответы модели: {f}", f=filename), "",
             tr("Каждый запрос к модели во время аудита и сверки с эталоном: какая проверка спросила, что спросили "
                "(начало и конец запроса; документ целиком заменён пометкой) и что модель ответила — слово в слово, "
                "до того как MINT разобрал ответ на замечания."), "",
             tr("**Запросов:** {n}; **модель:** {m}.", n=len(raw), m=models), ""]
    for i, r in enumerate(raw, 1):
        when = time.strftime("%H:%M:%S", time.localtime(r.get("started", 0)))
        lines += [f"## {i}. {r.get('label') or tr('Без подписи')} · {r.get('model', '')} · {when} · ⏱ {r.get('seconds', 0)} с",
                  "", tr("**Системный промпт** ({n} симв.):", n=len(r.get("system", ""))),
                  RAW_FENCE + "text", _raw_clip(r.get("system", ""), document), RAW_FENCE, "",
                  tr("**Запрос** ({n} симв.):", n=len(r.get("user", ""))),
                  RAW_FENCE + "text", _raw_clip(r.get("user", ""), document), RAW_FENCE, "",
                  tr("**Ответ модели** ({n} симв., как есть):", n=len(r.get("answer", ""))),
                  RAW_FENCE + "text", r.get("answer", ""), RAW_FENCE, ""]
    return "\n".join(lines) + "\n"


def reference_markdown(cmp, filename, cfg):
    """Итог сверки с эталоном: строка итога и таблица «эталон — вердикт — что нашёл аудит»."""
    lines = [tr("# Сверка с эталоном: {filename}", filename=filename), "",
             tr("**Модель:** {m}; **найдено:** {f} из {n}, **частично:** {p}, **не найдено:** {x}.", m=cfg.llm_model,
                f=cmp["found"], n=cmp["total"], p=cmp["partial"], x=cmp["missed"]),
             tr("**Замечаний в заключении:** {t}; из них к эталону относятся {m} — остальные стоит просмотреть: "
                "полезные или лишние.", t=cmp.get("audit_total", 0), m=cmp.get("matched", 0)), ""]
    if cmp["error"]:
        lines += [tr("⚠️ Сверка не выполнена: {err}", err=cmp["error"]), ""]
    lines += [tr("| № | Эталон | Вердикт | Что нашёл аудит | Комментарий |"), "|---|---|---|---|---|"]
    cell = lambda t: (t or "").replace("|", "\\|").replace("\n", " ")  # noqa: E731
    for r in cmp["rows"]:
        lines.append(f"| {r['n']} | {cell(r['reference'])} | {tr(REFERENCE_VERDICT_LABELS[r['verdict']])} | "
                     f"{cell('; '.join(r['found']))} | {cell(r['comment'])} |")
    return "\n".join(lines) + "\n"


# ---------- Документ с замечаниями (.docx) ----------
# Каждое замечание привязывается к ОДНОМУ месту: по дословной цитате из самокритики, а без цитаты — к абзацу
# с наибольшим совпадением слов (если совпадение достаточное и однозначное). Остальные замечания — списком
# в начале документа. В python-docx ≥ 1.2 замечания добавляются примечаниями Word: текст документа не меняется.

GENERIC_FINDING_WORDS = frozenset(normalize_token(w) for w in (
    "требование требования требований требованиям необходимо указан указано указаны указать отсутствует "
    "отсутствуют описан описано описаны описать определен определено определены должен должна должно должны "
    "следует раздел раздела разделе документ документе тз система системы системе "
    "requirement requirements must should shall missing specified described defined document section system "
    "not specify describe define".split()))

_PLAIN_RUN_TAGS = None


def _norm_with_map(text):
    """Нормализация как в norm_for_match + карта «позиция в нормализованной строке → позиция в исходной»."""
    out, idx, pending_space = [], [], False
    for i, ch in enumerate(text):
        c = ch.lower()
        if c == "ё":
            c = "е"
        if len(c) == 1 and ("0" <= c <= "9" or "a" <= c <= "z" or "а" <= c <= "я"):
            if pending_space and out:
                out.append(" ")
                idx.append(i)
            pending_space = False
            out.append(c)
            idx.append(i)
        else:
            pending_space = True
    return "".join(out), idx


def _docx_paragraphs(doc):
    """Абзацы основного текста, включая таблицы (объединённые ячейки — один раз)."""
    seen, out = set(), []

    def walk(container):
        for par in container.paragraphs:
            if par._p not in seen:
                seen.add(par._p)
                out.append(par)
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell._tc not in seen:
                        seen.add(cell._tc)
                        walk(cell)

    walk(doc)
    return out


def _runs_text(paragraph):
    return "".join(r.text for r in paragraph.runs)


def _is_plain_text_run(run):
    """Run без рисунков, полей и разрывов — его можно безопасно разделить на два."""
    global _PLAIN_RUN_TAGS
    if _PLAIN_RUN_TAGS is None:
        _PLAIN_RUN_TAGS = {qn("w:t"), qn("w:tab"), qn("w:rPr")}
    return all(child.tag in _PLAIN_RUN_TAGS for child in run._r)


def _split_run(run, offset):
    """Делит run по смещению: левая часть остаётся в run, правая — новый run СРАЗУ после него
    (с тем же форматированием). Возвращает правый run или None."""
    text = run.text
    if offset <= 0 or offset >= len(text) or not _is_plain_text_run(run):
        return None
    new_r = OxmlElement("w:r")
    if run._r.rPr is not None:
        new_r.append(copy.deepcopy(run._r.rPr))
    run._r.addnext(new_r)
    right = DocxRun(new_r, run._parent)
    run.text = text[:offset]
    right.text = text[offset:]
    return right


def _isolate_range(paragraph, start, end):
    """Run'ы, покрывающие ровно [start, end) текста абзаца (крайние run'ы делятся, порядок текста сохраняется)."""
    covered, pos = [], 0
    for run in paragraph.runs:
        r_start, r_end = pos, pos + len(run.text)
        pos = r_end
        if r_end > start and r_start < end:
            covered.append([run, r_start, r_end])
    if not covered:
        return []
    last_run, ls, le = covered[-1]
    if end < le:
        _split_run(last_run, end - ls)          # правая часть (после цитаты) отделяется
    first_run, fs, _fe = covered[0]
    if start > fs:
        right = _split_run(first_run, start - fs)
        if right is not None:
            covered[0][0] = right
    return [c[0] for c in covered]


def _locate_by_quote(par_norm, quote):
    """(номер абзаца, start, end) цитаты в тексте абзаца или None. Цитату, охватывающую несколько абзацев,
    ищем по самой длинной её части."""
    nq = norm_for_match(quote)
    candidates = [nq] if len(nq) >= 12 else []
    pieces = sorted({norm_for_match(x) for x in re.split(r"\n+|(?<=[.!?;])\s+", quote or "")}, key=len, reverse=True)
    candidates += [x for x in pieces if len(x) >= 12 and x != nq]
    for cand in candidates:
        for pi, (ntext, imap) in enumerate(par_norm):
            k = ntext.find(cand)
            if k != -1:
                return pi, imap[k], imap[k + len(cand) - 1] + 1
    return None


def _binding_keys(text):
    """Ключи слов для привязки без цитаты: первые 5 букв основы, без служебных и общих слов замечаний."""
    return {t[:5] for t in bm25_tokenize(text)
            if len(t) >= 3 and t not in BM25_STOP_WORDS and t not in GENERIC_FINDING_WORDS}


def _locate_by_words(f_keys, par_keys, idf):
    """Абзац с наибольшим взвешенным совпадением слов; None, если совпадение слабое или неоднозначное."""
    if len(f_keys) < 2:
        return None
    need = max(2, -(-len(f_keys) * 3 // 10))      # не меньше 2 слов и 30 % слов замечания
    scored = []
    for pi, keys in enumerate(par_keys):
        shared = f_keys & keys
        if len(shared) >= need:
            scored.append((sum(idf[k] for k in shared), pi))
    if not scored:
        return None
    scored.sort(key=lambda x: (-x[0], x[1]))
    if len(scored) > 1 and scored[0][0] < 1.2 * scored[1][0]:
        return None                                  # два абзаца подходят почти одинаково — не угадываем
    return scored[0][1]


def _insert_block_at_start(doc, title, lines):
    """Блок «Замечания без привязки» в самом начале документа (до первой таблицы или абзаца)."""
    from docx.text.paragraph import Paragraph
    body = doc.element.body
    pos = 0

    def add(text, bold=False, size=None, color=(200, 0, 0)):
        nonlocal pos
        p_el = OxmlElement("w:p")
        body.insert(pos, p_el)
        pos += 1
        par = Paragraph(p_el, doc._body)
        if text:
            run = par.add_run(text)
            run.font.bold = bold
            run.font.color.rgb = RGBColor(*color)
            if size:
                run.font.size = Pt(size)
        return par

    add(title, bold=True, size=12)
    for line in lines:
        add(line, size=10.5)
    add("")


def annotate_docx_with_findings(docx_bytes, findings, use_comments=None):
    """Документ с замечаниями MINT. Возвращает (байты .docx, статистика {mode, bound, by_quote, unbound}).
    use_comments: None — примечания Word, если их поддерживает python-docx (≥ 1.2), иначе вставки в текст."""
    doc = docx.Document(io.BytesIO(docx_bytes))
    if use_comments is None:
        use_comments = hasattr(doc, "add_comment")
    mode = "comments" if use_comments else "inline"
    paragraphs = _docx_paragraphs(doc)
    texts = [_runs_text(p) for p in paragraphs]
    par_norm = [_norm_with_map(t) for t in texts]
    par_keys = [_binding_keys(t) if len(t.strip()) >= 20 else set() for t in texts]
    df = Counter(k for keys in par_keys for k in keys)
    n_par = max(1, len(paragraphs))
    idf = defaultdict(lambda: 1.0, {k: float(np.log(1 + n_par / v)) for k, v in df.items()})

    placements, unbound, by_quote = [], [], 0
    for f in findings:
        loc = _locate_by_quote(par_norm, f.get("quote") or "") if f.get("quote") else None
        if loc is not None:
            placements.append((loc[0], loc[1], loc[2], f))
            by_quote += 1
            continue
        pi = _locate_by_words(_binding_keys(f["text"]), par_keys, idf)
        if pi is not None:
            placements.append((pi, None, None, f))
        else:
            unbound.append(f)

    # Сначала замечания с цитатой (они делят run'ы), в каждом абзаце — справа налево: смещения левее не сдвигаются
    placements.sort(key=lambda x: (x[0], -(x[1] if x[1] is not None else -1)))
    bound = 0
    for pi, start, end, f in placements:
        paragraph = paragraphs[pi]
        note = f"[{severity_label(f['severity'])}] {f['text']}"
        if start is not None:
            runs = _isolate_range(paragraph, start, end)
        else:
            runs = [r for r in paragraph.runs if r.text]
        if not runs:
            unbound.append(f)
            continue
        if mode == "comments":
            doc.add_comment(runs=runs, text=note, author="MINT", initials="MINT")
        else:
            if start is not None:
                for r in runs:
                    r.font.color.rgb = RGBColor(180, 0, 0)
                    r.font.bold = True
            notice_run = paragraph.add_run(tr("\n⚠️ [ЗАМЕЧАНИЕ MINT]: {note}", note=note))
            notice_run.font.color.rgb = RGBColor(220, 50, 50)
            notice_run.font.bold = True
            notice_run.font.italic = True
            notice_run.font.size = Pt(10.5)
        bound += 1

    if unbound:
        _insert_block_at_start(
            doc, tr("⚠️ Замечания MINT без привязки к тексту ({n_unbound})", n_unbound=len(unbound)),
            [f"• [{severity_label(f['severity'])}] {f['text']}" for f in unbound],
        )
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), {"mode": mode, "bound": bound, "by_quote": by_quote, "unbound": len(unbound)}
