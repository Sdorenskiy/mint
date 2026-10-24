"""Контроль исправления замечаний в новой редакции документа.

Весь контроль без интерфейса — verify_remarks(); им пользуется вкладка «Контроль исправлений».
"""

import io
import logging
import time

import numpy as np

try:
    import docx
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

from .audit import build_norms_context, is_llm_error
from .bm25 import LightweightBM25, best_window, query_term_weights
from .i18n import tr
from .kb import kb_default_top_k
from .prompts import prompt_text, prompts_version, render, review_prompt
from .services import ask_llm, get_query_embedding, get_query_embeddings_batch, rerank_scores
from .text import extract_json, quote_in_text, split_into_sections

logger = logging.getLogger("MINT")


def parse_verification(raw, source_text):
    """Разбирает JSON-ответ проверки замечания. Возвращает (статус_для_таблицы, комментарий).
    «Исправлено»/«Частично» засчитываются только при цитате, реально найденной в фрагментах новой редакции."""
    unverified = tr("❓ Не удалось проверить")
    if is_llm_error(raw):
        return unverified, (raw or "").strip()
    data = extract_json(raw, kind="object")
    if not isinstance(data, dict):
        return unverified, tr("Ответ модели не удалось разобрать: {raw}", raw=(raw or "").strip()[:300])

    status = str(data.get("status", "")).upper().replace("Ё", "Е").replace(" ", "_").replace("-", "_")
    quote = str(data.get("quote") or "").strip()
    comment = str(data.get("comment") or "").strip()

    # статусы русских и английских промптов: ИСПРАВЛЕНО / FIXED, ЧАСТИЧНО / PARTIAL, НЕ_ИСПРАВЛЕНО / NOT_FIXED
    if status.startswith(("НЕ_ИСПРАВ", "NOT_FIX", "NOT_ADDRESS", "UNFIXED")) or "ИГНОР" in status or "IGNOR" in status:
        return tr("❌ Не исправлено"), comment
    # «ИСПРАВЛЕНО_ЧАСТИЧНО» / «FIXED_PARTIALLY» — частично, а не исправлено
    partial = "ЧАСТИЧ" in status or "PARTIAL" in status
    fixed = not partial and status.startswith(("ИСПРАВ", "FIXED", "ADDRESSED", "RESOLVED"))
    if fixed or partial:
        if quote_in_text(quote, source_text):
            label = tr("✅ Исправлено") if fixed else tr("⚠️ Частично")
            return label, tr("{comment} Цитата из новой редакции: «{quote}»", comment=comment, quote=quote).strip()
        return unverified, tr(
            "Модель выставила «{status}», но цитата не найдена в тексте новой редакции — проверьте вручную. "
            "{comment} Цитата модели: «{quote}»", status=status, comment=comment, quote=quote[:200]).strip()
    return unverified, tr("Неожиданный статус в ответе модели: {raw}", raw=(raw or "").strip()[:300])


def build_section_index(sections, cfg):
    """Индекс разделов документа для поиска: эмбеддинги + BM25 (один раз на документ).
    При недоступных эмбеддингах остаётся только BM25."""
    index = {"sections": sections, "mat": None, "ok": None, "bm25": None}
    if not sections:
        return index
    vecs = get_query_embeddings_batch([f"{cfg.embed_doc_prefix}{s[:cfg.section_embed_chars]}" for s in sections], cfg)
    dim = next((len(v) for v in vecs if v is not None), None)
    if dim:
        mat = np.zeros((len(sections), dim), dtype="float32")
        ok = np.zeros(len(sections), dtype=bool)
        for i, v in enumerate(vecs):
            if v is not None and len(v) == dim:
                mat[i] = np.array(v, dtype="float32")
                ok[i] = True
        norms = np.linalg.norm(mat, axis=1, keepdims=True)
        norms[norms == 0] = 1.0
        index["mat"], index["ok"] = mat / norms, ok
    index["bm25"] = LightweightBM25(sections)
    return index


def pick_sections_for_remark(remark, index, cfg, top_n=3, glossary=None):
    """Разделы новой редакции для замечания: dense + BM25 + реранкер (взвешенный RRF).
    Возвращает индексы разделов по убыванию релевантности."""
    sections = index["sections"]
    if not sections:
        return []
    pool = max(10, top_n * 3)

    dense_rank = {}
    if index.get("mat") is not None:
        qv = get_query_embedding(f"{cfg.embed_query_prefix}{remark[:1500]}", cfg)
        if isinstance(qv, list) and qv and not isinstance(qv[0], (dict, list)):
            try:
                q = np.array(qv, dtype="float32")
                if q.shape[0] == index["mat"].shape[1]:
                    q = q / (np.linalg.norm(q) or 1.0)
                    sims = np.where(index["ok"], index["mat"] @ q, -1.0)
                    ranked = [int(i) for i in np.argsort(-sims) if index["ok"][i]][:pool]
                    dense_rank = {i: r for r, i in enumerate(ranked, 1)}
            except Exception as e:
                logger.warning(f"Dense-подбор разделов не удался: {str(e)[:120]}")

    sparse_rank = {}
    extra = glossary.bm25_terms(remark) if glossary is not None else []
    if index.get("bm25") is not None:
        try:
            scores = index["bm25"].get_scores(remark, extra_terms=extra)
            ranked = [int(i) for i in np.argsort(-scores) if scores[i] > 0][:pool]
            sparse_rank = {i: r for r, i in enumerate(ranked, 1)}
        except Exception as e:
            logger.warning(f"BM25-подбор разделов не удался: {str(e)[:120]}")

    candidates = set(dense_rank) | set(sparse_rank)
    if not candidates:
        return []
    weights = cfg.rrf_weights

    def term(rank, weight):
        return weight / (cfg.rrf_constant + rank) if rank else 0.0

    base = {i: term(dense_rank.get(i), weights["dense"]) + term(sparse_rank.get(i), weights["bm25"])
            for i in candidates}
    ordered = sorted(base, key=base.get, reverse=True)

    # Реранкер: все кандидаты (пачками), по самому «плотному» фрагменту раздела, а не по его началу
    rr_ids, rr_docs = [], []
    if cfg.rerank_enabled:
        bm = index.get("bm25")
        weights_q = query_term_weights(remark, extra, idf=bm.idf if bm is not None else None)
        for i in ordered:
            doc, _score = best_window(sections[i], weights_q, cfg.rerank_doc_max_chars)
            if doc.strip():
                rr_ids.append(i)
                rr_docs.append(doc)
    q_text = remark[:cfg.rerank_query_max_chars]
    if glossary is not None:
        q_text = glossary.enrich_query(q_text, max_chars=cfg.rerank_query_max_chars)
    rr_scores = rerank_scores(q_text, rr_docs, cfg) if rr_docs else None

    final = dict(base)
    if rr_scores:
        voting = sorted(((i, s) for i, s in zip(rr_ids, rr_scores) if s is not None and s >= cfg.min_rerank_score),
                        key=lambda kv: kv[1], reverse=True)
        for r, (i, _s) in enumerate(voting, 1):
            final[i] += term(r, weights["rerank"])
    return sorted(final, key=final.get, reverse=True)[:top_n]


def verify_remark(remark, sections, section_index, cfg, use_wiki, k, glossary=None):
    """Проверка одного замечания по новой редакции. Возвращает (статус, комментарий)."""
    if use_wiki:
        wiki_norms, _ = build_norms_context(remark, cfg, k, glossary)
        wiki_part = (render("review_norms", norms=wiki_norms) if wiki_norms else prompt_text("review_no_norms")) + "\n"
    else:
        wiki_part = ""

    picked = pick_sections_for_remark(remark, section_index, cfg, top_n=cfg.review_top_sections, glossary=glossary)
    blocks, total_chars = [], 0
    for si in picked:
        block = render("review_section", n=si + 1, total=len(sections), text=sections[si])
        if blocks and total_chars + len(block) > cfg.review_context_max_chars:
            break
        blocks.append(block)
        total_chars += len(block)
    v2_text_context = "\n\n".join(blocks)
    if not v2_text_context:
        return tr("❓ Не удалось проверить"), tr("Не удалось подобрать релевантные разделы новой редакции "
                                                 "(поиск недоступен или документ пуст).")
    user_content = render("review_user", remark=remark, sections=v2_text_context)
    raw_result = ask_llm(review_prompt(wiki_part), user_content, cfg)
    return parse_verification(raw_result, v2_text_context)


MIN_REMARK_CHARS = 10      # строки короче — не замечания (пустые строки, нумерация, заголовки)


def parse_remarks(text):
    """Список замечаний из текста: каждое с новой строки, строки не длиннее MIN_REMARK_CHARS пропускаются."""
    return [r.strip() for r in (text or "").split("\n") if len(r.strip()) > MIN_REMARK_CHARS]


def is_unverified(status):
    """Статус «не удалось проверить» (сбой, нет цитаты, неразборчивый ответ)."""
    return str(status or "").startswith("❓")


def verify_remarks(remarks, md_text, cfg, use_wiki=True, glossary=None, on_progress=None):
    """Контроль исправлений: каждое замечание проверяется по новой редакции документа (Markdown после разбора
    .docx). remarks — список замечаний или текст (каждое с новой строки).

    on_progress(событие, **данные): "sections" (n), "index" (dense — подобраны ли эмбеддинги разделов; без них
    разделы подбираются только по словам) и после каждого замечания — "remark" (row, index, total).

    Возвращает словарь: rows ([{remark, status, comment}]), matrix_md (матрица, Markdown), unverified (сколько
    замечаний не удалось проверить), sections (число разделов), dense, prompts (версия набора промптов)."""
    def emit(event, **data):
        if on_progress is not None:
            on_progress(event, **data)

    if isinstance(remarks, str):
        remarks = parse_remarks(remarks)
    sections = split_into_sections(md_text, cfg.audit_map_size)
    emit("sections", n=len(sections))
    index = build_section_index(sections, cfg)
    dense = index["mat"] is not None
    emit("index", dense=dense)
    k = cfg.top_k or kb_default_top_k(cfg)
    rows = []
    for i, remark in enumerate(remarks):
        status, comment = verify_remark(remark, sections, index, cfg, use_wiki, k, glossary=glossary)
        rows.append({"remark": remark, "status": status, "comment": comment})
        emit("remark", row=rows[-1], index=i, total=len(remarks))
    return {"rows": rows, "matrix_md": build_matrix_markdown(rows),
            "unverified": sum(1 for r in rows if is_unverified(r["status"])),
            "sections": len(sections), "dense": dense, "prompts": prompts_version()}


def build_matrix_markdown(rows):
    def _md_cell(text):
        return str(text).replace("|", "/").replace("\n", " ")
    md_rows = [f"| {i} | {_md_cell(r['remark'])} | {r['status']} | {_md_cell(r['comment'])} |"
               for i, r in enumerate(rows, 1)]
    return (tr("# МАТРИЦА КОНТРОЛЯ УСТРАНЕНИЯ ЗАМЕЧАНИЙ\n\n")
            + tr("| № | Замечание | Статус ИИ | Заключение |") + "\n|---|---|---|---|\n" + "\n".join(md_rows))


def generate_word_matrix(matrix_data, filename="Матрица_контроля.docx"):
    if not HAS_DOCX:
        return None
    doc = docx.Document()
    font = doc.styles["Normal"].font
    font.name = "Arial"
    font.size = Pt(11)

    title = doc.add_paragraph()
    title_run = title.add_run(tr("МАТРИЦА КОНТРОЛЯ УСТРАНЕНИЯ ЗАМЕЧАНИЙ"))
    title_run.bold = True
    title_run.font.size = Pt(14)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(tr("Дата выгрузки отчета: {date}", date=time.strftime("%d.%m.%Y")))
    if filename:
        doc.add_paragraph(tr("Файл: {filename}", filename=filename))
    doc.add_paragraph("---")

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for cell, text in zip(hdr_cells, ("№", tr("Замечание архитектора"), tr("Статус ИИ"),
                                      tr("Анализ исправления в новой редакции"))):
        cell.text = text
        if cell.paragraphs[0].runs:
            cell.paragraphs[0].runs[0].font.bold = True

    for row_idx, data in enumerate(matrix_data, 1):
        row_cells = table.add_row().cells
        row_cells[0].text = str(row_idx)
        row_cells[1].text = data.get("remark", "")
        row_cells[2].text = data.get("status", "")
        row_cells[3].text = data.get("comment", "")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
