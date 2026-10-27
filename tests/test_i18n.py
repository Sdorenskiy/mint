"""Английский интерфейс и поиск по английским документам; полнота перевода.

Запуск из корня проекта:  python tests/test_i18n.py
"""
import ast
import json
import os
import re
import shutil
import string
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import harness  # noqa: E402
import mock_server  # noqa: E402
import docx as pydocx  # noqa: E402

check = harness.Checker()


def fields(s):
    return sorted((f, spec, conv) for _, f, spec, conv in string.Formatter().parse(s) if f is not None)


# ---------------------------------------------------------------- 1. Полнота перевода (статически)
print("1. [unit] Полнота перевода: tr() и N_() в коде ↔ mint/i18n/en.py")
sys.path.insert(0, str(harness.ROOT))
from mint.i18n.en import EN  # noqa: E402

used = {}
for f in sorted((harness.ROOT / "mint").rglob("*.py")) + [harness.ROOT / "check_install.py", harness.ROOT / "eval_search.py"]:
    for node in ast.walk(ast.parse(f.read_text(encoding="utf-8"))):
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id in ("tr", "N_") \
                and node.args and isinstance(node.args[0], ast.Constant) and isinstance(node.args[0].value, str):
            used.setdefault(node.args[0].value, f"{f.relative_to(harness.ROOT)}:{node.lineno}")
missing = {k: v for k, v in used.items() if k not in EN}
check(not missing, f"у всех {len(used)} строк есть перевод", "\n".join(f"{v}: {k[:80]!r}" for k, v in list(missing.items())[:15]))
bad_fields = [k for k in EN if k in used and fields(k) != fields(EN[k])]
check(not bad_fields, "плейсхолдеры {…} в переводе совпадают с оригиналом", str(bad_fields[:5]))
cyr = [k for k, v in EN.items() if re.search(r"[а-яА-ЯёЁ]", v) and "ЗПС" not in v]
check(not cyr, "в английских строках нет кириллицы", str(cyr[:5]))
stale = [k for k in EN if k not in used]
check(not stale, "в en.py нет устаревших строк (их больше нет в коде)", "\n".join(repr(k[:80]) for k in stale[:15]))

# ---------------------------------------------------------------- 2. Механика tr и выбор языка
print("2. [unit] tr(), выбор языка")
W = harness.make_project(prefix="mint_i18n_")
os.environ.update({"MINT_LANG": "ru"})
app = harness.load_mint(W)
i18n = sys.modules["mint.i18n"]
check(app.tr("Файл: {filename}", filename="a.docx") == "Файл: a.docx", "русский — исходный текст")
with i18n.use_lang("en"):
    check(app.tr("Файл: {filename}", filename="a.docx") == "File: a.docx", "английский — перевод с подстановкой")
    check(app.tr("Строки нет в словаре") == "Строки нет в словаре", "нет перевода — русский текст")
check(app.tr("Файл: {filename}", lang="en", filename="x") == "File: x", "язык можно передать явно")
check([i18n.normalize_lang(v) for v in ("en", "EN", "en_US.UTF-8", "English", "ru-RU", "de", "")] ==
      ["en", "en", "en", "en", "ru", None, None], "normalize_lang")
saved = {k: os.environ.get(k) for k in ("MINT_LANG", "LANG", "LC_ALL", "LC_MESSAGES", "LANGUAGE")}
for k in saved:
    os.environ.pop(k, None)
os.environ["LANG"] = "ru_RU.UTF-8"
check(i18n.default_lang() == "ru", "русская локаль системы → ru")
os.environ["LANG"] = "en_GB.UTF-8"
check(i18n.default_lang() == "en", "другая локаль → en")
os.environ["MINT_LANG"] = "ru"
check(i18n.default_lang() == "ru", "MINT_LANG важнее локали")
for k, v in saved.items():
    if v is None:
        os.environ.pop(k, None)
    else:
        os.environ[k] = v

# ---------------------------------------------------------------- 3. Промпты и разбор английских ответов
print("3. [unit] Английские промпты и разбор ответов модели")
with i18n.use_lang("en"):
    check("Answer in the language of the question" in app.chat_system_prompt(), "промпт чата на английском")
    check("Senior Controller" in app.critique_system() and "KEEP|DROP" in app.critique_system(),
          "промпт самокритики на английском")
    check("[CRITICAL]" in app.audit_map_prompt("", "focus") and "Section is clean" in app.audit_map_prompt("", "f"),
          "промпт проверки раздела на английском")
    check('"FIXED|PARTIAL|NOT_FIXED"' in app.review_prompt("WIKI") and "WIKI" in app.review_prompt("WIKI"),
          "промпт контроля исправлений на английском")
    check("--- QUESTION ---" in app.chat_user_content("ctx", "q"), "разметка вопроса на английском")
check("Ты — ведущий ИИ-архитектор" in app.chat_system_prompt(), "русский промпт без изменений (профиль по умолчанию)")
f1 = app.finding_from_line("- **[CRITICAL]** Encryption is not specified")
check(f1["severity"] == "КРИТИЧНЫЙ" and f1["text"] == "Encryption is not specified", f"английская критичность: {f1}")
check(app.finding_from_line("- [HIGH] x" * 1)["severity"] == "КРИТИЧНЫЙ", "HIGH → критичный")
check(app.is_clean_report("Section is clean.") and app.is_clean_report("Часть чистая"), "«чистая часть» на обоих языках")
check(app.finding_clause("Clause 4.3: personal data are not described") == "4.3"
      and app.finding_clause("2.2.3, step 4: no transfer") == "2.2.3", "пункт в начале английского замечания")
src = "The data is encrypted with AES-256 between subsystems."
st_, _c = app.parse_verification(json.dumps({"status": "FIXED", "quote": "encrypted with AES-256", "comment": "ok"}), src)
st2, _c2 = app.parse_verification(json.dumps({"status": "NOT_FIXED", "quote": "", "comment": "no"}), src)
st3, _c3 = app.parse_verification(json.dumps({"status": "PARTIAL", "quote": "nowhere", "comment": "?"}), src)
check(st_.startswith("✅") and st2.startswith("❌") and st3.startswith("❓"), f"английские статусы: {st_}, {st2}, {st3}")
with i18n.use_lang("en"):
    st4, _ = app.parse_verification(json.dumps({"status": "FIXED", "quote": "encrypted with AES-256", "comment": "ok"}), src)
    check(st4 == "✅ Fixed", f"статус на языке интерфейса: {st4}")

# ---------------------------------------------------------------- 4. Поиск по-английски
print("4. [unit] Английский стеммер, стоп-слова, разбор вопроса, сокращения")
s = app.english_stem
check(len({s(w) for w in ("requirement", "requirements", "required", "requiring")}) == 1, "формы requirement — одна основа")
check(len({s(w) for w in ("encryption", "encrypted", "encrypting")}) == 1, "формы encrypt — одна основа")
check(s("caresses") == "caress" and s("ponies") == "poni" and s("hopping") == "hop" and s("relational") == "relat",
      "эталонные примеры Porter2")
check(app.bm25_tokenize("The logs are stored") == ["the", "log", "are", "store"], "токенизация английского текста")
check(app.normalize_token("документов") == app.russian_stem("документов"), "русские слова — прежний стеммер")
check({"what", "which", "the"} <= set(app.BM25_STOP_WORDS), "английские стоп-слова")
plan = app.analyze_query("What is SLA?")
check(plan["kind"] == "term" and plan["target"] == "sla", f"«what is …» — вопрос об определении: {plan}")
plan2 = app.analyze_query("components of the reporting portal")
check(plan2["kind"] == "composition" and "reporting portal" in plan2["target"], f"«components of …» — состав: {plan2}")
ab = dict((a, f) for a, f, _w in app.extract_abbreviations(
    "The Service Level Agreement (SLA) is signed by the Ministry of Finance (MOF).\n- HSM — Hardware Security Module"))
check(ab.get("SLA") == "Service Level Agreement" and ab.get("MOF") == "Ministry of Finance"
      and ab.get("HSM") == "Hardware Security Module", f"английские сокращения: {ab}")

# ---------------------------------------------------------------- 5. Интерфейс на английском
print("5. [сквозной] Интерфейс на английском: чат, аудит, контроль исправлений, отзывы")
srv = mock_server.start()
BASE = f"http://127.0.0.1:{srv.server_address[1]}/v1"
os.environ.update({
    "LLM_API_URL": BASE, "EMBED_API_URL": BASE + "/embeddings", "RERANK_API_URL": BASE + "/rerank",
    "LLM_MODEL": "test-llm", "EMBED_MODEL": "emb-model", "RERANK_MODEL": "rr-model", "LLM_API_KEY": "k",
    "KB_DIR": str(W / "kb"), "LOG_DIR": str(W / "logs"), "MINT_CONFIG": str(W / "none.yaml"),
    "GLOSSARY_PATH": str(W / "none.json"), "MINT_LANG": "ru",
})
os.chdir(W)
import streamlit as st  # noqa: E402  (заглушка)


def make_docx(path, blocks):
    d = pydocx.Document()
    for kind, val in blocks:
        if kind == "h1":
            d.add_heading(val, level=1)
        elif kind == "h2":
            d.add_heading(val, level=2)
        else:
            d.add_paragraph(val)
    d.save(path)
    return path.read_bytes()


class Upl:
    def __init__(self, name, data):
        self.name, self._d = name, data

    def getbuffer(self):
        return memoryview(self._d)


def run_ui(chat=None, buttons=(), uploads=None):
    st.LOG.clear()
    st.SCRIPT["chat_input"] = chat
    st.SCRIPT["buttons"] = set(buttons)
    st.SCRIPT["file_uploader"] = uploads or {}
    try:
        app.main()
        return "ok"
    except st.RerunException:
        return "rerun"


def log_text():
    return " ".join(str(x[1]) for x in st.LOG)


kb_doc = make_docx(W / "policy.docx", [
    ("h1", "Abbreviations"), ("p", "PKI – Public Key Infrastructure;"),
    ("h1", "1. Information security"),
    ("h2", "1.1 Encryption"),
    ("p", "Personal data transferred between subsystems must be encrypted with certified algorithms."),
    ("h2", "1.2 Logging"),
    ("p", "Security event logs are retained for at least three years and access to them is role-based. " * 6),
    ("h1", "2. Availability"),
    ("p", "The service recovery time must not exceed four hours; availability is 99.9 percent per month."),
])
spec_doc = make_docx(W / "spec.docx", [
    ("h1", "1. General"), ("p", "The system transfers personal data between subsystems over the internal network."),
    ("h2", "1.1 Performance"), ("p", "The system must respond quickly under load."),
])
spec_v2 = make_docx(W / "spec_v2.docx", [
    ("h1", "1. General"),
    ("p", "The system transfers personal data between subsystems; the channel is encrypted with TLS 1.3."),
])

try:
    check(run_ui() == "ok" and "Интеллектуальная RAG-система: ответы по базе знаний" in log_text(), "по умолчанию (MINT_LANG=ru) — русский")
    st.session_state["ui_lang"] = "en"
    check(run_ui() == "ok", "переключение на английский")
    txt = log_text()
    check("An intelligent RAG system: answers from your knowledge base" in txt and "⚙️ Settings" in txt, f"английский заголовок: {txt[:200]}")
    check(not re.search(r"[а-яА-ЯёЁ]{4,}", " ".join(str(x[1]) for x in st.LOG if x[0] in
                                                    ("title", "header", "subheader", "button", "caption", "slider",
                                                     "number_input", "checkbox", "selectbox", "text_area",
                                                     "file_uploader", "expander")
                                                    and "Язык / Language" not in str(x[1]))),
          "в элементах интерфейса нет русского текста",
          str([x for x in st.LOG if re.search(r"[а-яА-ЯёЁ]{4,}", str(x[1])) and "Язык / Language" not in str(x[1])][:5]))

    check(run_ui(uploads={"kb_uploader": [Upl("policy.docx", kb_doc)]}, buttons={"kb_index_all_btn"}) == "rerun",
          "загрузка английского документа")
    res = st.session_state.get("_kb_batch_results") or []
    check(res and res[0][1] and res[0][2].startswith("Added"), f"сообщение о загрузке на английском: {res}")
    kb_lang = app.kb_language(app.load_settings()[0])
    check(kb_lang[0] == "en", f"язык базы — английский: {kb_lang}")
    check(run_ui(buttons={"detect_limits_btn"}) == "rerun"
          and "measured on English text" in st.session_state.get("_limits_msg", "")
          and "Server: " in st.session_state["_limits_msg"],
          f"«символов на токен» измерено на английском: {st.session_state.get('_limits_msg', '')[:300]}")

    check(run_ui(chat="How long are security logs retained?") == "ok", "вопрос в чате")
    last = st.session_state["messages"][-1]
    check(last["content"].startswith("Answer from the context") and last["sources"], f"английский ответ: {last['content']}")
    sent = mock_server.STATE["last_chat"]["messages"]
    check("Answer in the language of the question" in sent[0]["content"] and "=== SOURCE:" in sent[-1]["content"],
          "модель получила английский промпт и разметку источников")
    check("retained for at least three years" in sent[-1]["content"] and "policy.docx" in last["sources"][0],
          f"в контексте — фрагмент с ответом: {last['sources'][:2]}")
    check(any(x == ("button", "👍") for x in st.LOG), "кнопки отзыва")

    gl = app.load_glossary(app.load_settings()[0])
    check(gl.entries.get("PKI") == "Public Key Infrastructure", f"сокращение из английского документа: {gl.entries}")

    check(run_ui(uploads={"tab2_audit_uploader": Upl("spec.docx", spec_doc)}, buttons={"tab2_run_btn"}) == "ok",
          "аудит английского документа")
    report = st.session_state.get("tab2_final_report_data", "")
    check("EXPERT REVIEW REPORT" in report and "THE DOCUMENT NEEDS REWORK" in report and "[CRITICAL]" in report,
          f"заключение на английском: {report[:300]}")
    check("FINDINGS BY DOCUMENT SECTION" in report and "Information security" not in report,
          "замечания по разделам документа, без тем")
    check("# Review log: spec.docx" in st.session_state.get("tab2_audit_log_md", "")
          and "MINT version" in st.session_state["tab2_audit_log_md"], "лог проверки на английском")
    docx_bytes, stats = st.session_state.get("tab2_docx_review") or (None, {})
    if docx_bytes and stats.get("mode") == "comments":
        import zipfile, io
        comments = zipfile.ZipFile(io.BytesIO(docx_bytes)).read("word/comments.xml").decode("utf-8")
        check("[CRITICAL]" in comments, "примечания Word с английской критичностью")

    st.session_state["tab3_remarks_area"] = "- [CRITICAL] The channel encryption for personal data is not specified"
    check(run_ui(uploads={"tab3_v2_uploader": Upl("spec_v2.docx", spec_v2)}, buttons={"tab3_run_btn"}) == "ok",
          "контроль исправлений по-английски")
    rows = (st.session_state.get("tab3_result") or {}).get("rows") or []
    check(rows and rows[0]["status"] == "✅ Fixed" and "Quote from the new revision" in rows[0]["comment"],
          f"статус и комментарий на английском: {rows}")
    check("# REVIEW COMMENT RESOLUTION MATRIX" in st.session_state["tab3_result"]["matrix_md"], "матрица на английском")

    # проверка качества по-английски: вопросы на языке фрагмента, подписи и отчёт — английские
    st.session_state["qa_n"] = 5
    check(run_ui(buttons={"qa_open_btn"}) == "ok" and ("button", "✨ Write questions") in st.LOG,
          "проверка качества открывается по-английски")
    check(run_ui(buttons={"qa_gen_btn"}) == "rerun", "вопросы составлены")
    qrows = st.session_state["qa"]["rows"]
    check(qrows and qrows[0]["question"].startswith("What does it say about"), f"английский промпт генерации: {qrows[:1]}")
    check(run_ui() == "ok" and any(x[0] == "expander" and "Fragment the question was written from" in str(x[1])
                                   for x in st.LOG), "шаг «Вопросы» по-английски")
    check(run_ui(buttons={"qa_run_btn"}) == "rerun" and st.session_state["qa"]["stage"] == "results", "прогон")
    check(run_ui() == "ok", "результаты по-английски")
    qa_txt = " ".join(str(x[1]) for x in st.LOG if x[0] in ("markdown", "button", "caption", "radio", "download_button",
                                                             "subheader", "expander"))
    check("in the model context" in qa_txt and "What to change" in qa_txt, "итоги и рекомендации по-английски")
    check(not re.search(r"[а-яА-ЯёЁ]{4,}", qa_txt), "в проверке качества нет русского текста",
          str(re.findall(r"[^.]{0,40}[а-яА-ЯёЁ]{4,}[^.]{0,40}", qa_txt)[:5]))
    rep = st.session_state["qa"].get("report", "")
    check(rep.startswith("# MINT search evaluation") and "Configurations compared" in rep, f"отчёт по-английски: {rep[:120]}")
    check(run_ui(buttons={"qa_close_btn"}) == "ok", "закрыть проверку")

    # страница промптов по-английски: подписи, названия промптов, тексты и разметка — английские
    check(run_ui(buttons={"pr_open_btn"}) == "ok" and ("subheader", "📝 Prompts") in st.LOG
          and ("button", "Check the edit") in st.LOG, "страница промптов открывается по-английски")
    pr_txt = " ".join(str(x[1]) for x in st.LOG if x[0] in ("markdown", "button", "caption", "subheader", "expander",
                                                             "selectbox", "selectbox_options", "text_area", "code",
                                                             "warning", "error", "info")
                      and "Язык / Language" not in str(x[1]) and not str(x[1]).startswith("Русский"))  # выбор языка — на двух
    check("Chat system prompt" in pr_txt and "Document review · Section review" in pr_txt
          and "=== SOURCE:" in pr_txt and st.session_state.get("pr_text_en_chat_system_0", "").startswith("You are"),
          "названия промптов и тексты — английские")
    check(not re.search(r"[а-яА-ЯёЁ]{4,}", pr_txt), "на странице промптов нет русского текста",
          str(re.findall(r"[^.]{0,40}[а-яА-ЯёЁ]{4,}[^.]{0,40}", pr_txt)[:5]))
    check(run_ui(buttons={"pr_close_btn"}) == "ok", "закрыть промпты")

    # черновик карточки системы по-английски: страница и сам черновик
    st.session_state["cards_system"] = "System A"
    check(run_ui(buttons={"cards_open_btn"}) == "ok" and run_ui(buttons={"cards_run_btn"}) == "ok"
          and ("subheader", "🗂️ System card draft") in st.LOG, "страница черновика карточки — по-английски")
    card_md = (st.session_state.get("cards_result") or {}).get("markdown", "")
    cd_txt = " ".join(str(x[1]) for x in st.LOG if x[0] in ("button", "caption", "subheader", "text_input",
                                                             "download_button", "warning", "error")
                      and "Язык / Language" not in str(x[1]))
    check(card_md.startswith('# System card "System A"') and "## What it can do now" in card_md
          and app.card_names(card_md, "Card_System_A.md")[0] == "System A"
          and not re.search(r"[а-яА-ЯёЁ]{4,}", cd_txt),
          "черновик и подписи — английские; название узнаётся так же, как у карточки",
          str(re.findall(r"[^.]{0,40}[а-яА-ЯёЁ]{4,}[^.]{0,40}", cd_txt)[:5]))
    check(run_ui(buttons={"cards_close_btn"}) == "ok", "закрыть черновик карточки")

    st.session_state["ui_lang"] = "ru"
    check(run_ui() == "ok" and "Интеллектуальная RAG-система: ответы по базе знаний" in log_text(), "обратно на русский")
    check(st.session_state["tab2_user_focus"] == app.DEFAULT_AUDIT_FOCUS, "критерии аудита по умолчанию сменили язык")
except Exception:
    import traceback
    traceback.print_exc()
    check.fails.append("исключение в английском интерфейсе")

srv.shutdown()
os.chdir(harness.ROOT)
shutil.rmtree(W, ignore_errors=True)
check.finish()
