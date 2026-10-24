"""Проверка ядра и интерфейса MINT на заглушках (Streamlit, Docling) и имитации API-шлюза.

Разделы помечены типом: [unit] — одна функция отдельно; [интеграция] — несколько модулей вместе
(.docx → база → поиск → модель); [сквозной] — действия пользователя на странице.

Запуск из корня проекта:  python tests/test_app.py
"""
import io
import json
import os
import shutil
import sys
import traceback
import zipfile
from dataclasses import replace
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import harness  # noqa: E402  (заглушки и пути — до импорта mint)
import mock_server  # noqa: E402
import docx as pydocx  # noqa: E402

WORK = harness.make_project()
srv = mock_server.start()
BASE = f"http://127.0.0.1:{srv.server_address[1]}/v1"

os.environ.update({
    "LLM_API_URL": BASE, "EMBED_API_URL": BASE + "/embeddings", "RERANK_API_URL": BASE + "/rerank",
    "LLM_MODEL": "test-llm", "EMBED_MODEL": "emb-model", "RERANK_MODEL": "rr-model",
    "KB_DIR": str(WORK / "kb"), "LOG_DIR": str(WORK / "logs"),
    "MINT_CONFIG": str(WORK / "config.yaml"), "GLOSSARY_PATH": str(WORK / "glossary.json"),
    "LLM_API_KEY": "test-key",
})
os.chdir(WORK)
(WORK / "config.yaml").write_text(
    "llm_url: http://bad\n"
    "llm:\n  max_tokens: 1200\n"
    "search:\n  top_k: 0\n  unknown_param: 5\n"
    "models:\n  chars_per_token: 0\n", encoding="utf-8")
(WORK / "glossary.json").write_text(json.dumps({
    "rerank_abbreviations": {"ГИС": "государственная информационная система"},
    "bm25_expansions": {"сеуд": ["система", "единого", "учёта", "документов"]},
    "bm25_phrase_boost": {"гис": ["государ", "информац"]},
    "query_suffixes": ["в системе"],
}, ensure_ascii=False), encoding="utf-8")

app = harness.load_mint(WORK)
import streamlit as st  # noqa: E402  (заглушка)

check = harness.Checker()
FAILS = check.fails

def make_docx(path, blocks):
    d = pydocx.Document()
    for kind, val in blocks:
        if kind == "h1":
            d.add_heading(val, level=1)
        elif kind == "h2":
            d.add_heading(val, level=2)
        elif kind == "p":
            d.add_paragraph(val)
        elif kind == "table":
            t = d.add_table(rows=len(val), cols=len(val[0]))
            for r, row in enumerate(val):
                for c, cell in enumerate(row):
                    t.cell(r, c).text = cell
    d.save(path)
    return path.read_bytes()


print("1. [unit] Стеммер")
groups = [["документ", "документа", "документов"], ["требование", "требования", "требованиями"],
          ["средства", "средствам", "средствах"], ["шифрование", "шифрования"]]
for g in groups:
    stems = {app.russian_stem(w) for w in g}
    check(len(stems) == 1, f"{g} → {stems}")

print("2. [unit] Извлечение сокращений")
cases = {
    "Требования к средствам защиты информации (СЗИ) приведены ниже.": ("СЗИ", "средствам защиты информации"),
    "СЗИ – средства защиты информации;": ("СЗИ", "средства защиты информации"),
    "| НСД | Несанкционированный доступ |": ("НСД", "несанкционированный доступ"),
    "Сертификация по государственным стандартам (ГОСТ) обязательна.": ("ГОСТ", "государственным стандартам"),
    "ФСТЭК — Федеральная служба по техническому и экспортному контролю": (
        "ФСТЭК", "федеральная служба по техническому и экспортному контролю"),
    "Научно-исследовательские и опытно-конструкторские работы (НИОКР) ведутся.": (
        "НИОКР", "научно-исследовательские и опытно-конструкторские работы"),
}
for text, (abbr, full) in cases.items():
    got = [(a, f) for a, f, _ in app.extract_abbreviations(text)]
    check((abbr, full) in got, f"{text[:50]}… → {got}")
negatives = ["Требования приведены в ТЗ (см. раздел 3).", "Система хранения (СУБД) развёрнута.",
             "Подробнее в разделе ВНИМАНИЕ (ОБЯЗАТЕЛЬНО)."]
for text in negatives:
    got = app.extract_abbreviations(text)
    check(not got, f"нет ложных срабатываний: {text} → {got}")

print("3. [unit] Настройки")
cfg0, notes = app.load_settings()
check(cfg0.llm_url == BASE, "адрес LLM взят из окружения, а не из config.yaml")
check(any("llm_url" in n for n in notes), "предупреждение: подключения в config.yaml не используются")
check(any("unknown_param" in n for n in notes), "предупреждение: неизвестный параметр")
check(cfg0.max_tokens == 1200 and "max_tokens" in cfg0.explicit_fields, "max_tokens из config.yaml")
check(cfg0.embed_key == "test-key" and cfg0.rerank_enabled, "ключи и реранкер")

print("4. [интеграция] Автоопределение лимитов у сервера")
server = app.detect_server_info(cfg0)
check("test-llm" in server["llm_models"], f"модели: {server['llm_models']}")
check(server["llm_context"].get("test-llm") == 8192, "окно test-llm = 8192")
check(server["embed_max_tokens"] == 512 and server["rerank_max_tokens"] == 1024, "лимиты эмбеддингов и реранкера")
cpt, cpt_src = app.measure_chars_per_token(cfg0, "test-llm")
check(cpt is not None and 2.8 <= cpt <= 3.2 and cpt_src == "LLM", f"символов на токен: {cpt} ({cpt_src})")
_orig_measure = app._measure_cpt
app._measure_cpt = lambda c, m, sample: ((2.6, "LLM") if sample == app.CPT_SAMPLE else (3.9, "LLM"))
by_lang = {lang: app.measure_chars_per_token(cfg0, "test-llm", lang)[0] for lang in ("ru", "en", "mixed", None)}
app._measure_cpt = _orig_measure
check(by_lang == {"ru": 2.6, "en": 3.9, "mixed": 2.6, None: 2.6},
      f"символов на токен по языку базы: русская — русский текст, английская — английский, иначе меньшее: {by_lang}")
check(app.kb_language(cfg0) == (None, 0.0), "язык пустой базы не определён")
_orig_snap = app.load_kb_snapshot
app.load_kb_snapshot = lambda c: app.KBSnapshot(
    {"a": {"text": "Журналы событий хранятся не менее трёх лет."}, "b": {"text": "Event logs are kept for three years."}},
    [], None)
mixed = app.kb_language(cfg0)
app.load_kb_snapshot = _orig_snap
check(mixed[0] == "mixed" and 0.2 < mixed[1] < 0.8, f"смешанная база: {mixed}")

print("5. [unit] Расчёт зависимых параметров")
limits = app.resolve_model_limits(cfg0, server, cpt, {})
check(limits["llm_context_tokens"] == (8192, "сервер"), f"окно: {limits['llm_context_tokens']}")
cfg, info = app.resolve_settings(cfg0, limits, {"top_k": 4})
print("   ", {f: i["value"] for f, i in info.items()})
check(cfg.child_size == 400, "child = 400 (вход 512 токенов позволяет)")
check(cfg.section_embed_chars < 4000, f"кусок раздела для эмбеддинга урезан под 512 токенов: {cfg.section_embed_chars}")
check(cfg.review_context_max_chars < 12000, f"контекст вкладки 3 урезан под окно 8k: {cfg.review_context_max_chars}")
check(cfg.rerank_query_max_chars < 1500, f"запрос реранкера урезан под 1024 токена: {cfg.rerank_query_max_chars}")
check(cfg.llm_context_known, "окно известно (от сервера)")
big_cfg, big_info = app.resolve_settings(
    cfg0, app.resolve_model_limits(cfg0, {}, None, {}), {"top_k": 5})
diff = {f: (i["value"], getattr(app.DEFAULTS, f)) for f, i in big_info.items() if i["value"] != getattr(app.DEFAULTS, f)}
check(not diff, f"при типовых лимитах (32k / 8k / 8k) все значения совпадают с v17 {diff}")
check(not big_cfg.llm_context_known, "типовое окно не считается известным")
small_cfg, small_info = app.resolve_settings(cfg0, app.resolve_model_limits(
    cfg0, server, 3.0, {"llm_context_tokens": 4096}), {"top_k": 6})
warns = app.limits_warnings(small_info)
check(small_cfg.smart_truncate_chars < 2000, f"окно 4k: фрагмент чата {small_cfg.smart_truncate_chars}")
print("    предупреждения при 4k:", warns[:2])
man_cfg, man_info = app.resolve_settings(cfg0, limits, {"top_k": 4, "parent_size": 3000})
check(man_info["parent_size"]["source"] == "вручную" and man_cfg.parent_size == 3000, "ручное переопределение")

print("6. [unit] Проверка «помещается ли запрос»")
tiny = replace(cfg, llm_context_tokens=1000, llm_context_known=True, chars_per_token=3.0)
r = app.ask_llm("system", "x" * 5000, tiny)
check(r.startswith("❌ [КОНТЕКСТ НЕ ПОМЕЩАЕТСЯ]"), r[:60])
r = app.ask_llm("system", "вопрос", replace(cfg, llm_context_known=False))
check(not r.startswith("❌"), "обычный запрос проходит")

print("7. [интеграция] Индексация документов")
docs_dir = WORK / "docs"
docs_dir.mkdir()
long_par = ("Журналы событий безопасности хранятся не менее трёх лет. Доступ к журналам разграничивается по ролям. "
            "Администратор обязан ежедневно проверять целостность журналов. ") * 12
data_a = make_docx(docs_dir / "Регламент_ИБ.docx", [
    ("h1", "Перечень сокращений"),
    ("p", "СЗИ – средства защиты информации;"),
    ("p", "ПДн — персональные данные;"),
    ("p", "ЦОД - центр обработки данных."),
    ("h1", "1. Общие положения"),
    ("p", "Регламент определяет требования к системе управления базами данных (СУБД) и к "
          "автоматизированному рабочему месту (АРМ) оператора."),
    ("h2", "1.1 Шифрование"),
    ("p", "При передаче ПДн между подсистемами применяются средства криптографической защиты информации (СКЗИ), "
          "сертифицированные по ГОСТ. СЗИ устанавливаются на каждом АРМ."),
    ("h2", "1.2 Журналирование"),
    ("p", long_par),
    ("table", [["Сокращение", "Расшифровка"], ["НСД", "Несанкционированный доступ"], ["ВКС", "Видеоконференцсвязь"]]),
])
data_b = make_docx(docs_dir / "Регламент_SLA.docx", [
    ("h1", "Показатели надёжности"),
    ("p", "Время восстановления сервиса не превышает четырёх часов. Доступность системы — 99,9 % в месяц."),
    ("p", "Резервное копирование выполняется ежедневно в ЦОД."),
])
ok, msg = app.process_and_add_to_kb("Регламент_ИБ.docx", data_a, cfg)
check(ok, f"ИБ: {msg}")
ok, msg = app.process_and_add_to_kb("Регламент_SLA.docx", data_b, cfg)
check(ok, f"SLA: {msg}")
man = app.read_manifest(cfg)
check(set(man.get("documents", {})) == {"Регламент_ИБ.docx", "Регламент_SLA.docx"}, "манифест: оба документа")
check(man["documents"]["Регламент_ИБ.docx"]["embed_model"] == "emb-model", "в манифесте записана модель эмбеддингов")
check(app.source_path(cfg, "Регламент_ИБ.docx").exists(), "оригинал сохранён для переиндексации")
rows, nvec = app.kb_document_status(cfg)
check(nvec > 5 and all(not r["issues"] for r in rows), f"векторов {nvec}, расхождений нет")
kb_lang = app.kb_language(cfg)
check(kb_lang[0] == "ru", f"язык базы — русский (английские аббревиатуры не мешают): {kb_lang}")

print("8. [интеграция] Глоссарий")
gl = app.load_glossary(cfg)
print("   ", dict(sorted(gl.entries.items())))
for a in ("СЗИ", "ПДн", "ЦОД", "СУБД", "АРМ", "СКЗИ", "НСД", "ВКС"):
    check(a in gl.entries, f"автоглоссарий: {a} → {gl.entries.get(a)}")
check(gl.entries.get("СЗИ") == "средства защиты информации", "СЗИ — форма из перечня (именительный падеж)")
check(gl.entries.get("ГИС") == "государственная информационная система" and gl.origin["ГИС"] == "свой",
      "старый формат glossary.json: rerank_abbreviations")
check("СЕУД" in gl.entries, "старый формат glossary.json: bm25_expansions")
check(gl.find("поставили по ГОСТ") == [] or all(a != "ПО" for a, _ in gl.find("поставили по ГОСТ")),
      "предлог «по» не путается с сокращением")
terms = gl.bm25_terms("какие требования к сзи")
check(any(t.startswith("защит") for t, _ in terms), f"BM25-расширение: {terms}")

print("9. [интеграция] Поиск и чат")
ctx, srcs = app.search_parent_child("Какие требования к СЗИ на АРМ?", cfg, k_neighbors=4, glossary=gl,
                                    annotate_low_confidence=True)
check(srcs and "Регламент_ИБ.docx" in srcs[0], f"источник №1: {srcs[0] if srcs else None}")
check(any("dense #" in s for s in srcs) and any("rerank" in s for s in srcs), "dense и реранкер работают")
ref = gl.reference_block("Какие требования к СЗИ на АРМ?", ctx)
check("СЗИ — средства защиты информации" in ref, "справочник сокращений для модели")
answer = app.ask_llm(app.CHAT_SYSTEM_PROMPT, app.chat_user_content(ctx, "вопрос", ref), cfg)
check("Справочник: да" in answer, f"модель получила справочник: {answer}")
rr = [p for m, path, p in mock_server.STATE["requests"] if path.endswith("/rerank") and p and "АРМ" in p["query"]]
check(rr and "средства защиты информации" in rr[-1]["query"], "расшифровка добавлена в запрос реранкера")
check(all(len(d) <= cfg.rerank_doc_max_chars for d in rr[-1]["documents"]), "документы реранкера в пределах лимита")
eng = app.LightweightBM25(["Порядок хранения журналов событий", "Порядок резервного копирования данных"])
sc = eng.get_scores("Сколько хранятся журналы?")
check(sc[0] > 0 and sc[1] == 0, f"морфология BM25: «журналы» находит «журналов» ({sc})")
ctx2, srcs2 = app.search_parent_child("Сколько хранятся журналы?", cfg, k_neighbors=3, glossary=gl)
check("Журналы событий" in ctx2, "вопрос о журналах: нужный фрагмент в контексте")
ctx3, _ = app.search_parent_child("рецепт борща со свёклой", cfg, k_neighbors=3, glossary=gl,
                                  annotate_low_confidence=True)
check(ctx3.startswith("[СИГНАЛ ПОИСКА"), "нерелевантный вопрос → сигнал низкой уверенности")
check(app.is_low_confidence(0.05, 0.95, True, 0.3, 0.5) and not app.is_low_confidence(0.8, 0.0, False, 0.3, 0.5)
      and app.is_low_confidence(None, 0.2, False, 0.3, 0.5) and not app.is_low_confidence(None, 0.2, True, 0.3, 0.5)
      and not app.is_low_confidence(None, 0.9, False, 0.3, 0.5),
      "[unit] правило уверенности: с реранкером — его балл; без — доля слов и согласие dense и BM25")
frag = "# Перечень\n\nСЗИ — средства.\n\n# 1. Общие\n\nТекст.\n\n## 1.2 Журналирование\n\nЖурналы хранятся три года."
check(app.section_at(frag, "Журналы хранятся три года.") == "1.2 Журналирование"
      and app.section_at(frag, "## 1.2 Журналирование\n\nЖурналы") == "1.2 Журналирование"
      and app.section_at(frag, "СЗИ — средства.") == "Перечень" and app.section_at(frag, "нет такого") is None,
      "[unit] подпись источника — раздел, где найдено место, а не раздел в начале фрагмента")
# вопрос, слова которого в базе есть: решает балл реранкера, а не совпадение слов
q_words = "Сколько хранятся журналы событий безопасности?"
_orig_rr = app.rerank_scores
diag_lo, diag_hi, diag_nr, diag_nr2 = {}, {}, {}, {}
app.rerank_scores = lambda q, docs, c: [0.05] * len(docs)
app.search_parent_child(q_words, cfg, k_neighbors=3, glossary=gl, diagnostics=diag_lo)
app.rerank_scores = lambda q, docs, c: [0.9] * len(docs)
app.search_parent_child(q_words, cfg, k_neighbors=3, glossary=gl, diagnostics=diag_hi)
app.rerank_scores = _orig_rr
check(diag_lo.get("low_conf") and diag_lo["lexical_cov"] >= cfg.low_conf_lexical,
      f"реранкер считает фрагменты нерелевантными → «уверенность низкая», хотя слова вопроса в контексте есть "
      f"(прежнее правило сигнал бы погасило): {diag_lo}")
check(diag_hi.get("low_conf") is False, f"высокий балл реранкера → уверенность нормальная: {diag_hi}")
cfg_nr = replace(cfg, rerank_url="")
app.search_parent_child("рецепт борща со свёклой", cfg_nr, k_neighbors=3, glossary=gl, diagnostics=diag_nr)
app.search_parent_child(q_words, cfg_nr, k_neighbors=3, glossary=gl, diagnostics=diag_nr2)
check(diag_nr.get("low_conf") and diag_nr["best_rerank"] is None and diag_nr2.get("low_conf") is False,
      f"без реранкера — по доле слов и согласию dense и BM25: {diag_nr}, {diag_nr2}")
emb = [p for m, path, p in mock_server.STATE["requests"] if path.endswith("/embeddings") and p]
check(all(len(p["input"]) <= cfg.embed_input_max_chars for p in emb if isinstance(p["input"], str)),
      "вход эмбеддингов не превышает лимит модели")

print("10. [интеграция] Смена модели эмбеддингов")
cfg2 = replace(cfg, embed_model="emb-model-2")
rows2, _ = app.kb_document_status(cfg2)
check(all(any(l == "error" for l, _ in r["issues"]) for r in rows2), "все документы помечены «другая модель»")
ctx4, srcs4 = app.search_parent_child("требования к СЗИ", cfg2, k_neighbors=3, glossary=gl)
check(srcs4 and all("dense —" in s for s in srcs4), "размерность другая → поиск продолжает работать по BM25")
ok, msg = app.process_and_add_to_kb("Регламент_SLA.docx", data_b, cfg2)
check(not ok and "Размерность" in msg, f"добавление в базу другой размерности отклонено: {msg[:70]}")
res = app.reindex_documents(cfg2)
check(res and all(ok for _, ok, _ in res), f"переиндексация: {[(n, ok) for n, ok, _ in res]}")
rows3, nvec3 = app.kb_document_status(cfg2)
check(nvec3 > 0 and all(not r["issues"] for r in rows3), "после переиндексации расхождений нет")
_, srcs5 = app.search_parent_child("требования к СЗИ", cfg2, k_neighbors=3, glossary=app.load_glossary(cfg2))
check(any("dense #" in s for s in srcs5), "векторный поиск снова работает")

print("11. [интеграция] Аудит документа")
tz_bytes = make_docx(docs_dir / "ТЗ.docx", [
    ("h1", "Требования к системе"),
    ("p", "Система передаёт ПДн между подсистемами. Время отклика должно быть приемлемым."),
])
md = app.compile_docx_to_markdown("ТЗ.docx", tz_bytes)
sections = app.split_into_sections(md, cfg.audit_map_size)
res = app.audit_section(sections[0], cfg2, app.DEFAULT_AUDIT_FOCUS, True, 3, glossary=app.load_glossary(cfg2))
check(not res["failed"] and len(res["findings"]) == 2 and len(res["dropped"]) == 1,
      f"черновик 3 замечания → самокритика оставила 2: {res['report']!r}")
f1, f3 = res["findings"]
check(f1["quote"] and app.quote_in_text(f1["quote"], sections[0]) and "category" not in f1,
      f"проверенная цитата от самокритики, без темы: {f1['quote']!r}")
check(f3["text"].startswith("152-ФЗ") and f3["severity"] == "КРИТИЧНЫЙ", f"цифры в начале сохранены: {f3['text']!r}")
check(res["sources"], "нормы подобраны из базы")
by_section = app.group_by_section(app.collect_unique_findings(res["findings"]), {})
report = app.build_final_report("ТЗ.docx", by_section, [], len(sections))
check("ТРЕБУЕТ ДОРАБОТКИ" in report and "152-ФЗ" in report and "АРХИТЕКТУРН" not in report
      and "## 📋 ЗАМЕЧАНИЯ ПО РАЗДЕЛАМ ДОКУМЕНТА" in report, "итоговое заключение: замечания по разделам, без тем")
annotated, ann_stats = app.annotate_docx_with_findings(tz_bytes, res["findings"])
check(len(annotated) > 1000 and ann_stats["by_quote"] == 1, f"ТЗ с замечаниями (.docx): {ann_stats}")

print("12. [интеграция] Контроль исправлений")
v2 = make_docx(docs_dir / "ТЗ_v2.docx", [
    ("h1", "Требования к системе"),
    ("p", "При передаче ПДн применяется шифрование по ГОСТ 34.12-2018 на всех каналах."),
    ("h1", "Производительность"),
    ("p", "Время отклика не превышает 2 секунд для 95 % запросов."),
])
sections_v2 = app.split_into_sections(app.compile_docx_to_markdown("ТЗ_v2.docx", v2), 200)
idx = app.build_section_index(sections_v2, cfg2)
status_lbl, comment = app.verify_remark("- [КРИТИЧНЫЙ] Не указан алгоритм шифрования ПДн", sections_v2, idx, cfg2,
                                         True, 3, glossary=app.load_glossary(cfg2))
check(status_lbl.startswith("✅"), f"{status_lbl}: {comment[:80]}")
matrix = [{"remark": "x", "status": status_lbl, "comment": comment}]
check(app.generate_word_matrix(matrix) and "МАТРИЦА" in app.build_matrix_markdown(matrix), "матрица (.md и .docx)")

print("13. [интеграция] Удаление документа")
check(app.delete_document_from_kb("Регламент_SLA.docx", cfg2), "удалён")
docs_left, _ = app.list_kb_documents(cfg2)
check(docs_left == ["Регламент_ИБ.docx"], f"осталось: {docs_left}")
check(not app.source_path(cfg2, "Регламент_SLA.docx").exists(), "оригинал удалён")
check("Регламент_SLA.docx" not in app.read_manifest(cfg2)["documents"], "манифест обновлён")


print("14. [unit + интеграция] Расширение запроса, реранкер, нормировка, префиксы, аудит, Word, блокировка")
# --- анализ запроса ---
check(app.analyze_query("срок составления акта")["target"] is None, "«срок составления акта» — не вопрос о составе")
pl = app.analyze_query("Что такое ФИС?")
check(pl["kind"] == "term" and pl["target"] == "фис", f"«что такое ФИС» → {pl}")
check(any(w == 0.3 for _, w in pl["terms"]) and all(w <= 0.5 for _, w in pl["terms"]), "слова расширения с малым весом")
pl2 = app.analyze_query("какая архитектура у системы хранения")
check(pl2["target"] == "системы хранения", f"предмет без предлога: {pl2['target']!r}")
check(app.analyze_query("архитектура " * 60)["target"] is None, "длинный запрос не разбирается")
# --- BM25: фраза и близость ---
eng = app.LightweightBM25(["Подсистема хранения резервных копий и система учёта",
                           "Система хранения данных состоит из трёх узлов"])
sc = eng.get_scores("какая архитектура у системы хранения", phrase=pl2["target"])
check(sc[1] > sc[0] + 9, f"буст точной фразы по предмету вопроса: {sc}")
eng2 = app.LightweightBM25(["Время восстановления сервиса не превышает четырёх часов",
                            "Время работы оператора ограничено, а для восстановления данных и сервиса нужен регламент"])
sc2 = eng2.get_scores("время восстановления сервиса")
check(sc2[0] > sc2[1] + 3, f"близость соседних слов вопроса: {sc2}")
# --- dense и multi-query получают исходный вопрос ---
mock_server.STATE["requests"].clear()
app.search_parent_child("Что такое СЗИ?", cfg2, k_neighbors=3, glossary=app.load_glossary(cfg2), use_multi_query=True)
emb_q = [pl_["input"] for m, path, pl_ in mock_server.STATE["requests"]
         if path.endswith("/embeddings") and pl_ and isinstance(pl_["input"], str)]
check(emb_q and emb_q[0] == cfg2.embed_query_prefix + "Что такое СЗИ?", f"вектор запроса — по исходному вопросу: {emb_q[:1]}")
para = [pl_ for m, path, pl_ in mock_server.STATE["requests"]
        if path.endswith("/chat/completions") and pl_ and "альтернативную" in pl_["messages"][0]["content"]]
check(para and para[0]["messages"][1]["content"] == "Исходный запрос: Что такое СЗИ?", "перефразируется исходный вопрос")
# --- реранкер: все кандидаты, пачками ---
cfg_rr = replace(cfg2, rerank_char_budget=1500, rerank_max_docs=1)
mock_server.STATE["requests"].clear()
_, srcs_rr = app.search_parent_child("требования к журналам и СЗИ", cfg_rr, k_neighbors=4, glossary=app.load_glossary(cfg2))
rr_reqs = [pl_ for m, path, pl_ in mock_server.STATE["requests"] if path.endswith("/rerank") and pl_]
n_docs = sum(len(r["documents"]) for r in rr_reqs)
snap = app.load_kb_snapshot(cfg2)
check(len(rr_reqs) == n_docs >= 2 and all(len(r["documents"]) == 1 for r in rr_reqs),
      f"пачки по ограничениям шлюза: {[len(r['documents']) for r in rr_reqs]}")
check(n_docs == min(cfg_rr.rerank_top_k, len(snap.parent_store)), f"на реранк ушли все кандидаты: {n_docs}")
check(all("rerank —" not in s_ for s_ in srcs_rr), "у всех выбранных источников есть балл реранкера")
filler = "Общие положения регламента описывают порядок работы подразделений. " * 40
pd_ = {"text": filler + "Особое требование: ключи СКЗИ хранятся в сейфе ответственного лица. " + filler,
       "source": "x.docx", "heading_path": "Раздел 5"}
tw = app.query_term_weights("где хранятся ключи СКЗИ")
doc_rr, origin = app.rerank_document(pd_, cfg2, tw)
check("ключи СКЗИ хранятся" in doc_rr and origin == "окно по словам" and len(doc_rr) <= cfg2.rerank_doc_max_chars,
      f"реранкеру — окно с нужной фразой ({origin}, {len(doc_rr)} симв.)")
child = "Особое требование: ключи СКЗИ хранятся в сейфе ответственного лица."
doc_rr2, origin2 = app.rerank_document(pd_, cfg2, {}, child_text=child)
check(child in doc_rr2 and origin2 == "child", "реранкеру — найденный child с окружением")
meta = app.load_kb_snapshot(cfg2).metadata
check(all(m.get("text") for m in meta), "текст child хранится в метаданных")
# --- нормировка векторов ---
kb_raw = WORK / "kb_raw"
cfg_raw = replace(cfg2, embed_model="emb-raw", kb_dir=str(kb_raw))
ok, msg = app.process_and_add_to_kb("Регламент_ИБ.docx", data_a, cfg_raw)
idx_raw = app.faiss.read_index(str(kb_raw / app.KB_FAISS))
norms = app.np.linalg.norm(app.index_vectors(idx_raw), axis=1)
check(ok and app.np.allclose(norms, 1.0, atol=1e-4), f"векторы в индексе нормированы ({norms.min():.3f}…{norms.max():.3f})")
raw_index = app.faiss.IndexFlatL2(idx_raw.d)
raw_index.add(app.index_vectors(idx_raw) * 5.0)              # база v17 с ненормированными векторами
app.faiss.write_index(raw_index, str(kb_raw / app.KB_FAISS))
snap_raw = app.load_kb_snapshot(cfg_raw)
check(app.np.allclose(app.np.linalg.norm(app.index_vectors(snap_raw.index), axis=1), 1.0, atol=1e-4),
      "ненормированная база нормируется в памяти при чтении")
_, s_raw = app.search_parent_child("требования к СЗИ", cfg_raw, k_neighbors=2, glossary=app.load_glossary(cfg_raw))
check(s_raw and "dense #" in s_raw[0], "поиск по ненормированной базе работает")
# --- префиксы эмбеддингов по имени модели ---
lim = app.resolve_model_limits(cfg0, {}, None, {})
cq, _ = app.resolve_settings(replace(cfg0, embed_model="qwen3-embedding-4b"), lim, {})
check(cq.embed_query_prefix.startswith("Instruct:") and cq.embed_doc_prefix == "" and "Qwen3" in cq.embed_prefix_source,
      f"Qwen3: инструкция запроса ({cq.embed_prefix_source})")
ce, _ = app.resolve_settings(replace(cfg0, embed_model="intfloat/multilingual-e5-large"), lim, {})
check((ce.embed_query_prefix, ce.embed_doc_prefix) == ("query: ", "passage: "), "e5: query:/passage:")
cb, _ = app.resolve_settings(replace(cfg0, embed_model="bge-m3"), lim, {})
check(cb.embed_query_prefix == "" and cb.embed_prefix_source == "не нужен", "bge-m3: без префикса")
cx, _ = app.resolve_settings(replace(cfg0, embed_model="qwen3-embedding-4b", embed_query_prefix="Q: ",
                                         explicit_fields=frozenset({"embed_query_prefix"})), lim, {})
check(cx.embed_query_prefix == "Q: " and "config.yaml" in cx.embed_prefix_source, "префикс из config.yaml важнее")
cq_kb = replace(cq, kb_dir=str(WORK / "kb_qwen"), llm_url=cfg0.llm_url, rerank_url="")
mock_server.STATE["requests"].clear()
app.process_and_add_to_kb("Регламент_SLA.docx", data_b, cq_kb)
app.search_parent_child("время восстановления", cq_kb, k_neighbors=2)
emb = [pl_["input"] for m, path, pl_ in mock_server.STATE["requests"] if path.endswith("/embeddings") and pl_]
doc_inputs = [x for inp in emb if isinstance(inp, list) for x in inp]
q_inputs = [inp for inp in emb if isinstance(inp, str)]
check(doc_inputs and not any(x.startswith("Instruct:") for x in doc_inputs), "документы кодируются без инструкции")
check(q_inputs and q_inputs[-1].startswith("Instruct:") and q_inputs[-1].endswith("время восстановления"),
      "запрос кодируется с инструкцией")
_, s_norr = app.search_parent_child("время восстановления", cq_kb, k_neighbors=2)
check(s_norr and "rerank —" in s_norr[0], "без реранкера поиск работает")
# --- разбор замечаний ---
ff = app.finding_from_line("- **[СРЕДНИЙ]** 152-ФЗ не учтён")
check(ff["text"] == "152-ФЗ не учтён" and ff["severity"] == "СРЕДНИЙ", f"разбор строки: {ff}")
check(app.finding_from_line("3. [НИЗКИЙ] — Неоднозначная формулировка")["text"] == "Неоднозначная формулировка",
      "нумерация и тире убираются")
uniq = app.collect_unique_findings([
    {"severity": "НИЗКИЙ", "text": "Не указан алгоритм шифрования ПДн при передаче", "quote": ""},
    {"severity": "КРИТИЧНЫЙ", "text": "Не указан алгоритм шифрования ПДн при передаче.", "quote": "q"},
    {"severity": "СРЕДНИЙ", "text": "Не указан алгоритм резервного копирования журналов", "quote": ""},
])
check(len(uniq) == 2 and uniq[0]["severity"] == "КРИТИЧНЫЙ" and uniq[0]["quote"] == "q",
      f"повторы сливаются с наибольшей критичностью: {[(u['severity'], u['text'][:25]) for u in uniq]}")
# --- Word: примечания, порядок текста, одно место на замечание ---
wd = pydocx.Document()
wd.add_heading("Требования к системе", 1)
wp = wd.add_paragraph()
wp.add_run("Система передаёт ")
wp.add_run("ПДн").bold = True
wp.add_run(" между подсистемами по защищённым каналам.")
wd.add_paragraph("Время отклика интерфейса должно быть приемлемым для пользователей и операторов.")
wd.add_paragraph("Журналы событий хранятся на сервере приложений.")
wpath = WORK / "w.docx"
wd.save(wpath)
orig_texts = [x.text for x in pydocx.Document(wpath).paragraphs]
fnd = [
    {"severity": "КРИТИЧНЫЙ", "text": "Не указан алгоритм шифрования", "category": None,
     "quote": "Система передаёт ПДн между подсистемами"},
    {"severity": "НИЗКИЙ", "text": "Время отклика интерфейса не задано числом: «приемлемым» неоднозначно",
     "category": None, "quote": ""},
    {"severity": "СРЕДНИЙ", "text": "Не определён порядок резервного копирования", "category": None, "quote": ""},
]
out_bytes, st_ = app.annotate_docx_with_findings(wpath.read_bytes(), fnd)
check(st_ == {"mode": "comments", "bound": 2, "by_quote": 1, "unbound": 1}, f"статистика: {st_}")
od = pydocx.Document(io.BytesIO(out_bytes))
texts_after = [x.text for x in od.paragraphs]
check(texts_after[0].startswith("⚠️ Замечания MINT без привязки") and "резервного копирования" in texts_after[1],
      "непривязанные — списком в начале")
check(texts_after[-len(orig_texts):] == orig_texts, "текст ТЗ не изменился (примечания)")
xml = od.element.xml
check(xml.count("commentRangeStart") == 2, f"два примечания: {xml.count('commentRangeStart')}")
first_run_texts = [r.text for r in od.paragraphs[-3].runs if r.text]
check("".join(first_run_texts) == orig_texts[-3] and "Система передаёт ПДн между подсистемами" in first_run_texts[0] + "".join(first_run_texts[1:3]),
      f"run'ы разделены по цитате, порядок сохранён: {first_run_texts}")
out2, st2 = app.annotate_docx_with_findings(wpath.read_bytes(), fnd, use_comments=False)
od2 = pydocx.Document(io.BytesIO(out2))
t2 = [x.text for x in od2.paragraphs]
check(st2["mode"] == "inline" and t2[-3].startswith(orig_texts[-3]) and "[ЗАМЕЧАНИЕ MINT]" in t2[-3],
      f"режим вставок: текст абзаца не переставлен: {t2[-3][:70]!r}")
check(sum("[ЗАМЕЧАНИЕ MINT]" in x for x in t2) == 2, "каждое замечание — в одном месте")
# --- очистка таблиц ---
row = "| Требование к шифрованию | Требование к шифрованию | Описание канала |"
check(app.clean_docling_markdown(row).count("Требование к шифрованию") == 1, "соседний повтор ячейки убран")
row2 = "| Значение по умолчанию | Да | Значение по умолчанию |"
check(app.clean_docling_markdown(row2).count("Значение по умолчанию") == 2, "несоседние одинаковые ячейки сохранены")
# --- блокировка записи: два документа одновременно ---
import threading as _th
kb_lock = WORK / "kb_lock"
cfg_lock = replace(cfg2, kb_dir=str(kb_lock))
results_ = []
threads = [_th.Thread(target=lambda n, d: results_.append(app.process_and_add_to_kb(n, d, cfg_lock)),
                      args=(n_, d_)) for n_, d_ in (("A.docx", data_a), ("B.docx", data_b))]
for t_ in threads:
    t_.start()
for t_ in threads:
    t_.join()
docs_lock, nvec_lock = app.list_kb_documents(cfg_lock)
parents_lock = app._read_parent_store(cfg_lock)
check(docs_lock == ["A.docx", "B.docx"] and len({v["source"] for v in parents_lock.values()}) == 2,
      f"параллельная загрузка: оба документа в базе ({docs_lock}, {nvec_lock} векторов)")
from docling.document_converter import DocumentConverter as _DC
check(_DC.created == 1, f"Docling-конвертер создан один раз ({_DC.created})")


print("15. [unit + интеграция] Рассуждения, ссылки на разделы, вопросительные слова, сбой реранкера, история чата")
# --- 1. рассуждения ---
check(app.strip_reasoning("<think>\n- мысль\n</think>\nОтвет") == "Ответ", "закрытый <think> убран")
check(app.strip_reasoning("мысли без начала\n</think>\nОтвет") == "Ответ", "только закрывающий тег — убрано всё до него")
check(app.strip_reasoning("<think>не успел") == "", "незакрытый <think> — ответа нет")
check(app.strip_reasoning("Обычный ответ с <b>тегом</b>") == "Обычный ответ с <b>тегом</b>", "обычный текст не трогается")
cfg_r1 = replace(cfg2, llm_model="r1-model")
r1_answer = app.ask_llm(app.CHAT_SYSTEM_PROMPT, "вопрос", cfg_r1)
check("<think>" not in r1_answer and "рассуждение" not in r1_answer and r1_answer.startswith("Ответ"), f"ответ R1 без рассуждений: {r1_answer[:60]!r}")
res_r1 = app.audit_section(sections[0], cfg_r1, app.DEFAULT_AUDIT_FOCUS, True, 3, glossary=app.load_glossary(cfg2))
check(len(res_r1["findings"]) == 2 and not any("рассужд" in f["text"] for f in res_r1["findings"]) and res_r1["dropped"],
      f"аудит с R1: рассуждения не стали замечаниями, самокритика сработала ({len(res_r1['findings'])})")
cut = app.ask_llm("s", "вопрос", cfg_r1, max_tokens=10)
check(cut.startswith("❌ [ОТВЕТ ОБРЕЗАН]"), f"обрезанный ответ R1 → понятная ошибка: {cut[:40]}")
check(app.ask_llm("s", "вопрос", replace(cfg2, llm_model="r1-noopen")).startswith("Ответ"), "шаблон без <think> в начале")
tr = app.ask_llm(app.CHAT_SYSTEM_PROMPT, "вопрос", replace(cfg2, llm_model="trunc-model"))
check(tr.startswith("Ответ") and "⚠️ Ответ обрезан" in tr, "обрезанный ответ помечен")
# --- 2. разделы внутри фрагмента ---
pd_h = {"source": "r.docx", "heading_path": "1. Общие → 1.1 Шифрование",
        "text": "## 1.1 Шифрование\n\nТекст.\n\n## 1.2 Журналирование\n\nТекст.\n\n### 1.2.1 Хранение\n\nТекст."}
hdr = app._source_header(pd_h)
check("Раздел в начале фрагмента: 1. Общие → 1.1 Шифрование" in hdr and "далее во фрагменте: 1.2 Журналирование; 1.2.1 Хранение" in hdr,
      f"заголовок источника: {hdr}")
check("ближайший заголовок" in app.CHAT_SYSTEM_PROMPT, "правило ссылок в промпте")
items_h, _ = app.search_parent_child("Сколько хранятся журналы?", cfg2, k_neighbors=2, return_items=True,
                                     glossary=app.load_glossary(cfg2))
check(any("далее во фрагменте:" in it["block"] and "Журналирование" in it["block"].split("===")[1] for it in items_h),
      "в контексте модели указаны разделы внутри фрагмента")
# --- 3. вопросительные слова ---
check(all(app.normalize_token(w) in app.BM25_STOP_WORDS for w in ("сколько", "чем", "каков", "почему")), "вопросительные слова — стоп-слова")
diag_q = {}
app.search_parent_child("Сколько хранятся журналы?", cfg2, k_neighbors=2, glossary=app.load_glossary(cfg2), diagnostics=diag_q)
check(diag_q.get("lexical_cov") == 1.0, f"доля слов вопроса в контексте: {diag_q.get('lexical_cov')}")
# --- 4. частичный сбой реранкера ---
cfg_rr1 = replace(cfg2, rerank_max_docs=1)
mock_server.STATE["rerank_fail_from"] = mock_server.STATE["rerank_count"] + 2
sc_part = app.rerank_scores("запрос", ["документ один", "документ два", "документ три"], cfg_rr1)
mock_server.STATE["rerank_fail_from"] = None
check(sc_part is None, "часть запросов к реранкеру упала → баллы не используются")
# --- 5. история чата ---
msgs = [{"role": "user", "content": "Какие требования к журналам?"},
        {"role": "assistant", "content": "Журналы хранятся три года. " * 80},
        {"role": "user", "content": "Что такое СЗИ?"},
        {"role": "assistant", "content": "❌ [ТАЙМАУТ MINT]: модель не ответила"},
        {"role": "user", "content": "А шифрование?"},
        {"role": "assistant", "content": "Используется ГОСТ."}]
h = app.chat_history(msgs, 2)
check(len(h) == 4 and h[0]["content"] == "Какие требования к журналам?" and h[1]["content"].endswith("…")
      and len(h[1]["content"]) <= app.CHAT_HISTORY_ANSWER_CHARS + 2 and h[3]["content"] == "Используется ГОСТ.",
      "история: последние пары, ошибки пропущены, длинный ответ обрезан")
check(app.chat_history(msgs, 0) == [], "история выключается (0)")
cq_ = app.condense_question(h[:2], "А сколько они хранятся?", cfg2)
check(cq_ == "А сколько они хранятся — про журналы событий?", f"уточняющий вопрос переписан (без <think>): {cq_!r}")
check(app.condense_question([], "Вопрос", cfg2) == "Вопрос", "без истории вопрос не меняется")
mock_server.STATE["requests"].clear()
ans_h = app.ask_llm(app.CHAT_SYSTEM_PROMPT, app.chat_user_content("ctx", "А сколько?", "", cq_), cfg2, history=h[:2])
sent = mock_server.STATE["last_chat"]["messages"]
check([m["role"] for m in sent] == ["system", "user", "assistant", "user"] and "с учётом предыдущих реплик" in sent[-1]["content"],
      "модель получает историю и переписанный вопрос")

print("16. [unit] FAISS: совместимость формата и путь с кириллицей")
import faiss  # noqa: E402  (настоящий faiss-cpu или заглушка)
import numpy as np  # noqa: E402
_idx = faiss.IndexFlatL2(4)
_idx.add(np.eye(4, dtype="float32"))
_old = WORK / "faiss_old.index"
faiss.write_index(_idx, str(_old))                       # так писали базы до 20.0
check(app.read_faiss(_old).ntotal == 4, "база, записанная faiss.write_index, читается (старые базы открываются)")
_new = WORK / "faiss_new.index"
app.write_faiss(_idx, _new)
check(faiss.read_index(str(_new)).ntotal == 4, "новый файл индекса читается и faiss.read_index (формат тот же)")
_cyr = WORK / "база знаний Иванов"
_cyr.mkdir()
app.write_faiss(_idx, _cyr / "child_index.faiss")
_back = app.read_faiss(_cyr / "child_index.faiss")
check(_back.ntotal == 4 and _back.d == 4, f"индекс пишется и читается по пути с кириллицей (faiss: {getattr(faiss, '__file__', '?')})")

print("17. [unit + интеграция] Нарезка текста, порядок эмбеддингов, статусы, глоссарий")
# --- обрезка фрагмента: ранняя граница не съедает текст ---
long_t = "Заголовок\n" + " ".join(["слово"] * 400)
check(len(app.smart_truncate(long_t, 1500)) >= 1400,
      f"«Заголовок + длинный абзац» обрезается по лимиту, а не до заголовка: {len(app.smart_truncate(long_t, 1500))}")
check(app.smart_truncate("Первое. " * 300, 1000).endswith("."), "обычная обрезка — по границе предложения")
# --- чанки: абзац, предложение, заголовок, перекрытие с целого слова ---
p1 = ("Первый абзац о шифровании данных. " * 20).strip()
p2 = ("Второй абзац про резервное копирование и время восстановления сервиса " * 3).strip()
ch = app.split_with_overlap_pos(p1 + "\n\n" + p2, 800, 0)
check(ch[0][2] == p1, f"чанк заканчивается на границе абзаца: …{ch[0][2][-30:]!r}")
p3 = " ".join(f"Предложение номер {i} о регламенте хранения журналов." for i in range(40))
check(all(c.endswith(".") for _, _, c in app.split_with_overlap_pos(p3, 500, 0)),
      "без абзацев чанки заканчиваются на границе предложения")
stream = ("Текст раздела один. " * 25).strip() + "\n\n## 1.2 Журналирование\n\n" + ("Журналы хранятся три года. " * 20).strip()
ch2 = app.split_with_overlap_pos(stream, 600, 100)
check(not any(c.rstrip().splitlines()[-1].startswith("#") for _, _, c in ch2[:-1]),
      "заголовок не остаётся в конце чанка — переносится к своему тексту")
check(any(c.startswith("## 1.2 Журналирование") or "\n## 1.2 Журналирование\n" in c for _, _, c in ch2[1:]),
      "заголовок — в чанке со своим текстом")
check(all(stream[s:e] == c for s, e, c in ch2) and all(s == 0 or stream[s - 1] in " \n" for s, _, _ in ch2),
      "смещения верны; перекрытие начинается с целого слова")
# --- разделы аудита: большая таблица режется, шапка повторяется ---
big_table = "| Параметр | Значение |\n|---|---|\n" + "\n".join(
    f"| Параметр номер {i} системы | Значение {i} в миллисекундах |" for i in range(400))
secs = app.split_into_sections("## Раздел 1\n\nКороткий текст.\n\n" + big_table + "\n\n## Раздел 2\n\nЕщё текст.", 4500)
check(max(len(s) for s in secs) <= 4500, f"разделы не больше лимита: {[len(s) for s in secs]}")
check(all("| Параметр | Значение |" in s for s in secs[:-1]) and "Параметр номер 399" in "".join(secs),
      "каждая часть таблицы — с шапкой, строки не потеряны")
secs2 = app.split_into_sections("Вступление. " * 30 + "\n\n## 2. Шифрование\n\n" + "Данные шифруются по ГОСТ. " * 10, 400)
check(not any(s.rstrip().endswith("## 2. Шифрование") for s in secs2) and any(s.startswith("## 2. Шифрование") for s in secs2),
      f"заголовок раздела — вместе со своим текстом: {[s[:20] for s in secs2]}")
# --- эмбеддинги: порядок по index, а не по положению в ответе ---
texts_o = ["журналы событий хранятся три года", "резервное копирование ежедневно", "шифрование по ГОСТ"]
mock_server.STATE["shuffle_embeddings"] = True
emb_calls = lambda: sum(1 for _m, path_, _p in mock_server.STATE["requests"] if path_.endswith("/embeddings"))  # noqa: E731
n_emb0 = emb_calls()
try:
    batch_o = app.get_query_embeddings_batch(texts_o, cfg2)
finally:
    mock_server.STATE["shuffle_embeddings"] = False
n_emb1 = emb_calls()
single_o = [app.get_query_embedding(t, cfg2) for t in texts_o]
check(batch_o == single_o, "ответ шлюза в другом порядке → векторы по полю index")
check(n_emb1 - n_emb0 == 1, f"пачка текстов — одним запросом к эмбеддингам, а не по одному: {n_emb1 - n_emb0}")
# --- контроль исправлений: «исправлено частично» — это частично ---
src_v = "При передаче ПДн применяется шифрование по ГОСТ 34.12-2018 на всех каналах."
for st_ in ("ИСПРАВЛЕНО_ЧАСТИЧНО", "PARTIALLY_FIXED", "ЧАСТИЧНО"):
    lbl_, _ = app.parse_verification(json.dumps({"status": st_, "quote": src_v, "comment": "к"}, ensure_ascii=False), src_v)
    check(lbl_.startswith("⚠️"), f"статус {st_} → {lbl_}")
lbl_, _ = app.parse_verification(json.dumps({"status": "ИСПРАВЛЕНО", "quote": src_v, "comment": "к"}, ensure_ascii=False), src_v)
check(lbl_.startswith("✅"), f"ИСПРАВЛЕНО → {lbl_}")
# --- глоссарий: полная форма в вопросе → сокращение в BM25 ---
gl_r = app.Glossary({}, {"СЗИ": "средства защиты информации", "ПО": "программное обеспечение"})
t_r = gl_r.bm25_terms("Требования к средствам защиты информации")
check(("сзи", 0.5) in t_r, f"«средства защиты информации» → СЗИ: {t_r}")
check(not any(t == "по" for t, _ in gl_r.bm25_terms("обновление программного обеспечения")),
      "сокращение-стоп-слово («ПО») не добавляется")
check(not any(t == "сзи" for t, _ in gl_r.bm25_terms("требования к СЗИ")), "сокращение из вопроса не дублируется")
eng_r = app.LightweightBM25(["СЗИ устанавливаются на каждом АРМ", "Порядок резервного копирования"])
q_r = "Требования к средствам защиты информации"
check(eng_r.get_scores(q_r)[0] == 0 and eng_r.get_scores(q_r, extra_terms=gl_r.bm25_terms(q_r))[0] > 0,
      "документ, где есть только сокращение, находится по полной форме")

print("18. [unit + интеграция] Контекст: склейка соседних фрагментов, окно вокруг найденного, проверка ссылок")
# --- склейка: соседние фрагменты одного документа → один блок без повтора перекрытия ---
seq = " ".join(f"Пункт {i}: требование номер {i} к подсистеме описано в этом абзаце регламента." for i in range(1, 60))
pp = app.split_with_overlap_pos(seq, 900, 200)
st_m = {f"m{i}": {"text": t, "source": "r.docx", "type": "paragraph"} for i, (_, _, t) in enumerate(pp)}
st_m["tb"] = {"text": "Таблица №1\n| a | b |\n|---|---|\n| 1 | 2 |", "source": "r.docx", "type": "table"}
ord_m = {p: i for i, p in enumerate(st_m)}
g_m = app.merge_adjacent(["m2", "m1"], st_m, ord_m)
check(len(g_m) == 1 and g_m[0][0] == ["m1", "m2"] and g_m[0][1] == seq[pp[1][0]:pp[2][1]],
      "соседи склеены в порядке текста, перекрытие не повторяется")
check([g[0] for g in app.merge_adjacent(["m3", "m0", "m1"], st_m, ord_m)] == [["m3"], ["m0", "m1"]],
      "блоки — в порядке релевантности лучшего участника")
check([g[0] for g in app.merge_adjacent(["m0", "m2"], st_m, ord_m)] == [["m0"], ["m2"]], "не соседи — не склеиваются")
check(app.join_overlapping("Первый фрагмент текста без связи.", "Совсем другой фрагмент, без общего начала.") is None,
      "без перекрытия не склеивается")
last_m = f"m{len(pp) - 1}"            # последний текстовый фрагмент — сосед таблицы по порядку в базе
check([g[0] for g in app.merge_adjacent([last_m, "tb"], st_m, ord_m)] == [[last_m], ["tb"]],
      "таблица не склеивается с текстом")
# --- окно вокруг найденного: длинный фрагмент режется не с начала ---
long_doc = ("## 1. Общие положения\n\n" + "Общие положения регламента описывают порядок работы. " * 30
            + "\n\n## 2. Хранение ключей\n\n" + "Порядок учёта носителей определяет администратор. " * 10
            + "Ключи СКЗИ хранятся в сейфе ответственного лица. " + "Порядок учёта носителей определяет администратор. " * 3)
target = "Ключи СКЗИ хранятся в сейфе ответственного лица."
check(target not in app.smart_truncate(long_doc, 1200), "(исходное положение: обрезка с начала теряет нужное)")
fw = app.focus_passage(long_doc, 1200, term_weights=app.query_term_weights("где хранятся ключи СКЗИ"))
fc = app.focus_passage(long_doc, 1200, anchor=target)
check(target in fw and target in fc and len(fc) <= 1200 + 60, f"окно вокруг найденного (по словам и по child): {len(fc)} симв.")
check(fc.startswith("## 1. Общие положения\n\n… ") and "\n## 2. Хранение ключей\n" in fc,
      f"окно начинается в разделе 1 — перед ним его заголовок, заголовок раздела 2 внутри: {fc[:40]!r}")
fn = app.focus_passage(long_doc, 500, anchor=target)
check(target in fn and fn.startswith("## 2. Хранение ключей\n\n… "),
      f"окно внутри раздела 2 — перед ним заголовок раздела 2 (ссылка на раздел будет верной): {fn[:40]!r}")
check(app.focus_passage(long_doc, 1200, term_weights=app.query_term_weights("общие положения регламента"))
      == app.smart_truncate(long_doc, 1200), "найденное в начале — обрезка как раньше")
import random  # noqa: E402
rnd = random.Random(11)
fz_words = "система журнал ключи доступ порядок хранение данные сервис требования ролям".split()
fz_bad = fz_n = 0
for _ in range(300):
    fz_parts = []
    for h_ in range(rnd.randint(2, 6)):
        fz_parts += [f"## {h_ + 1}. Раздел. Подраздел {h_ + 1}",
                     "".join(" ".join(rnd.choice(fz_words) for _ in range(rnd.randint(5, 15))).capitalize()
                             + rnd.choice([". ", "; "]) for _ in range(rnd.randint(3, 25))).strip()]
    fz_text = "\n\n".join(fz_parts)
    fz_heads = {ln for ln in fz_text.splitlines() if ln.startswith("#")}
    fz_limit = rnd.choice([300, 500, 900, 1500])
    if len(fz_text) <= fz_limit:
        continue
    a0_ = rnd.randrange(len(fz_text) - 60)
    fz_anchor = fz_text[a0_:a0_ + rnd.randint(40, 200)]
    fz_out = app.focus_passage(fz_text, fz_limit, anchor=fz_anchor)
    fz_n += 1
    fz_bad += (fz_anchor.strip() not in fz_out or len(fz_out) > fz_limit + 160
               or any(ln.startswith("#") and ln not in fz_heads for ln in fz_out.splitlines()))
check(fz_n > 200 and fz_bad == 0, f"окно на {fz_n} случайных текстах: найденное внутри, лимит, заголовки целые ({fz_bad} нарушений)")
big_tab = "Таблица №2. Сроки\n| Объект | Срок |\n|---|---|\n" + "\n".join(f"| Объект {i} | {i} дней |" for i in range(200)) \
          + "\n| Журналы аудита | 3 года |"
ft = app.focus_passage(big_tab, 600, term_weights=app.query_term_weights("срок хранения журналов аудита"), is_table=True)
check(ft.startswith("Таблица №2. Сроки\n| Объект | Срок |\n|---|---|") and "| Журналы аудита | 3 года |" in ft
      and len(ft) <= 600, f"таблица: шапка + строки вокруг найденной ({len(ft)} симв.)")
# --- поиск: склейка и окно в настоящем контексте ---
kb_m = WORK / "kb_merge"
cfg_m = replace(cfg2, kb_dir=str(kb_m), parent_size=700, parent_overlap=150, child_size=250, child_overlap=60)
merge_bytes = make_docx(docs_dir / "Ключи.docx", [
    ("h1", "1. Общие положения"),
    ("p", " ".join(f"Пункт {i}: порядок работы подразделения номер {i} определяется приказом." for i in range(1, 16))),
    ("h1", "2. Хранение ключей"),
    ("p", " ".join(f"Пункт {i}: носитель номер {i} учитывается в журнале." for i in range(16, 26))
          + " Пункт 26: ключи СКЗИ хранятся в сейфе ответственного лица."),
])
ok_m, msg_m = app.process_and_add_to_kb("Ключи.docx", merge_bytes, cfg_m)
n_par = len(app.load_kb_snapshot(cfg_m).parent_store)
items_m, _ = app.search_parent_child("где хранятся ключи СКЗИ", cfg_m, k_neighbors=n_par, return_items=True,
                                     glossary=app.load_glossary(cfg_m))
blk = "\n".join(it["block"] for it in items_m)
check(ok_m and n_par >= 3 and any(len(it["ids"]) > 1 for it in items_m), f"соседние фрагменты склеены: {[len(it['ids']) for it in items_m]} из {n_par}")
check(all(blk.count(f"Пункт {i}:") == 1 for i in range(1, 27)), "в контексте нет повторов перекрытия")
check(sum(len(it["ids"]) for it in items_m) == n_par, "все найденные фрагменты на месте (ids)")
ctx_f, _ = app.search_parent_child("где хранятся ключи СКЗИ", replace(cfg_m, smart_truncate_chars=300), k_neighbors=2,
                                   glossary=app.load_glossary(cfg_m))
check("ключи СКЗИ хранятся в сейфе" in ctx_f, "короткий лимит: в контексте — найденное место, а не начало фрагмента")
# --- проверка ссылок в ответе ---
ctx_c = ("=== ИСТОЧНИК: Регламент_ИБ.docx | Раздел в начале фрагмента: 1. Общие → 1.2 Журналирование ===\n"
         "Текст.\n\n## 1.3 Резервное копирование\n\nТекст.\n\n"
         "=== ИСТОЧНИК: Регламент_SLA.docx | Таблица №1. Показатели (часть 1/2) ===\nТаблица №1. Показатели\n| a |")
ans_c = ("A [Источник: Регламент_ИБ.docx, 1.2 Журналирование]. B [Источник: Регламент ИБ, раздел 1.3]. "
         "C [Источник: Регламент_SLA.docx, Таблица №1. Показатели]. D [Источник: Регламент_ИБ.docx]. "
         "E [Источник: Регламент_ИБ.docx, 4.1 Шифрование каналов]. F [Источник: Справочник сокращений]. "
         "G [Источник: Регламент_ИБ.docx, 1.2 Журналирование; Регламент_SLA.docx]. H [Source: Регламент_SLA.docx].")
cc = app.check_citations(ans_c, ctx_c)
check([c["ok"] for c in cc] == [True, True, True, True, False, False, True, True],
      f"ссылки сверены с контекстом: {[(c['text'][:30], c['ok'], c['reason']) for c in cc]}")
check([c["reason"] for c in cc if not c["ok"]] == ["section", "file"], "причины: нет раздела / нет файла")
marked = app.mark_invalid(ans_c, cc)
check(marked.count("⚠️") == 2 and "Шифрование каналов] ⚠️" in marked and "Справочник сокращений] ⚠️" in marked,
      "пометка ⚠️ — только у ссылок мимо контекста")
check(app.citation_warning([]) == "" and "(2)" in app.citation_warning([cc[4]["text"], cc[5]["text"]]),
      "предупреждение под ответом")

print("18a. [интеграция] Сценарии без интерфейса: answer_question, run_audit, verify_remarks")
gl2 = app.load_glossary(cfg2)
ver = app.prompts_version()
qa = app.answer_question("Какие требования к СЗИ?", cfg2, glossary=gl2)
check(qa["answer"].startswith("Ответ по контексту") and "Справочник: да" in qa["answer"] and qa["sources"]
      and qa["search_query"] is None and not qa["error"] and not qa["low_conf"] and qa["prompts"] == ver
      and qa["context"].startswith("=== ИСТОЧНИК:"), f"ответ чата одним вызовом: {qa['answer'][:60]!r}, {qa['prompts']}")
qa2 = app.answer_question("А сколько они хранятся?", cfg2, glossary=gl2, history=h[:2])
check(qa2["search_query"] == "А сколько они хранятся — про журналы событий?", f"уточняющий вопрос: {qa2['search_query']!r}")
qa3 = app.answer_question("рецепт борща со свёклой", replace(cfg2, top_k=0))
check(qa3["low_conf"] and qa3["context"].startswith("[СИГНАЛ ПОИСКА") and qa3["sources"],
      "вопрос не по базе: сигнал «уверенность низкая»; top_k = 0 — глубина по размеру базы")
check(app.is_no_info_answer("В предоставленных документах информация отсутствует. Уточните вопрос.")
      and app.is_no_info_answer("В предоставленных регламентах информация отсутствует.")     # прежняя фраза (до 21.0)
      and app.is_no_info_answer("The provided documents do not contain this information.")
      and not app.is_no_info_answer("Журналы хранятся три года [Источник: Регламент_ИБ.docx]."),
      "[unit] ответ «информации нет» узнаётся на обоих языках")
events = []
au = app.run_audit(md, "ТЗ.docx", cfg2, focus=app.DEFAULT_AUDIT_FOCUS, glossary=gl2,
                   on_progress=lambda ev, **d: events.append(ev))
check(len(au["findings"]) == 1 and "ТРЕБУЕТ ДОРАБОТКИ" in au["report"] and au["sections"] == len(sections)
      and not au["failed_sections"] and au["prompts"] == ver and au["document_mode"] == "single",
      f"аудит документа одним вызовом: замечаний {len(au['findings'])}, разделов {au['sections']}")
check([f["unverified"] for f in au["unverified"]] == [["152-ФЗ"]] and "НЕ ПОДТВЕРЖДЕНО ИСТОЧНИКОМ" in au["report"]
      and "152-ФЗ" in au["report"].split("НЕ ПОДТВЕРЖДЕНО")[1],
      f"ссылки на акт нет ни в документе, ни в нормах — замечание в блоке «Не подтверждено»: {au['unverified']}")
n_deep = events.count("deep_task")
check(events == ["sections", "checks"] + ["section"] * au["sections"] + ["deep"] + ["deep_task"] * n_deep + ["merge", "summary"]
      and n_deep == len(au["deep"]) >= 2, f"ход аудита сообщается по разделам и вопросам глубокой проверки: {events}")
check(f"**Промпты:** {ver}" in au["log_md"] and "Часть 1/" in au["log_md"], "в логе проверки — версия промптов")
au2 = app.run_audit(md, "ТЗ.docx", cfg2, focus="  ", special=" шифрование ПДн ", use_wiki=False, glossary=gl2)
check(au2["focus"] == app.GENERIC_AUDIT_FOCUS and au2["special"] == "шифрование ПДн" and not au2["focus_items"]
      and "**Нормативная база:** нет" in au2["log_md"], "пустые критерии → общий аудит; без базы нормы не ищутся")
check(app.parse_remarks("- [КРИТИЧНЫЙ] Не указан алгоритм шифрования ПДн\n\nкоротко\n  2) Не задан SLA доступности  ")
      == ["- [КРИТИЧНЫЙ] Не указан алгоритм шифрования ПДн", "2) Не задан SLA доступности"],
      "[unit] замечания — по строкам, короткие строки пропускаются")
ev_r = []
rv = app.verify_remarks("- [КРИТИЧНЫЙ] Не указан алгоритм шифрования ПДн\nкоротко",
                        app.compile_docx_to_markdown("ТЗ_v2.docx", v2), replace(cfg2, audit_map_size=200), glossary=gl2,
                        on_progress=lambda ev, **d: ev_r.append(ev))
check(len(rv["rows"]) == 1 and rv["rows"][0]["status"].startswith("✅") and rv["unverified"] == 0 and rv["dense"]
      and "МАТРИЦА" in rv["matrix_md"] and rv["prompts"] == ver and ev_r == ["sections", "index", "remark"],
      f"контроль исправлений одним вызовом: {[(r['status'], r['remark'][:20]) for r in rv['rows']]}, {ev_r}")

print("19. [сквозной] Интерфейс: действия пользователя на странице (заглушка Streamlit)")
os.environ["EMBED_MODEL"] = "emb-model-2"


def run_ui(chat=None, buttons=(), uploads=None):
    st.LOG.clear()
    st.SCRIPT["chat_input"] = chat
    st.SCRIPT["buttons"] = set(buttons)
    st.SCRIPT["file_uploader"] = uploads or {}
    try:
        app.main()
        return "ok"
    except st.RerunException:
        return "rerun"


try:
    check(run_ui() == "ok", "первый запуск страницы")
    sliders = [x for x in st.LOG if x[0] in ("slider", "number_input")]
    check(len(sliders) > 30, f"виджетов настроек: {len(sliders)}")
    i_panel = next((i for i, x in enumerate(st.LOG) if x == ("expander", "⚙️ Настройки аудита")), -1)
    in_panel = [x for x in st.LOG[i_panel:] if x[0] in ("slider", "checkbox")] if i_panel >= 0 else []
    audit_labels = ("Проходов проверки", "Аудит с рассуждениями", "«Главное» в начале", "🛡️ Самокритика",
                    "Раздел документа на одну проверку", "Аудит: запросов к модели одновременно",
                    "🧠 Подключать нормативную базу", "🔬 Глубокая проверка")
    check(i_panel > 0 and all(sum(str(x[1]).startswith(a) for x in in_panel) == 1 for a in audit_labels)
          and not any(str(x[1]).startswith(audit_labels) for x in st.LOG[:i_panel] if x[0] in ("slider", "checkbox")),
          "настройки аудита — во вкладке аудита, одним блоком (в боковой панели их нет)")
    check(run_ui(chat="Какие требования к СЗИ?") == "ok", "вопрос в чате")
    last = st.session_state["messages"][-1]
    check("Справочник: да" in last["content"] and last["sources"], f"ответ в чате: {last['content']}")
    check(run_ui(chat="А сколько хранятся журналы?") == "ok", "уточняющий вопрос в чате")
    last2 = st.session_state["messages"][-1]
    chat_msgs = mock_server.STATE["last_chat"]["messages"]
    check(last2.get("search_query") and len(chat_msgs) == 4, f"поиск по переписанному вопросу, история передана: {last2.get('search_query')!r}")
    fb_file = WORK / "logs" / "feedback.jsonl"
    aid = last2.get("id")
    check(bool(aid) and ("button", "👍") in st.LOG and ("button", "👎") in st.LOG, "под ответом есть 👍/👎")
    check(run_ui(buttons={f"fb_down_chat_{aid}"}) == "ok" and fb_file.exists(), "оценка 👎 записана в feedback.jsonl")
    fb_rows = [json.loads(x) for x in fb_file.read_text(encoding="utf-8").splitlines()]
    r0 = fb_rows[-1] if fb_rows else {}
    check(r0.get("kind") == "chat" and r0.get("rating") == "down" and r0.get("id") == aid
          and r0.get("question") == "А сколько хранятся журналы?" and r0.get("answer") == last2["content"]
          and r0.get("sources") and r0.get("settings", {}).get("llm_model") == "test-llm"
          and r0.get("mint_version") == app.MINT_VERSION, f"в отзыве вопрос, ответ, источники, настройки: {r0}")
    check(r0.get("prompts") == app.prompts_version() and r0.get("settings", {}).get("lang") == "ru",
          f"в отзыве — версия промптов и язык: {r0.get('prompts')}, {r0.get('settings', {}).get('lang')}")
    check(any(x[0] == "text_input" for x in st.LOG), "после оценки предлагается комментарий")
    st.session_state[f"fb_text_chat_{aid}"] = "Не тот раздел регламента"
    check(run_ui(buttons={f"fb_send_chat_{aid}"}) == "rerun", "комментарий отправлен")
    fb_rows = [json.loads(x) for x in fb_file.read_text(encoding="utf-8").splitlines()]
    check(fb_rows[-1].get("rating") == "comment" and fb_rows[-1].get("comment") == "Не тот раздел регламента"
          and fb_rows[-1].get("id") == aid, "комментарий связан с оценкой")
    check(run_ui() == "ok" and any("Оценка и комментарий записаны" in str(x[1]) for x in st.LOG)
          and sum(1 for x in st.LOG if x == ("button", "👍")) == 1, "после комментария — спасибо, кнопки только у другого ответа")
    check(any(x[0] == "expander" and str(x[1]).startswith("💬 Отзывы: 1") for x in st.LOG)
          and ("download_button", "⬇️ Скачать отзывы (feedback.jsonl)") in st.LOG, "сводка отзывов в боковой панели")
    summ, n_comments = app.feedback_summary(fb_rows + [{"kind": "chat", "id": aid, "rating": "up"}])
    check(summ == {"chat": {"up": 1}} and n_comments == 1, f"сводка считает последнюю оценку: {summ}")
    check(run_ui(buttons={"chat_clear_btn"}) == "ok" and st.session_state["messages"] == [], "кнопка «Новый диалог»")
    check(run_ui(chat="Какие требования к СЗИ? ВЫДУМАЙ ССЫЛКУ") == "ok", "ответ с выдуманной ссылкой")
    am = st.session_state["messages"][-1]
    warns_ = [x[1] for x in st.LOG if x[0] == "warning" and "Несуществующий.docx" in str(x[1])]
    check(warns_ and am["citations_invalid"] == ["[Источник: Несуществующий.docx, 9.9 Раздел]"]
          and am["shown"].count("⚠️") == 1 and "⚠️" not in am["content"],
          f"ссылка мимо контекста помечена, под ответом — предупреждение, в истории ответ без пометок: {am.get('shown')!r}")
    check(run_ui() == "ok" and any(x[0] == "warning" and "Несуществующий.docx" in str(x[1]) for x in st.LOG),
          "предупреждение о ссылке не пропадает после перерисовки")
    check(run_ui(buttons={"chat_clear_btn"}) == "ok", "новый диалог")
    check(run_ui(buttons={"detect_limits_btn"}) == "rerun", "кнопка «Определить у сервера»")
    check("Сервер:" in st.session_state.get("_limits_msg", "") and "база на русском" in st.session_state["_limits_msg"],
          st.session_state.get("_limits_msg", "")[:300])
    check(run_ui() == "ok", "страница после определения лимитов")
    st.session_state["mint_manual"]["parent_size"] = 1800
    st.session_state["mint_manual_limits"][("test-llm", "llm_context_tokens")] = 4096
    check(run_ui() == "ok", "ручные значения (parent 1800, окно 4096)")
    texts = " ".join(str(x[1]) for x in st.LOG)
    check("· вручную" in texts, "ручные значения помечены в интерфейсе")
    check(run_ui(buttons={"reset_settings_btn"}) == "ok" and not st.session_state["mint_manual"], "сброс сеанса")

    class Upl:
        def __init__(self, name, data):
            self.name, self._d = name, data
        def getbuffer(self):
            return memoryview(self._d)

    check(run_ui(uploads={"kb_uploader": [Upl("Регламент_SLA.docx", data_b)]},
                 buttons={"kb_index_all_btn"}) == "rerun", "загрузка документа через интерфейс")
    check(all(ok for _, ok, _ in st.session_state["_kb_batch_results"]), str(st.session_state["_kb_batch_results"]))
    check(run_ui(uploads={"tab2_audit_uploader": Upl("ТЗ.docx", tz_bytes)}, buttons={"tab2_run_btn"}) == "ok",
          "аудит через интерфейс")
    check("ТРЕБУЕТ ДОРАБОТКИ" in st.session_state.get("tab2_final_report_data", ""), "заключение аудита в сеансе")
    log_ = st.session_state.get("tab2_audit_log") or []
    check(log_ and "Часть 1/" in st.session_state.get("tab2_audit_log_md", "")
          and "Версия MINT" in st.session_state["tab2_audit_log_md"], "лог проверки сохранён и выгружается")
    check(run_ui(uploads={"tab2_audit_uploader": Upl("ТЗ.docx", tz_bytes)}) == "ok", "перерисовка после аудита")
    audit_id = st.session_state.get("tab2_audit_id")
    check(run_ui(uploads={"tab2_audit_uploader": Upl("ТЗ.docx", tz_bytes)}, buttons={f"fb_up_audit_{audit_id}"}) == "ok",
          "оценка аудита")
    ra = json.loads(fb_file.read_text(encoding="utf-8").splitlines()[-1])
    check(ra.get("kind") == "audit" and ra.get("rating") == "up" and ra.get("file") == "ТЗ.docx"
          and "ТРЕБУЕТ ДОРАБОТКИ" in ra.get("report", "") and ra.get("sections")
          and ra.get("prompts") == app.prompts_version(), f"в отзыве об аудите — заключение и версия промптов: {list(ra)}")

    exp_ = [x for x in st.LOG if x[0] == "expander" and "Лог проверки: Часть" in str(x[1])]
    dls_ = [x[1] for x in st.LOG if x[0] == "download_button"]
    check(len(exp_) == len(log_) and "📥 Скачать лог проверки (.md)" in dls_,
          f"лог разделов не пропал после перерисовки: {len(exp_)} из {len(log_)}")
    check("📥 Скачать сырые ответы модели (.md)" in dls_, "рядом с логом — кнопка «Скачать сырые ответы модели»")

    # ---- сохранённый аудит и разбор замечаний экспертом ----
    saved_id = st.session_state.get("tab2_saved_id")
    audit_dir = WORK / "logs" / "audits" / str(saved_id)
    check(saved_id and (audit_dir / "meta.json").is_file() and (audit_dir / "document.docx").is_file()
          and (audit_dir / "raw.md.gz").is_file()
          and json.loads((audit_dir / "meta.json").read_text(encoding="utf-8"))["file"] == "ТЗ.docx",
          "аудит сохранён в LOG_DIR/audits: сведения, документ, сырые ответы")
    meta_ = json.loads((audit_dir / "meta.json").read_text(encoding="utf-8"))
    check(st.session_state.get("mint_profile") == "general" and ("selectbox", "🧭 Профиль:") in st.LOG
          and st.session_state.get("tab2_user_focus") == app.DEFAULT_AUDIT_FOCUS
          and meta_.get("profile") == "general",
          f"профиль по умолчанию — универсальный, он в сведениях аудита: {meta_.get('profile')}")
    rows_ = st.session_state.get("tab2_rows") or []
    check(len(rows_) >= 2 and ("expander", "✍️ Разбор замечаний экспертом и «Мои замечания»") in st.LOG
          and sum(1 for x in st.LOG if x == ("radio", "Отметка")) == min(len(rows_), 15)
          and ("expander", "📂 Мои аудиты") in st.LOG, f"под заключением — разбор замечаний: {len(rows_)}")
    kv0, kv1 = (f"rv::{saved_id}::{rows_[i]['key']}::v" for i in (0, 1))
    st.session_state[kv0] = "ok"
    st.session_state[f"rv::{saved_id}::{rows_[0]['key']}::c"] = "главное замечание"
    st.session_state[kv1] = "wrong"
    st.session_state[f"rv::{saved_id}::{rows_[1]['key']}::r"] = "present"
    st.SCRIPT["data_editor"][f"rv_added::{saved_id}"] = lambda data: list(data) + [
        {"section": "4. Восстановление", "severity": "КРИТИЧНЫЙ", "text": "Нет плана восстановления после сбоя"}]
    check(run_ui(uploads={"tab2_audit_uploader": Upl("ТЗ.docx", tz_bytes)}) == "ok", "отметки эксперта")
    st.SCRIPT["data_editor"].clear()
    rv_ = json.loads((audit_dir / "review.json").read_text(encoding="utf-8"))
    check(rv_["marks"].get(rows_[0]["key"]) == {"verdict": "ok", "comment": "главное замечание"}
          and rv_["marks"].get(rows_[1]["key"]) == {"verdict": "wrong", "reason": "present"}
          and rv_["added"] == [{"section": "4. Восстановление", "severity": "КРИТИЧНЫЙ",
                                "text": "Нет плана восстановления после сбоя"}]
          and rv_["progress"]["marked"] == 2, f"разбор сохранён сразу: {rv_.get('progress')}")
    check(any(x[0] == "selectbox" and x[1] == "Причина" for x in st.LOG)
          and {"📥 Мои замечания (.docx)", "📥 Мои замечания (.md)", "📤 Скачать разбор (.json)"}
          <= {x[1] for x in st.LOG if x[0] == "download_button"}, "у «неверно» — причина; «Мои замечания» и выгрузка")
    check(run_ui(uploads={"tab2_audit_uploader": Upl("ТЗ.docx", tz_bytes)}, buttons={f"rv_doc_btn::{saved_id}"}) == "ok"
          and ("download_button", "📥 Скачать документ (.docx)") in st.LOG, "документ с моими замечаниями собирается")
    with zipfile.ZipFile(io.BytesIO(st.session_state[f"rv_doc::{saved_id}"][1])) as z_:
        xml_ = "".join(z_.read(n).decode("utf-8", "replace") for n in z_.namelist() if n.endswith(".xml"))
    check("Нет плана восстановления после сбоя" in xml_ and rows_[0]["text"][:40] in xml_
          and rows_[1]["text"][:40] not in xml_, "в документе — ✅ и свои замечания, неверных нет")
    check(json.loads((audit_dir / "review.json").read_text(encoding="utf-8"))["added"] == rv_["added"],
          "своё замечание не задваивается после перерисовки")
    for k_ in [k_ for k_ in st.session_state if str(k_).startswith(("tab2_", "rv"))]:
        del st.session_state[k_]        # как новый сеанс: заключения на экране нет
    check(run_ui() == "ok" and ("markdown", "## 📄 Итоговое экспертное заключение") not in st.LOG
          and any(x[0] == "selectbox_options" and "ТЗ.docx · test-llm" in str(x[1]) and "разобрано 2 из" in str(x[1])
                  for x in st.LOG), "в «Моих аудитах» — аудит с прогрессом разбора")
    st.session_state["tab2_saved_pick"] = saved_id
    check(run_ui(buttons={"tab2_saved_open"}) == "ok" and st.session_state.get("tab2_reopened")
          and st.session_state.get("tab2_rows") == rows_ and ("markdown", "## 📄 Итоговое экспертное заключение") in st.LOG,
          "сохранённый аудит открывается без загрузки документа")
    check(run_ui() == "ok" and st.session_state.get(kv0) == "ok" and st.session_state.get(kv1) == "wrong"
          and "📥 Скачать документ с замечаниями (.docx)" in [x[1] for x in st.LOG if x[0] == "download_button"]
          and "📥 Скачать сырые ответы модели (.md)" in [x[1] for x in st.LOG if x[0] == "download_button"]
          and not any(x == ("button", "🔁 Перепроверить непроверенные разделы (1)") for x in st.LOG),
          "у открытого аудита — отметки, документ с замечаниями и сырые ответы; без перепроверки")
    check(run_ui(buttons={"tab2_export_all"}) == "ok" and st.session_state["_tab2_export_all"][2] == 1
          and ("download_button", "📥 Скачать разборы (1, .jsonl)") in st.LOG, "все разборы — одним файлом")
    one_ = json.loads(st.session_state["_tab2_export_all"][1].splitlines()[0])
    check(one_["format"] == "mint-review/1" and one_["audit"]["file"] == "ТЗ.docx" and one_["texts"]
          and one_["added"][0]["text"] == "Нет плана восстановления после сбоя", "в файле разбора — тексты по умолчанию")
    st.session_state["tab2_export_notext"] = True
    check(run_ui(buttons={"tab2_export_all"}) == "ok"
          and "ТЗ.docx" not in st.session_state["_tab2_export_all"][1]
          and rows_[0]["text"][:30] not in st.session_state["_tab2_export_all"][1], "флажок «без текстов»")
    st.session_state["tab2_special_focus"] = "шифрование персональных данных при передаче"
    mock_server.STATE["requests"].clear()
    check(run_ui(uploads={"tab2_audit_uploader": Upl("ТЗ.docx", tz_bytes)}, buttons={"tab2_run_btn"}) == "ok",
          "аудит с «особым вниманием»")
    chats_ = [pl_ for m_, path_, pl_ in mock_server.STATE["requests"] if path_.endswith("/chat/completions") and pl_]
    map_ = [c for c in chats_ if "ОБЯЗАТЕЛЬНЫЕ ФОКУСЫ" in c["messages"][0]["content"]]
    focus_ = [c for c in chats_ if "Пользователь отметил пункт" in c["messages"][0]["content"]]
    check(len(focus_) == 1 and "шифрование персональных данных при передаче" in focus_[0]["messages"][-1]["content"],
          "на пункт особого внимания — один отдельный ответ по документу")
    rep_ = st.session_state.get("tab2_final_report_data", "")
    check("ОТВЕТЫ НА ПУНКТЫ ОСОБОГО ВНИМАНИЯ" in rep_ and "\nчастично — по пункту «шифрование" in rep_
          and "платформу интеграции; и кто его получает дальше" in rep_,
          "в заключении — ответ на пункт и замечание с подпунктом")
    check(any(x[0] == "expander" and "Особое внимание: «шифрование" in str(x[1]) for x in st.LOG),
          "ответ на пункт особого внимания — в логе проверки")
    crit_ = [c for c in chats_ if "Старший Контролер" in c["messages"][0]["content"]]
    check(map_ and all("ОСОБОЕ ВНИМАНИЕ" in c["messages"][0]["content"] and "шифрование персональных" in
                       c["messages"][0]["content"] for c in map_), "особое внимание — в промпте проверки")
    check(crit_ and all("ОСОБОЕ ВНИМАНИЕ ПОЛЬЗОВАТЕЛЯ" in c["messages"][-1]["content"] for c in crit_),
          "особое внимание — в самокритике")
    check("**Особое внимание**: шифрование" in st.session_state.get("tab2_final_report_data", ""),
          "особое внимание — в заключении")
    check("Особое внимание:** шифрование" in st.session_state.get("tab2_audit_log_md", ""), "особое внимание — в логе")
    emb_focus = [pl_ for m_, path_, pl_ in mock_server.STATE["requests"] if path_.endswith("/embeddings") and pl_
                 and any("шифрование персональных" in str(x) for x in (pl_.get("input") if isinstance(pl_.get("input"), list) else [pl_.get("input")]))]
    check(len(emb_focus) == 2, f"нормы по особому вниманию ищутся не по разделам: для норм к разделам и для ответа на "
                               f"пункт — {len(emb_focus)}")
    st.session_state["tab2_special_focus"] = ""
    st.session_state["tab2_whole"] = True
    check(run_ui(uploads={"tab2_audit_uploader": Upl("ТЗ.docx", tz_bytes)}, buttons={"tab2_run_btn"}) == "ok"
          and any(x[0] == "write" and "проверяется одним запросом" in str(x[1]) for x in st.LOG)
          and not any(x[0] == "write" and "Документ разбит" in str(x[1]) for x in st.LOG),
          "весь документ одним запросом: сообщение без «разбит на части до N симв.»",
          str([x for x in st.LOG if x[0] in ("write", "warning") and "📜" in str(x[1])]))
    st.session_state["tab2_whole"] = False
    manual_ = st.session_state["mint_manual"]
    manual_.update(audit_passes=2, audit_summary_max=0)
    check(run_ui(uploads={"tab2_audit_uploader": Upl("ТЗ.docx", tz_bytes)}, buttons={"tab2_run_btn"}) == "ok"
          and len((st.session_state["tab2_audit_result"]["state"].get("passes") or [])) == 1
          and any(x[0] == "write" and "🔂 Проход 2 из 2" in str(x[1]) for x in st.LOG)
          and "ГЛАВНОЕ" not in st.session_state["tab2_final_report_data"],
          "настройки из блока аудита действуют: два прохода, без «Главного»")
    manual_.pop("audit_passes")
    manual_.pop("audit_summary_max")
    st.session_state["tab2_reference"] = "Не указан алгоритм шифрования при передаче\nНет резервного копирования журналов"
    check(run_ui(uploads={"tab2_audit_uploader": Upl("ТЗ.docx", tz_bytes)}, buttons={"tab2_reference_btn"}) == "ok",
          "сверка с эталоном через интерфейс")
    ref_ = st.session_state.get("tab2_reference_result")
    check(ref_ and ref_[0] == "ТЗ.docx" and ref_[2]["total"] == 2 and ref_[2]["found"] == 1
          and any(x[0] == "markdown" and "**найдено:** 1 из 2" in str(x[1]) for x in st.LOG)
          and "📥 Скачать сверку (.md)" in [x[1] for x in st.LOG if x[0] == "download_button"],
          f"итог сверки на странице и для скачивания: {ref_ and ref_[2]['found']}")
    check(any(x[0] == "expander" and "Глубокая проверка: вопросов" in str(x[1]) for x in st.LOG)
          and any(x == ("checkbox", "🔬 Глубокая проверка (процесс по шагам, противоречия, цифры, опечатки)") for x in st.LOG),
          "глубокая проверка: переключатель и её лог на странице")
    st.session_state["tab3_remarks_area"] = "- [КРИТИЧНЫЙ] Не указан алгоритм шифрования ПДн"
    check(run_ui(uploads={"tab3_v2_uploader": Upl("ТЗ_v2.docx", v2)}, buttons={"tab3_run_btn"}) == "ok",
          "контроль исправлений через интерфейс")
    check(run_ui(uploads={"tab3_v2_uploader": Upl("ТЗ_v2.docx", v2)}) == "ok"
          and ("markdown", "## 📊 Результаты проверки") in st.LOG, "результат контр-проверки не пропал после перерисовки")
    rid = st.session_state["tab3_result"]["id"]
    check(run_ui(uploads={"tab3_v2_uploader": Upl("ТЗ_v2.docx", v2)}, buttons={f"fb_down_review_{rid}"}) == "ok",
          "оценка контр-проверки")
    rr_ = json.loads(fb_file.read_text(encoding="utf-8").splitlines()[-1])
    check(rr_.get("kind") == "review" and rr_.get("statuses") and "Не указан алгоритм" in rr_.get("remarks", "")
          and rr_.get("prompts") == app.prompts_version(), "в отзыве о контр-проверке — замечания, статусы, версия промптов")
    md_doc = ("# Требования к ТЗ\n\n## 1. Защита данных\n\nПерсональные данные передаются по открытому каналу.\n\n"
              "## 2. Надёжность\n\nРезервное копирование не описано.\n").encode("utf-8")
    check(run_ui(buttons={"kb_reset_btn"}) == "rerun", "«загрузить новые документы»")
    check(run_ui(uploads={"kb_uploader": [Upl("Памятка.md", md_doc)]}, buttons={"kb_index_all_btn"}) == "rerun"
          and all(ok for _, ok, _ in st.session_state["_kb_batch_results"])
          and "Структура: заголовки Markdown" in st.session_state["_kb_batch_results"][0][2],
          f"загрузка .md в базу через интерфейс: {st.session_state.get('_kb_batch_results')}")
    run_ui(buttons={"kb_reset_btn"})
    check(run_ui(uploads={"kb_uploader": [Upl("старый.doc", b"x")]}, buttons={"kb_index_all_btn"}) == "rerun"
          and not st.session_state["_kb_batch_results"][0][1]
          and "не поддерживается" in st.session_state["_kb_batch_results"][0][2], "неподдерживаемый файл — понятная ошибка")
    st.session_state.pop("tab2_final_report_data", None)
    check(run_ui(uploads={"tab2_audit_uploader": Upl("ТЗ.md", md_doc)}, buttons={"tab2_run_btn"}) == "ok"
          and st.session_state.get("tab2_final_report_data") and st.session_state.get("tab2_docx_review") is None,
          "аудит .md: заключение есть, документ Word не собирается")
    check(run_ui(uploads={"tab2_audit_uploader": Upl("ТЗ.md", md_doc)}) == "ok"
          and any(x[0] == "caption" and "только для .docx" in str(x[1]) for x in st.LOG)
          and "📥 Скачать документ с замечаниями (.docx)" not in [x[1] for x in st.LOG if x[0] == "download_button"],
          "для .md — подпись вместо скачивания Word")
    mock_server.STATE["fail_words"] = {"ЗАМЕДЛЕНИЕ": 2}         # шлюз отвечает 503 на обе попытки
    slow_doc = "# Требования\n\nЗАМЕДЛЕНИЕ. Персональные данные передаются по открытому каналу.\n".encode()
    check(run_ui(uploads={"tab2_audit_uploader": Upl("Медленный.md", slow_doc)}, buttons={"tab2_run_btn"}) == "ok"
          and st.session_state["tab2_audit_result"]["failed_sections"]
          and not any("проверяю их ещё раз" in str(x[1]) for x in st.LOG),
          "аудит: единственный раздел не проверен — автоповтора нет (шлюз, похоже, недоступен)")
    check(run_ui(uploads={"tab2_audit_uploader": Upl("Медленный.md", slow_doc)}) == "ok"
          and ("button", "🔁 Перепроверить непроверенные разделы (1)") in st.LOG, "кнопка «Перепроверить непроверенные»")
    slow_id = st.session_state.get("tab2_saved_id")
    check(run_ui(uploads={"tab2_audit_uploader": Upl("Медленный.md", slow_doc)}, buttons={"tab2_recheck_btn"}) == "ok"
          and not st.session_state["tab2_audit_result"]["failed_sections"]
          and "НЕ УДАЛОСЬ ПРОВЕРИТЬ" not in st.session_state["tab2_final_report_data"]
          and st.session_state["tab2_audit_log"][0]["attempts"] == 2
          and "с 2-й попытки" in st.session_state["tab2_audit_log_md"], "перепроверка: раздел проверен, заключение обновлено")
    check(slow_id and st.session_state.get("tab2_saved_id") == slow_id
          and json.loads((WORK / "logs" / "audits" / slow_id / "meta.json").read_text(encoding="utf-8"))["failed"] == 0,
          "перепроверка сохраняется в ту же папку аудита")
    mock_server.STATE["fail_words"] = {}
    check(run_ui(uploads={"tab3_v2_uploader": Upl("ТЗ_v2.txt", "1. Защита данных\nДанные шифруются по ГОСТ.\n"
                                                                .encode("cp1251"))}, buttons={"tab3_run_btn"}) == "ok"
          and st.session_state["tab3_result"]["rows"], "контроль исправлений по новой редакции в .txt")
    check(run_ui(buttons={"reindex_btn"}) == "rerun", "переиндексация через интерфейс")
    check(run_ui() == "ok", "страница после переиндексации")
except Exception:
    traceback.print_exc()
    FAILS.append("исключение в интерфейсе")

print("21. [сквозной] Проверка качества: вопросы, правка в таблице, прогон, рекомендации, прошлые проверки")
try:
    from mint.ui import quality
    st.session_state.pop("qa", None)
    check(run_ui(buttons={"qa_open_btn"}) == "ok" and st.session_state.get("qa_open"),
          "кнопка на боковой панели открывает проверку")
    check(("button", "✨ Составить вопросы") in st.LOG
          and ("subheader", "Консультация по накопленным знаниям") not in st.LOG, "проверка — вместо вкладок")
    st.session_state["qa_n"] = 5
    check(run_ui(buttons={"qa_gen_btn"}) == "rerun" and st.session_state["qa"]["stage"] == "review",
          "вопросы составлены → шаг «Вопросы»")
    qa_ = st.session_state["qa"]
    gen_rows = qa_["rows"]
    check(gen_rows and all(r["file"] and r["id"] in qa_["fragments"] for r in gen_rows),
          f"в таблице {len(gen_rows)} вопросов с документом и фрагментом")
    f0, c0 = gen_rows[0]["file"], gen_rows[0]["contains"]

    def edit_(data):
        data = [dict(r) for r in data]
        data[0]["use"] = False
        data += [{"use": True, "question": "Свой вопрос про то же место", "file": f0, "section": "", "contains": c0,
                  "note": float("nan"), "id": None},
                 {"use": True, "question": "Какой курс валют на завтра?", "file": quality.tr(app.NO_ANSWER),
                  "section": "", "contains": "", "note": "", "id": None},
                 {"use": True, "question": "Вопрос с фразой мимо документа", "file": f0, "section": "",
                  "contains": "этой фразы нет ни в одном документе", "note": "", "id": None},
                 {"use": None, "question": "Вопрос без документа", "file": None, "section": None, "contains": None,
                  "note": None, "id": None}]
        return data

    ekey = f"qa_editor_{qa_['ver']}"
    st.SCRIPT["data_editor"] = {ekey: edit_}
    check(run_ui() == "ok", "таблица вопросов с правками")
    fix_ = [x for x in st.LOG if x[0] == "expander" and str(x[1]).startswith("Что поправить")]
    check(fix_ and "(2)" in fix_[0][1], f"подсвечены строка без документа и фраза мимо документа: {fix_}")
    check(("data_editor", ekey) in st.LOG and any("mint-tally" in str(x[1]) for x in st.LOG if x[0] == "markdown"),
          "таблица и счётчики вопросов на странице")
    check(any(x[0] == "expander" and "Фрагмент, по которому" in str(x[1]) for x in st.LOG), "просмотр фрагмента вопроса")
    check(any(x[0] == "download_button" and str(x[1]).startswith("⬇ Документ целиком") for x in st.LOG),
          "из фрагмента можно скачать документ целиком")
    check(run_ui(buttons={"qa_run_btn"}) == "rerun" and st.session_state["qa"]["stage"] == "results",
          "прогон → шаг «Результаты»")
    run_ = st.session_state["qa"].get("run")
    qs_ = run_.questions if run_ else []
    check(len(qs_) == len(gen_rows) - 1 + 3 and sum(1 for q in qs_ if not q.answerable) == 1,
          f"в прогоне — отмеченные строки с документом: {len(qs_)}")
    saved_ = Path(st.session_state["qa"].get("saved") or "")
    check((saved_ / "report.md").exists() and (saved_ / "questions.csv").exists()
          and json.loads((saved_ / "summary.json").read_text(encoding="utf-8")).get("current") == run_.current,
          f"проверка сохранена в LOG_DIR/eval: {saved_}")
    check(saved_.parent == WORK / "logs" / "eval", "папка проверки — в LOG_DIR/eval")
    check(run_ui() == "ok", "страница результатов")
    md_ = " ".join(str(x[1]) for x in st.LOG if x[0] == "markdown")
    check("mint-scorecard" in md_ and "mint-ladder" in md_ and "mint-steps" in md_, "итоги, сравнение настроек и шаги")
    dls_ = [x[1] for x in st.LOG if x[0] == "download_button"]
    check({"⬇ Отчёт (.md)", "⬇ По вопросам (.csv)", "⬇ Вопросы (.csv)", "⬇ summary.json"} <= set(dls_),
          f"отчёты для скачивания: {dls_}")
    check(("dataframe", len(qs_)) in st.LOG, "таблица по вопросам")
    frag_ = quality.highlight("# 1. Журналы\n\nЖурналы хранятся не менее\nтрёх лет. <b>", "не менее трёх лет")
    check("\n" not in frag_ and '<b class="mint-fragment__h">1. Журналы</b>' in frag_
          and "<mark>не менее<br>трёх лет</mark>" in frag_ and "&lt;b&gt;" in frag_,
          f"фрагмент: без пустых строк (они рвут HTML в Markdown), заголовок, цитата через перенос: {frag_}")
    long_ = quality.highlight("а " * 2000 + "искомая фраза" + " б" * 2000, "искомая фраза")
    check(long_.startswith("… ") and "<mark>искомая фраза</mark>" in long_ and len(long_) < 2000,
          "длинный фрагмент обрезается вокруг цитаты")
    check('aria-current="step"' in quality.steps_html("review") and quality.steps_html("review").count("is-done") == 1,
          "шаги: на таблице вопросов отмечен второй шаг")
    lad_ = quality.ladder_html(run_)
    check(lad_.count("<li") == len(run_.configs) + 1 and "сейчас" in lad_ and "★" in lad_, "сравнение настроек: строка на настройку")

    orig_rec_ = quality.ev.recommendations
    quality.ev.recommendations = lambda r: {"weight_rerank": 1.5, "low_conf_rerank": 0.2}
    check(run_ui() == "ok" and ("button", "Применить в этой сессии") in st.LOG
          and ("code", "search:\n  weight_rerank: 1.5\nconfidence:\n  low_conf_rerank: 0.2") in st.LOG,
          "рекомендации: кнопка и фрагмент config.yaml по разделам")
    st.session_state["x_weight_rerank_b1.0"] = 1.0
    quality._apply({"weight_rerank": 1.5, "low_conf_rerank": 0.2})
    check(run_ui() == "ok" and ("button", "✓ Применено в этой сессии") in st.LOG, "применённые рекомендации отмечены")
    quality.ev.recommendations = orig_rec_
    st.session_state["mint_manual"].pop("low_conf_rerank", None)
    quality._apply({"weight_rerank": 1.5})
    check(st.session_state["mint_manual"].get("weight_rerank") == 1.5
          and "x_weight_rerank_b1.0" not in st.session_state, "рекомендация применяется к сессии")
    check(run_ui() == "ok" and any(x[0] == "toast" and "weight_rerank = 1.5" in str(x[1]) for x in st.LOG),
          "сообщение о применённой рекомендации")
    st.session_state["mint_manual"].pop("weight_rerank", None)

    check(run_ui(buttons={"qa_edit_btn"}) == "rerun" and st.session_state["qa"]["stage"] == "review"
          and any(r.get("question") == "Свой вопрос про то же место" for r in st.session_state["qa"]["rows"]),
          "«Изменить вопросы» возвращает таблицу со своими вопросами")
    st.session_state["qa"]["opts"] = {"compare": False, "answers": True, "multi": False}
    st.SCRIPT["data_editor"] = {}
    check(run_ui(buttons={"qa_run_btn"}) == "rerun", "повторный прогон с ответами чата и без сравнения")
    run2_ = st.session_state["qa"]["run"]
    check(run2_.answers and len(run2_.answers["rows"]) == len(run2_.questions) and len(run2_.configs) <= 4,
          f"ответы чата проверены, настроек меньше: {len(run2_.configs)}")
    check((Path(st.session_state["qa"]["saved"]) / "answers.csv").exists(), "answers.csv сохранён")
    check(st.session_state["qa"].get("prev", {}).get("current") == run_.current, "есть прошлая проверка для сравнения")
    check(run_ui() == "ok" and any("mint-answers" in str(x[1]) for x in st.LOG if x[0] == "markdown"),
          "результаты с ответами чата")
    log_txt_ = " ".join(str(x[1]) for x in st.LOG)
    check(f"Промпты: {app.prompts_version()}" in log_txt_ and "Это первая проверка ответов" in log_txt_
          and run2_.answers["summary"]["prompts"] == app.prompts_version(),
          "у ответов — версия промптов; первая проверка ответов — сравнивать не с чем")
    st.session_state["qa_show"] = "bad"
    check(run_ui() == "ok" and ("radio", "Показать") in st.LOG, "фильтр «ответы с ошибками»")
    bad_rows_ = quality.result_rows(run2_, "bad")
    check(bad_rows_ and all(r["Ответ: итог"].startswith("✗") for r in bad_rows_)
          and len(bad_rows_) == sum(1 for q in run2_.questions if not run2_.answers["rows"][q.n][0]["ok"]),
          "в таблице — ответы с ошибками и причина", str([r["Ответ: итог"] for r in bad_rows_]))
    st.session_state.pop("qa_show", None)
    check(run_ui(buttons={"qa_edit_btn"}) == "rerun", "к вопросам — для нового прогона")
    st.session_state["qa"]["opts"] = {"compare": False, "answers": True, "multi": False, "repeat": 2}
    check(run_ui(buttons={"qa_run_btn"}) == "rerun", "прогон ответов чата с повторами")
    run3_ = st.session_state["qa"]["run"]
    check(all(len(v) == 2 for v in run3_.answers["rows"].values()) and run3_.answers["summary"]["repeat"] == 2
          and st.session_state["qa"].get("prev_answers", {}).get("repeat") == 1,
          "каждый вопрос задан дважды; есть прошлая проверка ответов")
    check(run_ui() == "ok" and any("С прошлой проверкой ответов" in str(x[1]) and "повторов: 1 → 2" in str(x[1])
                                   for x in st.LOG if x[0] == "markdown"),
          "сравнение ответов с прошлой проверкой: что изменилось и вердикт")
    rep_ = st.session_state["qa"].get("report", "")
    check("### Сравнение ответов с прошлой проверкой" in rep_, "сравнение ответов — и в отчёте")
    check(quality.answer_cell([{"ok": True, "reason": None}, {"ok": False, "reason": "no_ref"}])
          == "1 из 2 ✓ · модель: нет ссылки на нужный документ", "[unit] итог ответов на вопрос с повторами")

    check(run_ui(buttons={"qa_new_btn"}) == "rerun" and st.session_state["qa"]["stage"] == "setup", "новая проверка")
    check(run_ui(buttons={"qa_prev_btn"}) == "rerun" and len(st.session_state["qa"]["rows"]) == len(run2_.questions),
          "вопросы прошлой проверки — без CSV")
    check(run_ui(buttons={"qa_close_btn"}) == "ok" and not st.session_state.get("qa_open"), "закрыть проверку")
    check(run_ui() == "ok" and ("subheader", "Консультация по накопленным знаниям") in st.LOG, "вкладки снова на месте")
except Exception:
    traceback.print_exc()
    FAILS.append("исключение в проверке качества")

print("22. [сквозной] Промпты: правка, проверка формата, сеанс, prompts.yaml, проверка правки")
try:
    from mint import prompts as P_
    ppath_ = WORK / "prompts.yaml"
    check(run_ui(buttons={"pr_open_btn"}) == "ok" and st.session_state.get("prompts_open")
          and ("subheader", "📝 Промпты") in st.LOG and ("subheader", "Консультация по накопленным знаниям") not in st.LOG,
          "кнопка на боковой панели открывает страницу промптов — вместо вкладок")
    st.session_state["pr_key"] = "chat_system"
    check(run_ui() == "ok" and ("text_area", "Текст") in st.LOG, "выбран системный промпт чата, текст в поле")
    tkey_ = f"pr_text_ru_chat_system_{st.session_state.get('prompts_ver', 0)}"
    base_ = P_.default_text("chat_system", "ru")
    v0_ = P_.prompts_version("ru")
    st.session_state[tkey_] = base_.replace("[Источник: имя_файла, раздел]", "(источник)")
    check(run_ui(buttons={"pr_save_btn", "pr_try_btn"}) == "ok"
          and any(x[0] == "error" and "нет обязательной части формата" in str(x[1]) for x in st.LOG)
          and ("button_disabled", "Сохранить") in st.LOG and not ppath_.exists() and P_.prompts_version("ru") == v0_,
          "правка без формата ссылки [Источник: …] не применяется и не сохраняется — страница говорит, чего нет")
    edited_ = base_ + "\n9. Отвечай кратко."
    st.session_state[tkey_] = edited_
    check(run_ui(buttons={"pr_try_btn"}) == "rerun" and st.session_state["prompt_session"]["ru"]["chat_system"] == edited_,
          "«Попробовать в сеансе» — правка в сеансе")
    check(run_ui() == "ok" and P_.custom_status("ru") == {"chat_system": "session"} and P_.prompts_version("ru") != v0_,
          "правка сеанса действует, версия промптов другая")
    v1_ = P_.prompts_version("ru")
    check(run_ui(buttons={"pr_close_btn"}) == "ok" and run_ui(chat="Какие требования к СЗИ?") == "ok",
          "закрыть страницу, вопрос в чате")
    sys_ = mock_server.STATE["last_chat"]["messages"][0]["content"]
    check("9. Отвечай кратко." in sys_ and st.session_state["messages"][-1]["prompts"] == v1_,
          "чат идёт с правкой сеанса, в ответе — её версия промптов")
    check(run_ui(buttons={"pr_open_btn"}) == "ok" and run_ui(buttons={"pr_save_btn"}) == "rerun" and ppath_.exists()
          and "Отвечай кратко" in ppath_.read_text(encoding="utf-8")
          and not st.session_state["prompt_session"]["ru"].get("chat_system"),
          "«Сохранить» — правка в prompts.yaml, из сеанса убрана")
    check(run_ui() == "ok" and P_.custom_status("ru") == {"chat_system": "file"} and P_.prompts_version("ru") == v1_,
          "сохранённая правка действует, версия та же")
    raw_ = ppath_.read_text(encoding="utf-8")
    ppath_.write_text(raw_.replace(f"base: {P_._short_hash(base_)}", "base: 00000000"), encoding="utf-8")
    check(run_ui() == "ok" and any(x[0] == "warning" and "изменился после того" in str(x[1]) for x in st.LOG),
          "если умолчание обновилось после правки — предупреждение")
    check(run_ui(buttons={"pr_reset_btn"}) == "rerun" and P_.custom_status("ru") == {}
          and "chat_system" not in ppath_.read_text(encoding="utf-8") and P_.prompts_version("ru") == v0_,
          "«Вернуть умолчание» — правки нет ни в сеансе, ни в файле")
    st.session_state["pr_key"] = "audit_system"
    check(run_ui() == "ok" and ("button_disabled", "Проверить правку") in st.LOG,
          "аудит проверкой ответов чата не измеряется — кнопка «Проверить правку» неактивна")
    st.session_state["pr_key"] = "chat_no_info"
    run_ui()
    nkey_ = f"pr_text_ru_chat_no_info_{st.session_state.get('prompts_ver', 0)}"
    st.session_state[nkey_] = "Сведений в регламентах нет"
    check(run_ui(buttons={"pr_check_btn"}) == "rerun" and st.session_state.get("qa_open")
          and not st.session_state.get("prompts_open") and st.session_state["qa"]["stage"] == "review"
          and st.session_state["qa"]["opts"] == {"compare": False, "answers": True, "multi": False, "repeat": 2},
          "«Проверить правку» — проверка качества с ответами чата на вопросах прошлой проверки")
    v2_ = P_.prompts_version("ru")
    st.SCRIPT["data_editor"] = {}
    check(run_ui(buttons={"qa_run_btn"}) == "rerun", "прогон проверки правки")
    a_ = st.session_state["qa"]["run"].answers["summary"]
    check(a_["prompts"] == v2_ and a_["prompts_custom"] == {"chat_no_info": "session"}
          and app.is_no_info_answer("Сведений в регламентах нет.") and st.session_state["qa"].get("prev_answers"),
          "в итогах — версия с правкой и что правлено; «информации нет» узнаётся по новой фразе; есть с чем сравнить")
    check(run_ui() == "ok" and any("версия промптов:" in str(x[1]) and v2_ in str(x[1]) for x in st.LOG
                                   if x[0] == "markdown"),
          "в сравнении с прошлой проверкой видна смена версии промптов")
    check(any(x[0] == "caption" and "свои правки: Фраза «информации нет» (в сеансе)" in str(x[1]) for x in st.LOG)
          and "Свои правки промптов: Фраза «информации нет» (в сеансе)." in st.session_state["qa"].get("report", ""),
          "в итогах и отчёте названо, какой промпт изменён")
    st.session_state["prompt_session"] = {}
    st.session_state["qa_open"] = False
    check(run_ui() == "ok" and P_.prompts_version("ru") == v0_, "правки сеанса убраны — снова умолчания")
except Exception:
    traceback.print_exc()
    FAILS.append("исключение на странице промптов")

print("19в. [сквозной] Черновик карточки системы по базе")
try:
    check(any(x == ("button", "🗂️ Собрать черновик карточки системы по базе") for x in st.LOG),
          "вход — в блоке «Для аудита» боковой панели")
    check(run_ui(buttons={"cards_open_btn"}) == "ok" and st.session_state.get("cards_open")
          and ("subheader", "🗂️ Черновик карточки системы") in st.LOG
          and ("button_disabled", "🗂️ Собрать черновик") in st.LOG
          and not any(x[0] == "file_uploader" and "согласование" in str(x[1]) for x in st.LOG),
          "страница карточки вместо вкладок; без названия системы собрать нельзя")
    st.session_state["cards_system"] = "Регламент ИБ"
    check(run_ui(buttons={"cards_run_btn"}) == "ok" and st.session_state.get("cards_result")
          and "📥 Скачать черновик карточки (.md)" in [x[1] for x in st.LOG if x[0] == "download_button"]
          and st.session_state["cards_result"]["markdown"].startswith("# Карточка системы «Регламент ИБ»"),
          "черновик собран и скачивается")
    check(run_ui(buttons={"qa_open_btn"}) == "ok" and st.session_state.get("qa_open")
          and not st.session_state.get("cards_open"), "открыть проверку качества — карточка закрывается")
    st.session_state["qa_open"] = False
    check(run_ui(buttons={"cards_open_btn"}) == "ok" and run_ui(buttons={"cards_close_btn"}) == "ok"
          and not st.session_state.get("cards_open"), "закрыть страницу карточки")
    # профиль сеанса (боковая панель): «Архитектура и ИТ-решения» — промпты MINT до профилей слово в слово
    st.session_state["mint_profile"] = "architecture"
    check(run_ui() == "ok" and app.get_profile() == "architecture" and app.prompts_version("ru") == "ru-8fd10871"
          and "Старший Контролер Архитектурного Комитета" in app.critique_system("ru"),
          f"профиль архитектуры в сеансе — набор промптов до профилей: {app.prompts_version('ru')}")
    st.session_state["mint_profile"] = "general"
    check(run_ui() == "ok" and app.get_profile() == "general" and app.prompts_version("ru") != "ru-8fd10871",
          "обратно универсальный профиль")
except Exception:
    traceback.print_exc()
    FAILS.append("исключение на странице карточки")

print("20. [unit] Оформление: штамп заключения, счётчики, источники, пустой чат")
try:
    from mint.ui import theme
    crit = [{"severity": "КРИТИЧНЫЙ"}, {"severity": "НИЗКИЙ"}]
    check(theme.audit_verdict(crit)[:2] == ("rework", "Требует доработки"), "критичное замечание → «Требует доработки»")
    check(theme.audit_verdict([{"severity": "СРЕДНИЙ"}])[0] == "remarks", "среднее → «Есть замечания»")
    check(theme.audit_verdict([], n_failed=1)[0] == "incomplete", "сбой разделов без замечаний — не «Согласовано»")
    check(theme.audit_verdict([])[0] == "approved", "без замечаний и сбоев → «Согласовано»")
    vh = theme.audit_verdict_html(crit, 1, 4)
    check("--mint-n:1" in vh and "mint-verdict--rework" in vh and "1 из 4" in vh and 'class="mint-sr">1<' in vh,
          "штамп: счётчики, вид, примечание о сбое, число для экранного диктора")
    rows_ = [{"status": "✅ Исправлено"}, {"status": "⚠️ Частично"}, {"status": "❌ Не исправлено"}, {"status": "❓ x"}]
    check(theme.review_tally(rows_) == {"fixed": 1, "partial": 1, "not_fixed": 1, "unverified": 1}, "счётчики контроля")
    check("mint-verdict--approved" in theme.review_tally_html([{"status": "✅"}]), "всё исправлено → зелёный штамп")
    lab = "📄 Регламент.docx › 1.2 Журналы (итог 0.0443 · dense #1 · BM25 #1 · rerank 0.80 ✓)"
    check(theme.split_source(lab) == ("Регламент.docx", "1.2 Журналы", "итог 0.0443 · dense #1 · BM25 #1 · rerank 0.80 ✓"),
          f"подпись источника разбирается: {theme.split_source(lab)}")
    check(theme.split_source("📄 a (b).docx") == ("a (b).docx", "", ""), "скобки в имени файла — не подробности")
    sh = theme.sources_html([lab, "📄 <b>x</b>.docx"])
    check("&lt;b&gt;" in sh and "<b>x" not in sh and "--i:1" in sh, "список источников: экранирование, задержки")
    check("пока пуста" in theme.empty_chat_html([]), "пустая база — подсказка, где добавить документы")
    many = theme.empty_chat_html([f"d{i}.docx" for i in range(9)])
    check(many.count("<li") == 7 and "и ещё 3" in many and "Документов в базе: 9" in many, "пустой чат: 6 документов и «и ещё 3»")
except Exception:
    traceback.print_exc()
    FAILS.append("исключение в оформлении")

srv.shutdown()
os.chdir(harness.ROOT)
shutil.rmtree(WORK, ignore_errors=True)
check.finish()
