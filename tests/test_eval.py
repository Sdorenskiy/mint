"""Оценка поиска (eval_search.py) на тестовой базе: генерация черновика вопросов, прогон, сравнение прогонов,
файл вопросов в Windows-1251.

Запуск из корня проекта:  python tests/test_eval.py            (с реранкером)
                          python tests/test_eval.py --no-rerank
"""
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import harness  # noqa: E402
import mock_server  # noqa: E402
import docx as pydocx  # noqa: E402

rerank = "--no-rerank" not in sys.argv
W = harness.make_project(prefix="mint_eval_", extra_files=("eval_search.py", "eval"))
srv = mock_server.start()
BASE = f"http://127.0.0.1:{srv.server_address[1]}/v1"
os.environ.update({
    "LLM_API_URL": BASE, "EMBED_API_URL": BASE + "/embeddings", "RERANK_API_URL": BASE + "/rerank" if rerank else "",
    "LLM_MODEL": "test-llm", "EMBED_MODEL": "emb-model", "RERANK_MODEL": "rr-model",
    "KB_DIR": str(W / "kb"), "LOG_DIR": str(W / "logs"), "MINT_CONFIG": str(W / "none.yaml"),
    "GLOSSARY_PATH": str(W / "none.json"), "MINT_LANG": "ru",
})
os.chdir(W)
app = harness.load_mint(W, include_ui=False)
check = harness.Checker()

cfg, _ = app.load_settings()
cfg, _ = app.resolve_settings(cfg, app.resolve_model_limits(cfg, {}, None, {}), {})


def make(path, blocks):
    d = pydocx.Document()
    for kind, val in blocks:
        if kind == "h1":
            d.add_heading(val, 1)
        elif kind == "h2":
            d.add_heading(val, 2)
        else:
            d.add_paragraph(val)
    d.save(path)
    return path.read_bytes()


long_par = ("Журналы событий безопасности хранятся не менее трёх лет. Доступ к журналам разграничивается по ролям. "
            "Администратор обязан ежедневно проверять целостность журналов. ") * 12
make(W / "a.docx", [
    ("h1", "Перечень сокращений"), ("p", "СЗИ – средства защиты информации;"), ("p", "ПДн — персональные данные;"),
    ("h1", "1. Общие положения"), ("p", "Регламент определяет требования к системе управления базами данных (СУБД)."),
    ("h2", "1.1 Шифрование"),
    ("p", "При передаче ПДн между подсистемами применяются средства криптографической защиты информации, "
          "сертифицированные по ГОСТ."),
    ("h2", "1.2 Журналирование"), ("p", long_par)])
make(W / "b.docx", [
    ("h1", "Показатели надёжности"),
    ("p", "Время восстановления сервиса не превышает четырёх часов. Доступность системы — 99,9 % в месяц."),
    ("p", "Резервное копирование выполняется ежедневно в центре обработки данных.")])
for name, f in (("Регламент_ИБ.docx", "a.docx"), ("Регламент_SLA.docx", "b.docx")):
    ok, msg = app.process_and_add_to_kb(name, (W / f).read_bytes(), cfg)
    check(ok, f"документ {name} загружен", msg)

child_env = {**os.environ, "PYTHONPATH": harness.child_pythonpath(W)}
eval_py = W / "eval_search.py"


def run(args, label):
    r = subprocess.run([sys.executable, str(eval_py), *args], env=child_env, capture_output=True, text=True,
                       timeout=600, cwd=W)
    err = "\n".join(line for line in r.stderr.splitlines() if "MINT_CORE" not in line)
    check(r.returncode == 0, label, (r.stdout[-1500:] + "\n" + err[-1500:]).strip())
    return r


run(["--generate", "6", "--out", str(W / "draft.csv")], "черновик вопросов по базе")
draft = (W / "draft.csv").read_text(encoding="utf-8-sig") if (W / "draft.csv").exists() else ""
check(len(draft.splitlines()) >= 3, f"в черновике есть вопросы ({max(len(draft.splitlines()) - 1, 0)})", draft[:500])

run([str(W / "eval" / "questions.example.csv"), "--out", str(W / "run1")], "прогон по примеру вопросов")
report = (W / "run1" / "report.md").read_text(encoding="utf-8") if (W / "run1" / "report.md").exists() else ""
check("hit@1" in report and (W / "run1" / "details.csv").exists() and (W / "run1" / "summary.json").exists(),
      "отчёт, подробности и итоги записаны", report[:500])
if rerank:
    check("реранкер" in report.lower(), "в отчёте есть конфигурации с реранкером")
    check("| low_conf_rerank | Нет ответа распознан | Ложная тревога |" in report
          and "Лучший балл реранкера: на вопросах с ответом" in report,
          "пороги уверенности подбираются по баллу реранкера", report[report.find("## Пороги"):][:800])
else:
    check("| low_conf_lexical | Нет ответа распознан | Ложная тревога |" in report and "Реранкера нет" in report,
          "без реранкера пороги подбираются по доле слов вопроса", report[report.find("## Пороги"):][:800])

import csv  # noqa: E402
check((W / "run1" / "details.csv").exists() and b"\r\r\n" not in (W / "run1" / "details.csv").read_bytes(),
      "details.csv без двойных переводов строк (на Windows — без пустых строк в Excel)")
details = list(csv.reader((W / "run1" / "details.csv").open(encoding="utf-8-sig"), delimiter=";")) \
    if (W / "run1" / "details.csv").exists() else []
check(details and all(cell.strip() for row in details for cell in row), "в details.csv нет пустых ячеек",
      str([row for row in details if not all(c.strip() for c in row)][:3]))
no_answer_rows = [row for row in details[1:] if row[2].startswith("—")]
check(no_answer_rows and all(c in ("✓ сигнал", "✗ без сигнала") for row in no_answer_rows for c in row[3:-5]),
      "у вопросов без ответа в колонках «Место» — сработал ли сигнал", str(no_answer_rows[:1]))
check("## Вопросы без ответа в базе" in report and "## Где конфигурации расходятся" in report
      and "критерий знаков" in report, "в отчёте — вопросы без ответа и сравнение конфигураций на одних вопросах")
sys.path.insert(0, str(W))
import eval_search  # noqa: E402
check(abs(eval_search.sign_test_p(7, 2) - 0.1797) < 1e-3 and eval_search.sign_test_p(0, 0) == 1.0
      and eval_search.sign_test_p(9, 0) < 0.01, "критерий знаков: 7 против 2 — p ≈ 0.18, 9 против 0 — p < 0.01")
check(eval_search.hit_ks(3) == [1, 3] and eval_search.hit_ks(5) == [1, 3, 5], "hit@K без повтора при top_k = 3")

# ядро проверки качества в интерфейсе (mint/evaluation.py): таблица вопросов, разметка, ответы чата
snap = app.load_kb_snapshot(cfg)
check(app.known_files(snap) == ["Регламент_SLA.docx", "Регламент_ИБ.docx"], f"документы базы: {app.known_files(snap)}")
na = app.NO_ANSWER
qs, issues = app.rows_to_questions([
    {"use": True, "question": "Сколько хранятся журналы?", "file": "Регламент_ИБ.docx", "section": "Журналирование",
     "contains": "не менее трёх лет", "note": float("nan")},
    {"use": True, "question": "Курс валют?", "file": na, "section": "x", "contains": "y"},
    {"use": False, "question": "Снятый вопрос", "file": "Регламент_ИБ.docx"},
    {"use": None, "question": "Без документа", "file": float("nan")},
    {"use": True, "question": "  ", "file": "Регламент_ИБ.docx"},
    {"use": True, "question": "Фраза мимо", "file": "Регламент_ИБ.docx", "contains": "такой фразы нет"},
    {"use": True, "question": "Раздел мимо", "file": "Регламент_SLA", "section": "Несуществующий раздел"},
    {"use": True, "question": "Нет файла", "file": "Другой.docx"},
])
check([q.n for q in qs] == [1, 2, 6, 7, 8] and issues == {4: "не выбран документ"},
      f"строки таблицы → вопросы: {[q.n for q in qs]}, {issues}")
check(qs[0].note == "" and not qs[1].answerable and qs[1].sections == [] and qs[1].contains == [],
      "пустая ячейка (NaN) → пусто; «ответа в базе нет» — без раздела и фразы")
probs = app.markup_problems(qs, snap)
check(set(probs) == {6, 7, 8} and "фраза не найдена" in probs[6] and "раздел не найден" in probs[7]
      and "нет в базе" in probs[8], f"разметка, которая никогда не засчитается: {probs}")
rows_back = app.question_rows([app.Question(1, "Q", ["Регламент_ИБ"], [], ["трёх лет"]), app.Question(2, "N")], snap)
check(rows_back[0]["file"] == "Регламент_ИБ.docx" and rows_back[1]["file"] == na,
      f"вопросы прошлой проверки → строки таблицы: {rows_back}")
check(app.is_refusal("В предоставленных документах информация отсутствует.")
      and app.is_refusal("В предоставленных регламентах информация отсутствует.")
      and app.is_refusal("The provided documents do not contain this information.") and not app.is_refusal("Да."),
      "отказ узнаётся на обоих языках")
ans = app.check_answers(cfg, [app.Question(1, "ВЫДУМАЙ ССЫЛКУ про журналы событий безопасности",
                                           ["Регламент_ИБ.docx"]), app.Question(2, "Курс валют на завтра?"),
                              app.Question(3, "Сколько хранятся журналы событий безопасности?", ["Регламент_ИБ"],
                                           [], ["не менее трёх лет"])],
                        app.load_glossary(cfg), repeat=2)
r1 = ans["rows"][1][0]
check(len(ans["rows"][1]) == 2 and r1["n_citations"] == 2 and len(r1["invalid"]) == 1 and r1["cites_expected"]
      and r1["reason"] == "bad_ref", f"ответ чата: ссылка на нужный документ и выдуманная ссылка — ошибка модели: {r1}")
r2, r3 = ans["rows"][2][0], ans["rows"][3][0]
check(r2["refused"] and r2["ok"] and r2["low_conf"], f"вопрос не по базе: честное «информации нет»: {r2['answer']!r}")
check(r3["ok"] and r3["in_context"] and r3["cites_expected"] and r3["reason"] is None,
      f"верный ответ: нужный фрагмент в контексте и ссылка на него: {r3['answer']!r}")
sm = ans["summary"]
check(sm["n_answerable"] == 4 and sm["n_no_answer"] == 2 and sm["ok_answerable"] == 2 and sm["clean"] == 4
      and sm["reasons"] == {"bad_ref": 2} and sm["repeat"] == 2 and sm["prompts"] == app.prompts_version()
      and sm["scores"] == {"ВЫДУМАЙ ССЫЛКУ про журналы событий безопасности": 0.0, "Курс валют на завтра?": 1.0,
                           "Сколько хранятся журналы событий безопасности?": 1.0},
      f"сводка по ответам (вопрос × повтор), причины, версия промптов, баллы вопросов: {sm}")
prev_sm = dict(sm, scores={**sm["scores"], "Курс валют на завтра?": 0.0}, prompts="ru-00000000", ok_answerable=1)
cmp_ = app.compare_answers(sm, prev_sm)
check(cmp_["better"] == 1 and cmp_["worse"] == 0 and cmp_["same"] == 2 and cmp_["meta"] == [("prompts", "ru-00000000",
      sm["prompts"])] and not cmp_["several"] and round(cmp_["deltas"]["ok_answerable"]) == 25,
      f"сравнение ответов с прошлой проверкой: {cmp_}")
check(app.compare_answers(sm, dict(prev_sm, llm_model="other", lang="en"))["several"]
      and not app.compare_answers(sm, dict(prev_sm, lang="en"))["several"],
      "больше одного изменённого параметра замечается; язык — часть версии промптов")
check(app.latest_answers([("a", {"configs": {}}), ("b", {"answers": sm}), ("c", {"answers": prev_sm})]) is sm,
      "прошлая проверка ответов — самая свежая, где ответы проверялись")
Q = app.Question
base_res = {"answer": "Хранятся три года [Источник: Регламент ИБ.docx, 1.2].", "error": False, "low_conf": False,
            "citations": [{"text": "[Источник: Регламент ИБ.docx, 1.2]", "ok": True}], "citations_invalid": [],
            "sources": ["📄 Регламент_ИБ.docx › 1.2"], "context": "=== ИСТОЧНИК: Регламент_ИБ.docx ===\nТри года."}
q_ans, q_none = Q(1, "q", ["Регламент_ИБ"], [], ["три года"]), Q(2, "q2")
check(app.judge_answer(q_ans, base_res)["ok"], "[unit] ссылка на нужный документ (имя записано иначе) — верно")
check(app.judge_answer(q_ans, dict(base_res, sources=["📄 Другой.docx"]))["reason"] is None,
      "[unit] поиск не нашёл, но ответ ссылается на нужный документ — засчитан")
check(app.judge_answer(q_ans, dict(base_res, sources=["📄 Другой.docx"], answer="Нет ссылки."))["reason"] == "search",
      "[unit] нужного фрагмента нет в контексте — ошибка поиска")
no_info_ = dict(base_res, answer="В предоставленных регламентах информация отсутствует.")
check(app.judge_answer(q_ans, no_info_)["reason"] == "no_info" and app.judge_answer(q_none, no_info_)["ok"]
      and app.judge_answer(q_none, base_res)["reason"] == "answered",
      "[unit] «информации нет»: ошибка, если ответ есть; верно, если нет; ответ не по базе — ошибка")
check(app.judge_answer(q_ans, dict(base_res, answer="❌ [ТАЙМАУТ MINT]", error=True))["reason"] == "error",
      "[unit] сбой модели")
gen, failed = app.generate_questions(cfg, 2, snap=snap)
gen2, _ = app.generate_questions(cfg, 2, snap=snap, exclude={g["pid"] for g in gen})
check(gen and not {g["pid"] for g in gen} & {g["pid"] for g in gen2}, "«Ещё вопросы» — по другим фрагментам")
with app.request_cache() as stats_:
    app.get_query_embedding("один и тот же запрос", cfg)
    app.get_query_embedding("один и тот же запрос", cfg)
check(stats_["embed_calls"] == 1, f"повторный эмбеддинг внутри прогона не запрашивается: {stats_}")

run([str(W / "draft.csv"), "--out", str(W / "run2"), "--compare", str(W / "run1" / "summary.json"), "--multi-query"],
    "прогон черновика со сравнением и multi-query")

run([str(W / "eval" / "questions.example.csv"), "--out", str(W / "ans1"), "--answers", "--quick"],
    "ответы чата из консоли (--answers --quick)")
run([str(W / "eval" / "questions.example.csv"), "--out", str(W / "ans2"), "--answers", "--quick", "--repeat", "2",
     "--compare", str(W / "ans1" / "summary.json")], "ответы чата с повторами и сравнением")
a_rep = (W / "ans2" / "report.md").read_text(encoding="utf-8") if (W / "ans2" / "report.md").exists() else ""
a_sum = json.loads((W / "ans2" / "summary.json").read_text(encoding="utf-8")) if (W / "ans2" / "summary.json").exists() else {}
check("## Ответы чата" in a_rep and "### Сравнение ответов с прошлой проверкой" in a_rep
      and "Изменилось: повторов: 1 → 2" in a_rep and app.prompts_version() in a_rep,
      "отчёт: ответы чата, версия промптов, сравнение с прошлым прогоном", a_rep[a_rep.find("## Ответы чата"):][:600])
check(a_sum.get("answers", {}).get("repeat") == 2 and a_sum.get("prompts") == app.prompts_version()
      and (W / "ans2" / "answers.csv").exists() and len(a_sum.get("configs", {})) <= 4,
      "итоги: ответы, повторы, версия промптов; без полного сравнения настроек")

(W / "q1251.csv").write_bytes("question,file,section,contains\nСколько хранятся журналы?,Регламент_ИБ,,трёх лет\n"
                              .encode("cp1251"))
r = run([str(W / "q1251.csv"), "--out", str(W / "run3")], "файл вопросов в Windows-1251 через запятую")
check("Сколько хранятся журналы" in (W / "run3" / "details.csv").read_text(encoding="utf-8-sig")
      if (W / "run3" / "details.csv").exists() else False, "вопрос из cp1251 прочитан")

srv.shutdown()
os.chdir(harness.ROOT)
shutil.rmtree(W, ignore_errors=True)
check.finish()
