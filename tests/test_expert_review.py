"""Разбор замечаний экспертом, сохранённые аудиты и сводка разборов: строки разбора в порядке заключения,
постоянные ключи, «Мои замечания» (Markdown, Word, документ с примечаниями), выгрузка с текстами и без, папка
аудита на диске, tools/review_stats.py.

Запуск из корня проекта:  python tests/test_expert_review.py
"""
import io
import json
import os
import subprocess
import sys
import zipfile
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parent))
import harness  # noqa: E402

W = harness.make_project(prefix="mint_review_")
os.environ.update({"MINT_LANG": "ru", "LOG_DIR": str(W / "logs"), "MINT_CONFIG": str(W / "none.yaml")})
os.chdir(W)
app = harness.load_mint(W, include_ui=False)
check = harness.Checker()
cfg = SimpleNamespace(log_path=W / "logs", llm_model="test-llm")


def finding(text, severity="СРЕДНИЙ", sections=("2. Требования",), **kw):
    return {"severity": severity, "text": text, "quote": kw.pop("quote", ""), "sections": list(sections),
            "origins": kw.pop("origins", ["s1"]), **kw}


f_main1 = finding("Не указан срок хранения журналов событий", "КРИТИЧНЫЙ", quote="журналы хранятся",
                  origins=["s1", "deep:numbers#1"])
f_main2 = finding("Шаг 3 не связан с шагом 4: не указано, кто передаёт заявку", source="deep", check="step",
                  sections=("3. Процесс",), origins=["deep:step#1"])
f_check = finding("Сокращение ЗКС не расшифровано", "НИЗКИЙ", sections=(), source="check", origins=[])
f_focus = finding("Нет требования к шифрованию при передаче [Источник: Основание_ТЗ.pdf, 4.2]", sections=(),
                  focus_item=0, point_kind="focus")
f_unv = finding("Нарушен пункт 9.9 стандарта", unverified=["9.9"])
f_minor = finding("Указать, каким образом обеспечивается контроль", "НИЗКИЙ", sections=())
f_typo = finding("«заявкя» → «заявка»", "НИЗКИЙ", check="language", source="deep", origins=["deep:language#1"])
result = {
    "findings": [f_main1, f_main2, f_check, f_focus, f_minor, f_typo],
    "by_section": [("2. Требования", [f_main1]), ("3. Процесс", [f_main2]), (app.GENERAL_SECTION, [f_check])],
    "focus_answers": [{"item": "шифрование при передаче", "failed": False, "answer": "нет"}],
    "summary": [{"ref": "№1", "finding": f_main1, "gist": "срок хранения"}],
    "unverified": [f_unv], "minor": [f_minor], "typos": [f_typo],
    "basis": {"docs": [{"source": "Основание_ТЗ.pdf", "by": "code", "ref": {}}]},
    "report": "# ЭКСПЕРТНОЕ ЗАКЛЮЧЕНИЕ\n", "log_md": "# Лог\n", "failed_sections": [], "sections": 3,
    "special": "шифрование при передаче", "focus": "общие критерии", "prompts": "p1", "document_mode": "full",
}

print("1. [unit] Строки разбора: порядок заключения, виды проверок, ключи")
rows = app.review_rows(result)
check([r["block"] for r in rows] == ["point", "main", "main", "main", "unverified", "minor", "typo"],
      "порядок: ответы на пункты, по разделам, не подтверждено, мелкие, опечатки", [r["block"] for r in rows])
check([r["ref"] for r in rows[:4]] == ["особое внимание 1, №1", "№1", "№2", "№3"],
      "номера — как в заключении", [r["ref"] for r in rows[:4]])
check([r["kind"] for r in rows] == ["focus", "section", "deep:step", "nomodel", "section", "section", "typo"],
      "виды проверок", [r["kind"] for r in rows])
check(rows[1]["star"] and not rows[2]["star"] and rows[1]["found_by"] == 2, "⭐ «Главное» и «нашли проверок»")
check(rows[0]["basis"] and not rows[1]["basis"], "замечание со ссылкой на документ-основание отмечено")
check(rows[3]["section"] == "Документ в целом" and rows[1]["quote"] == "журналы хранятся",
      "раздел «Документ в целом» и цитата")
check(len({r["key"] for r in rows}) == len(rows) and app.review_rows(result)[2]["key"] == rows[2]["key"],
      "ключи разные и не меняются от прогона к прогону")
f_main0 = finding("Новое замечание после перепроверки", sections=("1. Общие",))
result2 = dict(result, by_section=[("1. Общие", [f_main0])] + result["by_section"],
               findings=[f_main0] + result["findings"])
rows2 = app.review_rows(result2)
check(rows2[3]["key"] == rows[2]["key"] and rows2[3]["ref"] == "№3",
      "номер сдвинулся после перепроверки — ключ тот же (отметка эксперта не теряется)")

print("2. [unit] Отметки, «Мои замечания», выгрузка")
review = {"marks": {rows[1]["key"]: {"verdict": "ok", "comment": "согласен"},
                    rows[2]["key"]: {"verdict": "wrong", "reason": "present"},
                    rows[3]["key"]: {"verdict": "wrong"},
                    rows[5]["key"]: {"verdict": "minor"}},
          "added": [{"section": "4. Восстановление", "severity": "критичный", "text": "Нет плана восстановления"},
                    {"section": "", "severity": "", "text": "  "}]}
prog = app.review_progress(rows, review)
check(prog == {"total": 7, "marked": 4, "ok": 1, "minor": 1, "wrong": 2, "no_reason": 1, "added": 1},
      "прогресс разбора", prog)
mine = app.my_remarks(rows, review)
check([m["text"] for m in mine] == ["Не указан срок хранения журналов событий", "Нет плана восстановления"]
      and mine[1]["severity"] == "КРИТИЧНЫЙ" and mine[1]["own"], "«Мои замечания»: ✅ и свои, критичность приведена")
check(len(app.my_remarks(rows, review, include_minor=True)) == 3, "с флажком — и ➖ мелочь")
md = app.my_remarks_markdown("Документ.docx", rows, review)
check("Документ.docx" in md and "1. 🔥 **[КРИТИЧНЫЙ]** Не указан срок хранения журналов событий — *2. Требования* "
      "(согласен)" in md and "Шаг 3" not in md, "Markdown: номер, критичность, раздел, комментарий; неверных нет")
docx_bytes = app.my_remarks_docx("Документ.docx", rows, review)
with zipfile.ZipFile(io.BytesIO(docx_bytes)) as z:
    xml = z.read("word/document.xml").decode("utf-8")
check("Нет плана восстановления" in xml and "4. Восстановление" in xml and "Критичность" in xml,
      "Word: таблица с разделом и критичностью")
full = app.export_review({"id": "x1", "file": "Документ.docx", "model": "test-llm", "basis": ["Основание_ТЗ.pdf"]},
                         rows, review, texts=True)
bare = app.export_review({"id": "x1", "file": "Документ.docx", "model": "test-llm", "basis": ["Основание_ТЗ.pdf"],
                          "special": "шифрование"}, rows, review, texts=False)
check(full["format"] == "mint-review/1" and full["findings"][2]["text"].startswith("Шаг 3")
      and full["counts"]["deep:step"]["wrong"] == 1 and full["reasons"]["deep:step"] == {"present": 1}
      and full["added"][0]["text"] == "Нет плана восстановления", "выгрузка с текстами: тексты, счётчики, причины")
dump = json.dumps(bare, ensure_ascii=False)
check(not any(w in dump for w in ("Шаг 3", "Документ.docx", "журналы", "Основание_ТЗ", "шифрование",
                                  "Нет плана", "согласен", "2. Требования"))
      and bare["findings"][2]["verdict"] == "wrong" and bare["audit"]["model"] == "test-llm"
      and bare["audit"]["basis"] == 1 and bare["added"] == [{"severity": "КРИТИЧНЫЙ"}],
      "без текстов: ни текстов, ни цитат, ни разделов, ни имён файлов — только виды, отметки и причины")

print("3. [unit] Сохранённые аудиты: папка на диске, список, разбор, удаление")
doc = io.BytesIO()
import docx  # noqa: E402
d = docx.Document()
d.add_paragraph("Требования. Журналы хранятся в архиве подразделения без указания срока.")
d.add_paragraph("Восстановление после сбоя выполняет дежурный администратор по отдельной инструкции.")
d.save(doc)
aid = app.save_audit(cfg, result, "Документ.docx", doc_bytes=doc.getvalue(), settings={"use_kb": True},
                     raw_md="# Сырые ответы\nответ модели\n")
folder = W / "logs" / "audits" / aid
check(all((folder / n).is_file() for n in ("meta.json", "report.md", "log.md", "result.json", "rows.json",
                                            "review.json", "raw.md.gz", "document.docx")),
      "в папке аудита — сведения, заключение, лог, замечания, разбор, сырые ответы (сжаты), документ")
saved = app.load_audit(cfg, aid)
check(saved["meta"]["file"] == "Документ.docx" and saved["meta"]["use_kb"] is True and saved["meta"]["findings"] == 7
      and saved["rows"] == rows and saved["review"]["marks"] == {} and saved["report"].startswith("# ЭКСПЕРТНОЕ"),
      "аудит открывается с диска таким же")
check(app.load_raw(cfg, aid) == "# Сырые ответы\nответ модели\n" and app.load_document(cfg, aid) == doc.getvalue(),
      "сырые ответы и документ читаются")
check(len(saved["result"]["findings"]) == 6 and set(saved["result"]["findings"][0]) <= set(app._FINDING_FIELDS),
      "в result.json — только нужные поля замечаний")
app.save_review(cfg, aid, review)
listed = app.list_audits(cfg)
check(len(listed) == 1 and listed[0]["id"] == aid and listed[0]["progress"]["marked"] == 4,
      "в списке аудитов — прогресс разбора")
aid2 = app.save_audit(cfg, result2, "Документ.docx", aid=aid)
check(aid2 == aid and app.load_audit(cfg, aid)["review"]["marks"] == review["marks"]
      and len(app.load_audit(cfg, aid)["rows"]) == 8, "перепроверка — в ту же папку, разбор эксперта остался")
annotated, stats = app.annotate_docx_with_findings(doc.getvalue(), app.accepted_findings(rows, review))
with zipfile.ZipFile(io.BytesIO(annotated)) as z:
    names = z.namelist()
    allxml = "".join(z.read(n).decode("utf-8", "replace") for n in names if n.endswith(".xml"))
check("Нет плана восстановления" in allxml and "Не указан срок хранения" in allxml and stats["bound"] >= 1,
      "документ с замечаниями эксперта: ✅ и свои — примечаниями", stats)
try:
    app.load_audit(cfg, "../../etc")
    check(False, "чужой путь вместо id аудита отклоняется")
except ValueError:
    check(True, "чужой путь вместо id аудита отклоняется")
aid3 = app.save_audit(cfg, result, "Второй.md")
check(len(app.list_audits(cfg)) == 2 and app.list_audits(cfg)[0]["id"] == aid3 and app.load_document(cfg, aid3) is None,
      "новый аудит — сверху; у .md копии документа нет")
app.delete_audit(cfg, aid3)
check([m["id"] for m in app.list_audits(cfg)] == [aid], "аудит удаляется")

print("4. Сводка разборов: tools/review_stats.py")
out_dir = W / "reviews"
out_dir.mkdir()
(out_dir / "one.json").write_text(json.dumps(full, ensure_ascii=False), encoding="utf-8")
later = dict(full, exported="2099-01-01 00:00")
(out_dir / "all.jsonl").write_text(json.dumps(later, ensure_ascii=False) + "\n"
                                   + json.dumps(dict(bare, audit={"id": "x2", "model": "other-llm", "profile": "architecture"}),
                                                ensure_ascii=False) + "\n", encoding="utf-8")
tool = harness.ROOT / "tools" / "review_stats.py"
r = subprocess.run([sys.executable, str(tool), str(out_dir)], capture_output=True, text=True, encoding="utf-8",
                   env={**os.environ, "PYTHONUTF8": "1"})
txt = r.stdout
check(r.returncode == 0 and "Аудитов: 2 (с текстами — 1)" in txt, "один аудит в двух файлах — один раз", txt[:300])
check("| Глубокая: шаги процесса | 2 | 2 | 0 | 0 | 2 | 0% | 0% |" in txt
      and "| в документе это есть | 2 |" in txt and "| без причины | 2 |" in txt and "## По моделям" in txt
      and "## По профилям проверки" in txt and "| architecture |" in txt and "| general |" in txt,
      "по видам проверок, причинам, моделям и профилям проверки (без профиля — универсальный)", txt)
check("Шаг 3 не связан" in txt and "Нет плана восстановления" in txt, "примеры текстов — из разборов с текстами")
r0 = subprocess.run([sys.executable, str(tool), str(out_dir), "--examples", "0"], capture_output=True, text=True,
                    encoding="utf-8", env={**os.environ, "PYTHONUTF8": "1"})
check("Шаг 3 не связан" not in r0.stdout and "Нет плана восстановления" not in r0.stdout, "--examples 0 — без текстов")
r1 = subprocess.run([sys.executable, str(tool), str(W / "nothing")], capture_output=True, text=True, encoding="utf-8",
                    env={**os.environ, "PYTHONUTF8": "1"})
check(r1.returncode == 1, "нет разборов — код 1")

check.finish()
