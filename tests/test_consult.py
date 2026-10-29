"""Консультант (mint/consult.py и вкладка «Консультант»): разбор запроса, поиск по аспектам, разделы целиком,
карточки систем, бюджет по окну модели, ответ с 📄/💡, проверка ссылок по всему разговору, пункты без ссылки,
схемы Mermaid; вкладка на заглушке Streamlit.

Запуск из корня проекта:  python tests/test_consult.py
"""
import json
import os
import shutil
import sys
import traceback
from dataclasses import replace
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import harness  # noqa: E402  (заглушки и пути — до импорта mint)
import mock_server  # noqa: E402
import docx as pydocx  # noqa: E402

WORK = harness.make_project(prefix="mint_consult_")
srv = mock_server.start()
BASE = f"http://127.0.0.1:{srv.server_address[1]}/v1"
os.environ.update({
    "LLM_API_URL": BASE, "EMBED_API_URL": BASE + "/embeddings", "RERANK_API_URL": BASE + "/rerank",
    "LLM_MODEL": "test-llm", "EMBED_MODEL": "emb-model", "RERANK_MODEL": "rr-model", "LLM_API_KEY": "k",
    "KB_DIR": str(WORK / "kb"), "LOG_DIR": str(WORK / "logs"), "MINT_CONFIG": str(WORK / "config.yaml"),
    "GLOSSARY_PATH": str(WORK / "none.json"), "MINT_LANG": "ru",
})
os.chdir(WORK)
# мелкая нарезка: разделы документов — из нескольких фрагментов, чтобы было что «брать целиком»
(WORK / "config.yaml").write_text("indexing:\n  parent_size: 300\n  parent_overlap: 40\n  child_size: 120\n"
                                  "  child_overlap: 20\n", encoding="utf-8")
app = harness.load_mint(WORK)
check = harness.Checker()


def make_docx(path, blocks):
    d = pydocx.Document()
    for kind, val in blocks:
        if kind == "h1":
            d.add_heading(val, level=1)
        elif kind == "h2":
            d.add_heading(val, level=2)
        else:
            d.add_paragraph(val)
    d.save(path)
    return path.read_bytes()


QUEUE = ("Очереди сообщений шлюза обмена разделены по направлениям. Входящая очередь принимает сообщения от "
         "внешних участников, исходящая — передаёт их во внутренние системы. ")
docs = {
    "Описание_шлюза.docx": [
        ("h1", "1. Общие сведения"),
        ("p", "Шлюз обмена принимает сообщения внешних участников и передаёт их во внутренние системы организации."),
        ("h1", "2. Состав шлюза обмена"),
        ("h2", "2.1 Приём сообщений"),
        ("p", "Приём сообщений выполняет сервис приёма: он проверяет формат, подпись и размер сообщения и записывает "
              "его во входящую очередь. Сообщение больше 50 Мб отклоняется с кодом ошибки."),
        ("h2", "2.2 Очереди сообщений"),
        ("p", QUEUE * 4),
        ("p", "Глубина очереди контролируется: при превышении 10 000 сообщений приём временно замедляется."),
        ("h2", "2.3 Хранение"),
        ("p", "Обработанные сообщения хранятся в архиве 5 лет; журнал обработки — 3 года."),
        ("h1", "3. Мониторинг"),
        ("p", "Мониторинг шлюза показывает длину очередей, время обработки и число ошибок за сутки."),
    ],
    "Регламент_эксплуатации.docx": [
        ("h1", "1. Резервное копирование"),
        ("p", "Резервные копии архива сообщений создаются ежедневно и хранятся 30 дней."),
    ],
    "Карточка_Шлюз.docx": [
        ("p", "Названия: Шлюз обмена, ШО"),
        ("h1", "Карточка системы «Шлюз обмена»"),
        ("p", "Сейчас: приём сообщений до 50 Мб, две очереди. В планах на 2027 год: третья очередь для больших файлов."),
    ],
}
base_cfg, _ = app.load_settings()
cfg, _ = app.resolve_settings(base_cfg, app.resolve_model_limits(base_cfg, {}, None, {}), {})
for name, blocks in docs.items():
    ok, msg = app.process_and_add_to_kb(name, make_docx(WORK / name, blocks), cfg)
    assert ok, msg
app.save_roles(cfg, [], ["Карточка_Шлюз.docx"])
C = app  # краткость

print("1. [unit] Разбор запроса, тип запроса, схемы, пункты без ссылки, бюджет")
plan = C.parse_plan('```json\n{"kind": "diagram", "systems": ["Шлюз обмена"], "queries": ["состав шлюза", '
                    '"Состав шлюза!", "очереди", "мониторинг", "хранение", "лишний"]}\n```', "Нарисуй схему шлюза", 4)
check(plan["ok"] and plan["kind"] == "diagram" and plan["systems"] == ["Шлюз обмена"]
      and plan["queries"] == ["Нарисуй схему шлюза", "состав шлюза", "очереди", "мониторинг"],
      f"разбор: сначала сам запрос, повторы убраны, не больше 4 своих запросов: {plan['queries']}")
bad = C.parse_plan("❌ [ОШИБКА]: шлюз недоступен", "Почему сообщения идут медленно? Проблема с очередью", 4)
check(not bad["ok"] and bad["queries"] == ["Почему сообщения идут медленно? Проблема с очередью"]
      and bad["kind"] == "problem", "разбор не удался — поиск по самому запросу, тип — по словам")
check(C.parse_plan('{"kind": "чепуха", "queries": "не список"}', "Как устроен шлюз?", 4)["kind"] == "explain"
      and C.guess_kind("Нарисуй схему") == "diagram" and C.guess_kind("Сравни два подхода") == "compare",
      "неизвестный тип — по словам запроса")
text_ = ("Вступление без раздела.\n### 📄 Как устроено сейчас\n- Сервис приёма проверяет подпись [Источник: a.docx, 2.1]\n"
         "- Очереди разделены по направлениям, их две и они независимы.\n- Ниже — компоненты:\n"
         "- 💡 Вероятно, очередь общая для всех потоков данных шлюза.\n- Во фрагментах ничего не сказано про балансировку.\n"
         "```mermaid\nsequenceDiagram\n  A->>B: \"очень длинная строка схемы без ссылки\"\n```\n"
         "### Варианты решения\n- Разделить очереди по потокам — без ссылки, но это не раздел 📄.\n"
         "**📄 Ещё по документам**\n- Архив хранится пять лет, журнал обработки три года.\n")
marked, n_ = C.mark_unsupported(text_)
check(n_ == 2 and "независимы. ❔" in marked and "три года. ❔" in marked and "[Источник: a.docx, 2.1]\n" in marked
      and "по потокам — без ссылки, но это не раздел 📄." in marked and "💡 Вероятно" in marked
      and "балансировку.\n" in marked and "без ссылки\"\n" in marked,
      f"❔ — только у пунктов раздела 📄 без ссылки (не вступление, не 💡, не «ничего не», не код): {n_}")
two = C.diagrams("```mermaid\nflowchart LR\n  A-->B\n```\nтекст\n``` Mermaid\nsequenceDiagram\n  A->>B: x\n```\n```python\nx=1\n```")
check(len(two) == 2 and two[0].startswith("flowchart") and two[1].startswith("sequenceDiagram"), "код схем Mermaid")
small = replace(cfg, llm_context_tokens=8000, chars_per_token=2.5, consult_max_tokens=4000, consult_context_max_chars=60000)
check(C.context_budget(small, 1000) == int((8000 - 4000 - C.RESERVE_TOKENS) * 2.5 * C.LIMIT_SAFETY) - 1000
      and C.context_budget(replace(small, llm_context_tokens=1_000_000), 1000) == 60000
      and C.context_budget(replace(small, llm_context_tokens=4500), 1000) == C.MIN_CONTEXT_CHARS,
      "бюджет: по окну модели, не больше потолка и не меньше минимума")
store = {f"p{i}": {"source": "d", "heading_path": p, "text": "x" * 100}
         for i, p in enumerate(["2 Состав", "2 Состав → 2.1 Приём", "2 Состав → 2.2 Очереди", "2 Состав → 2.2 Очереди",
                                "2 Состав → 2.2 Очереди → 2.2.1 Глубина", "2 Состав → 2.3 Хранение", "3 Мониторинг"])}
order = {p: i for i, p in enumerate(store)}
by_source = {"d": list(store)}
found = {"id": "p2", "ids": ["p2"]}
check(C.section_ids(found, by_source, store, order, 10_000) == ["p2", "p3", "p4"]
      and C.section_ids({"id": "p0", "ids": ["p0"]}, by_source, store, order, 10_000) == ["p0", "p1", "p2", "p3", "p4", "p5"]
      and C.section_ids({"id": "p0", "ids": ["p0"]}, by_source, store, order, 250) == ["p0", "p1"]
      and C.section_ids({"id": "p3", "ids": ["p3"]}, by_source, store, order, 250) == ["p3", "p4"],
      "раздел целиком — с подразделами, без соседей; длинный — окно вокруг найденного (сначала текст после)")

print("2. [интеграция] Ответ консультанта: аспекты, карточка, разделы целиком, ссылки, ❔, схема")
mock_server.STATE["consult_plan_reply"] = json.dumps({"kind": "diagram", "systems": ["Шлюз обмена"],
                                                     "queries": ["очереди сообщений шлюза", "мониторинг шлюза",
                                                                 "приём сообщений"]}, ensure_ascii=False)
stages = []
req = "Нарисуй схему прохождения сообщения через шлюз обмена"
res = C.answer_consult(req, cfg, progress=lambda stage, **kw: stages.append(stage))
sent = mock_server.STATE["last_chat"]["messages"]
check(not res["error"] and res["kind"] == "diagram" and res["queries"][0] == req and len(res["queries"]) == 4
      and len(res["aspects"]) == 4 and all(a["found"] for a in res["aspects"]),
      f"разбор запроса и поиск по каждому аспекту: {[len(a['found']) for a in res['aspects']]}")
check(stages == ["plan"] + ["search"] * 4 + ["answer"], f"ход работы: {stages}")
kinds_ = [b["kind"] for b in res["blocks"]]
check(kinds_[0] == "card" and res["blocks"][0]["source"] == "Карточка_Шлюз.docx",
      f"карточка упомянутой системы — первой: {kinds_}")
check(any(b["expanded"] for b in res["blocks"]) and res["context"].count("входящая очередь") == 0
      and res["context"].count("Входящая очередь принимает") >= 3 and "10 000 сообщений" in res["context"],
      "найденный раздел — целиком: весь текст «2.2 Очереди сообщений», а не один фрагмент")
check(res["context_chars"] <= res["budget"] and res["budget"] >= C.MIN_CONTEXT_CHARS, "фрагменты — в пределах бюджета")
check("📄 Как устроено сейчас" in sent[0]["content"] and "--- ФРАГМЕНТЫ БАЗЫ ЗНАНИЙ ---" in sent[-1]["content"]
      and "(тип запроса по разбору: diagram" in sent[-1]["content"] and "Карточка системы" in sent[-1]["content"],
      "модель получила промпт консультанта, фрагменты и тип запроса")
check(res["citations_invalid"] == [] and res["unsupported"] == 1 and "получателю. ❔" in res["shown"]
      and "❔" not in res["answer"], "ссылка на найденный фрагмент верна; пункт 📄 без ссылки помечен ❔ (в ответе — нет)")
check(len(res["diagrams"]) == 1 and res["diagrams"][0].startswith("sequenceDiagram"), "схема Mermaid выделена")
check(res["prompts"].startswith("ru-consult-") and res["prompts"] == C.consult_prompts_version()
      and C.prompts_version() == "ru-8fd10871", f"своя версия текстов консультанта; версия чата прежняя: {res['prompts']}")
check(sent[0]["content"].startswith("Ты — ведущий ИИ-архитектор и консультант"), "профиль по умолчанию — роль архитектора")
with C.use_profile("general"):
    C.answer_consult(req, cfg)
    check(mock_server.STATE["last_chat"]["messages"][0]["content"].startswith("Ты — опытный консультант"),
          "универсальный профиль — своя роль")

print("3. [интеграция] Уточнения: история разговора, ссылки на прошлые фрагменты, мало места, разделы не целиком")
messages = [{"role": "user", "content": req}, {"role": "assistant", "content": res["answer"]}]
history = C.chat_history(messages, cfg.consult_history_turns, answer_chars=C.HISTORY_ANSWER_CHARS)
mock_server.STATE["consult_plan_reply"] = json.dumps({"kind": "explain", "systems": [], "queries": ["резервные копии"]})
old_block = C.render("context_block", header="Старый_документ.docx | раздел: 1. Очереди", body="Очередь одна.")
mock_server.STATE["consult_reply"] = ("### 📄 Как устроено сейчас\n- Очередь одна [Источник: Старый_документ.docx]\n"
                                      "### Ответ\n- 💡 Добавить очередь для больших файлов.")
with_prev = C.answer_consult("А как копируется архив?", cfg, history=history, prev_context=res["context"] + "\n\n" + old_block)
sent = mock_server.STATE["last_chat"]["messages"]
check(len(sent) == 4 and sent[2]["content"] == res["answer"], "прошлая реплика целиком — в запросе к модели")
check(with_prev["citations_invalid"] == [], "ссылка на фрагмент из прошлой реплики — не ошибка")
without = C.answer_consult("А как копируется архив?", cfg, history=history)
check(without["citations_invalid"] == ["[Источник: Старый_документ.docx]"] and "⚠️" in without["shown"],
      "без прошлых фрагментов та же ссылка — мимо найденного (⚠️)")
mock_server.STATE["consult_reply"] = None
plan_ = C.parse_plan('{"kind": "diagram", "queries": ["очереди сообщений", "приём сообщений"]}', req, 4)
big = C.gather_context(req, plan_, cfg, C.load_glossary(cfg), 60000)
few = C.gather_context(req, plan_, cfg, C.load_glossary(cfg), 600)
check(len(few["blocks"]) == 1 and few["chars"] <= 650 and len(big["blocks"]) > 1 and big["chars"] > few["chars"],
      f"мало места — фрагментов меньше, текст в пределах бюджета: {few['chars']} симв. против {big['chars']}")
flat = C.answer_consult(req, replace(cfg, consult_expand_sections=False))
check(not any(b["expanded"] for b in flat["blocks"]), "«разделы целиком» выключено — только найденные фрагменты")
mock_server.STATE["consult_plan_reply"] = "не JSON"
fallback = C.answer_consult("Нарисуй схему очередей", cfg)
check(not fallback["plan_ok"] and fallback["queries"] == ["Нарисуй схему очередей"] and fallback["kind"] == "diagram"
      and fallback["blocks"], "разбор не удался — ответ всё равно есть (поиск по самому запросу)")
mock_server.STATE["consult_plan_reply"] = None

print("4. [интеграция] Пустая база, английский язык")
empty_cfg = replace(cfg, kb_dir=str(WORK / "kb_empty"))
empty = C.answer_consult("Как устроен шлюз?", empty_cfg)
sent = mock_server.STATE["last_chat"]["messages"]
check(not empty["error"] and empty["low_found"] and not empty["blocks"]
      and C.prompt_text("consult_nothing_found") in sent[-1]["content"]
      and C.prompt_text("consult_low_found") in sent[-1]["content"], "пустая база: модель знает, что ничего не нашлось")
with C.use_lang("en"):
    en = C.answer_consult("Draw a diagram of the message path through the gateway", cfg)
    sent = mock_server.STATE["last_chat"]["messages"]
    check(not en["error"] and en["kind"] == "diagram" and en["prompts"].startswith("en-consult-")
          and "📄 How it works now" in sent[0]["content"] and "--- KNOWLEDGE BASE FRAGMENTS ---" in sent[-1]["content"]
          and en["diagrams"] and en["unsupported"] == 1, "по-английски: промпты, схема, ❔")

print("5. [сквозной] Вкладка «Консультант» на заглушке Streamlit")
import streamlit as st  # noqa: E402  (заглушка)


def run_ui(chat=None, buttons=()):
    st.LOG.clear()
    st.SCRIPT["chat_input"] = chat
    st.SCRIPT["buttons"] = set(buttons)
    st.SCRIPT["file_uploader"] = {}
    try:
        app.main()
        return "ok"
    except st.RerunException:
        return "rerun"


try:
    check(run_ui() == "ok" and ("expander", "⚙️ Настройки консультанта") in st.LOG
          and ("button", "➕ Новый разговор") in st.LOG
          and any(x[0] == "markdown" and "mint-empty__examples" in str(x[1]) for x in st.LOG)
          and sum(1 for x in st.LOG if x[0] == "slider" and str(x[1]).startswith("Сторон вопроса для поиска")) == 1,
          "вкладка: настройки, «Новый разговор», примеры запросов")
    check(run_ui(chat={"consult_field": req}) == "ok", "запрос консультанту")
    cm = st.session_state["consult_messages"]
    a0 = cm[-1] if cm else {}
    labels_ = [str(x[1]) for x in st.LOG if x[0] == "status.update"]
    check(len(cm) == 2 and a0.get("kind") == "diagram" and a0.get("diagrams") and a0.get("context")
          and a0.get("prompts", "").startswith("ru-consult-") and not st.session_state.get("messages"),
          f"ответ в разговоре консультанта, чат не тронут: {a0.get('kind')}, схем {len(a0.get('diagrams') or [])}")
    check(any(s.startswith("Ищу в базе: 1 из") for s in labels_) and any(s.startswith("Готово за") for s in labels_)
          and ("download_button", "⬇ Схема 1 (.mmd)") in st.LOG and ("expander", "🔎 Что нашёл в вики") in st.LOG
          and any(x[0] == "caption" and str(x[1]).startswith("❔ — утверждение") for x in st.LOG)
          and ("button", "👍") in st.LOG, f"ход работы, схема для скачивания, «Что нашёл в вики», ❔, оценки: {labels_}")
    fb_file = WORK / "logs" / "feedback.jsonl"
    check(run_ui(buttons={f"fb_down_consult_{a0['id']}"}) == "ok" and fb_file.exists(), "оценка 👎 ответа консультанта")
    fb_ = json.loads(fb_file.read_text(encoding="utf-8").splitlines()[-1])
    check(fb_.get("kind") == "consult" and fb_.get("request_kind") == "diagram" and fb_.get("question") == req
          and fb_.get("prompts") == a0["prompts"]
          and fb_.get("queries") and fb_.get("settings", {}).get("consult", {}).get("consult_queries") == 4,
          f"в отзыве — запрос, ответ, поисковые запросы, версия текстов консультанта: {sorted(fb_)}")
    st.session_state["mint_manual"]["consult_queries"] = 1
    mock_server.STATE["consult_plan_reply"] = json.dumps({"kind": "explain", "queries": ["резервные копии", "архив"]},
                                                        ensure_ascii=False)
    check(run_ui(chat={"consult_field": "А как копируется архив сообщений?"}) == "ok", "уточнение")
    a1 = st.session_state["consult_messages"][-1]
    sent_ = mock_server.STATE["last_chat"]["messages"]
    check(len(a1.get("queries") or []) == 2 and len(sent_) == 4 and sent_[2]["content"] == a0["content"],
          f"уточнение: настройка сеанса (1 сторона + сам запрос), прошлая реплика — в запросе: {a1.get('queries')}")
    mock_server.STATE["consult_plan_reply"] = None
    st.session_state["mint_manual"].pop("consult_queries")
    check(run_ui(chat="Как устроен шлюз обмена?") == "ok" and len(st.session_state["consult_messages"]) == 4
          and st.session_state["messages"] and st.session_state["messages"][-1]["role"] == "assistant",
          "вопрос во вкладке «Чат» — в чат, разговор консультанта не тронут")
    check(run_ui(buttons={"consult_clear_btn"}) == "ok" and st.session_state["consult_messages"] == []
          and st.session_state["messages"], "«Новый разговор» очищает только консультанта")
except Exception:
    traceback.print_exc()
    check(False, "исключение во вкладке «Консультант»")

srv.shutdown()
os.chdir(harness.ROOT)
shutil.rmtree(WORK, ignore_errors=True)
check.finish()
